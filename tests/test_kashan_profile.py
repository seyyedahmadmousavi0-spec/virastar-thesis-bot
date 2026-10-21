"""Independent checks of authority isolation, evidence and immutable documents."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document

from document_blocks import extract_document_blocks
from doc_guard import sha256_file
import kashan_checker


ROOT = Path(__file__).resolve().parents[1]


class KashanProfileTests(unittest.TestCase):
    def base_rules(self):
        return json.loads((ROOT / "rules.json").read_text(encoding="utf-8"))

    def profile(self, degree="master", scope="full"):
        mode = {"master": "kashan_master", "phd": "kashan_phd", "unknown": "kashan"}[degree]
        rules = kashan_checker.apply_profile(self.base_rules(), mode)
        rules["document_scope"] = scope
        return rules

    def review(self, texts, degree="master", scope="full", **stats):
        doc = Document()
        for text in texts:
            doc.add_paragraph(text)
        blocks = extract_document_blocks(doc)
        issues = kashan_checker.check_document(doc, blocks, stats, self.profile(degree, scope))
        return doc, blocks, issues, stats

    def test_profile_copy_does_not_modify_default_or_retain_tums_overrides(self):
        original = self.base_rules()
        original["tums_special_rules"] = {"et_al_forbidden": True}
        original["checklist_items"] = ["a TUMS-only rule"]
        snapshot = copy.deepcopy(original)
        result = kashan_checker.apply_profile(original, "kashan_master")
        result["modes"]["thesis"]["required_sections"].append({"name": "example", "patterns": ["example"]})
        self.assertEqual(original, snapshot)
        self.assertNotIn("tums_special_rules", result)
        self.assertNotIn("checklist_items", result)
        self.assertFalse(all(result["profile_checks"].values()))
        self.assertTrue(all(value is False for value in result["profile_checks"].values()))

    def test_no_unconfirmed_numeric_or_fixed_chapter_requirements(self):
        for mode in kashan_checker.MODES:
            profile = kashan_checker.apply_profile(self.base_rules(), mode)
            for name in ("abstract_words", "keywords_count", "min_chapters_thesis"):
                self.assertNotIn(name, profile["limits"])
            self.assertEqual(profile["modes"][mode]["required_sections"], [])
            self.assertIsNone(profile["reference_style"]["edition"])
        _, _, issues, stats = self.review(["فصل اول", "متن پژوهش"], scope="fragment")
        self.assertFalse(any(issue.rule_id in {"chapters", "abstract_length", "keywords_count", "size_fa", "line_spacing", "page_margin"} for issue in issues))
        self.assertIn("سه", next(row["detail"] for row in stats["university_checklist"] if row["id"] == "chapter_count"))

    def test_unknown_degree_is_manual_and_never_uses_master_sequence(self):
        _, _, issues, stats = self.review(["این فقط یک متن فارسی است."], degree="unknown")
        self.assertEqual(stats["university_profile"]["degree"], "unknown")
        self.assertTrue(any(row["id"] == "degree" and row["status"] == "manual" for row in stats["university_checklist"]))
        self.assertFalse(any(issue.rule_id == "KASHAN_SECTION_REVIEW" for issue in issues))
        self.assertFalse(any(row["id"] == "defense_minutes" for row in stats["university_checklist"]))

    def test_master_dedication_is_optional_missing_sections_never_fabricated(self):
        _, _, issues, stats = self.review(["چکیده", "متن چکیده", "منابع"])
        rows = {row["id"]: row for row in stats["university_checklist"]}
        self.assertEqual(rows["dedication"]["status"], "optional")
        self.assertEqual(rows["appendices"]["status"], "optional")
        self.assertEqual(rows["figures_list"]["status"], "manual")
        self.assertFalse(any("تقدیم»" in issue.message for issue in issues))
        self.assertTrue(all(issue.target_text == "" for issue in issues))
        self.assertTrue(all(issue.location_status != "exact" for issue in issues))

    def test_fragment_and_unreadable_input_do_not_report_missing_pages(self):
        _, _, issues, stats = self.review(["این بخشی از پایان‌نامه است."], scope="fragment")
        self.assertEqual(issues, [])
        self.assertFalse(any(row["status"] == "not_detected" for row in stats["university_checklist"]))
        stats = {"analysis_status": "failed", "is_pdf": True}
        issues = kashan_checker.check_document(None, [], stats, self.profile("phd"))
        self.assertEqual(issues, [])
        self.assertEqual(stats["university_profile"]["document_scope"], "fragment")

    def test_pdf_blocks_can_show_presence_without_claiming_absent_pages(self):
        stats = {"is_pdf": True}
        blocks = [{"text": "Abstract", "block_id": "pdf:block1", "story": "pdf",
                   "source_part": "pdf/page/2", "page_index": 1,
                   "location": "صفحه ۲", "extraction_status": "pdf_reading_order_unverified"}]
        issues = kashan_checker.check_document(None, blocks, stats, self.profile("master"))
        self.assertFalse(any(issue.rule_id == "KASHAN_SECTION_REVIEW" for issue in issues))
        rows = {row["id"]: row for row in stats["university_checklist"]}
        self.assertEqual(rows["abstract_en"]["status"], "detected")
        self.assertEqual(rows["abstract_fa"]["status"], "manual")
        self.assertEqual(stats["university_profile"]["document_scope"], "full")

    def test_master_and_phd_degree_findings_are_precisely_located_in_cover(self):
        for degree, wrong, correct in (("phd", "Thesis", "Ph.D. Dissertation"), ("master", "Dissertation", "Thesis")):
            _, blocks, issues, stats = self.review(["University of Kashan", wrong, "Faculty of Engineering", "Supervisor: Example"], degree=degree, scope="fragment")
            matches = [issue for issue in issues if issue.rule_id == "KASHAN_ENGLISH_DEGREE"]
            self.assertEqual(len(matches), 1)
            issue = matches[0]
            self.assertEqual(issue.target_text, wrong)
            self.assertEqual(issue.location_status, "exact")
            self.assertEqual(issue.confidence, "B")
            self.assertEqual(issue.block_id, blocks[1].block_id)
            self.assertIn(correct, issue.suggestion)
            self.assertFalse(issue.auto_applied)
            self.assertTrue(issue.validate_evidence(blocks[1].text))
            self.assertFalse(any(row["id"] == "defense_minutes" for row in stats["university_checklist"]) if degree == "phd" else False)

    def test_bibliographic_thesis_word_does_not_become_cover_finding(self):
        _, _, issues, _ = self.review(["References", "Smith, J. (2020). A Thesis on control. University of Kashan.", "Kashan University", "Research was submitted in 2020."], degree="phd", scope="fragment")
        self.assertFalse(any(issue.rule_id == "KASHAN_ENGLISH_DEGREE" for issue in issues))

    def test_explicit_opposite_persian_degree_has_exact_source_evidence(self):
        for degree, wrong, wanted, section in (("phd", "پایان‌نامه کارشناسی ارشد", "رساله دکتری", "۵.۲"),
                                                ("master", "رساله دکتری", "پایان‌نامه کارشناسی ارشد", "۵.۱")):
            _, blocks, issues, _ = self.review(["دانشگاه کاشان", "دانشکده مهندسی", wrong,
                                               "استاد راهنما: نمونه"], degree=degree, scope="fragment")
            matches = [issue for issue in issues if issue.rule_id == "KASHAN_PERSIAN_DEGREE"]
            self.assertEqual(len(matches), 1)
            issue = matches[0]
            self.assertEqual(issue.target_text, wrong)
            self.assertTrue(issue.validate_evidence(blocks[2].text))
            self.assertEqual(issue.block_id, blocks[2].block_id)
            self.assertEqual(issue.confidence, "B")
            self.assertIn(wanted, issue.suggestion)
            self.assertIn("بخش " + section, issue.message)
            self.assertFalse(issue.auto_applied)

    def test_persian_degree_needs_positive_cover_context_and_unquoted_title(self):
        cases = [
            ["دانشگاه کاشان", "پایان‌نامه کارشناسی ارشد"],
            ["دانشگاه کاشان", "دانشکده مهندسی", "در این نمونه عبارت پایان‌نامه کارشناسی ارشد آمده است."],
            ["دانشگاه کاشان", "دانشکده مهندسی", "«پایان‌نامه کارشناسی ارشد»"],
            ["منابع", "دانشگاه کاشان", "دانشکده مهندسی", "پایان‌نامه کارشناسی ارشد"],
            ["دانشگاه دیگر", "دانشکده مهندسی", "پایان‌نامه کارشناسی ارشد"],
        ]
        for texts in cases:
            with self.subTest(texts=texts):
                _, _, issues, _ = self.review(texts, degree="phd", scope="fragment")
                self.assertFalse(any(issue.rule_id == "KASHAN_PERSIAN_DEGREE" for issue in issues))
        _, _, issues, _ = self.review(["دانشگاه کاشان", "دانشکده مهندسی", "رساله دکتری"], degree="unknown", scope="fragment")
        self.assertFalse(any(issue.rule_id == "KASHAN_PERSIAN_DEGREE" for issue in issues))

    def test_toc_entry_and_body_mentions_are_not_independent_headings(self):
        doc = Document()
        doc.add_paragraph("چکیده\t۱", style="TOC Heading")
        doc.add_paragraph("در این پژوهش چکیده فارسی و منابع بررسی شده‌اند.")
        doc.add_paragraph("فصل اول...........۱")
        stats = {}
        issues = kashan_checker.check_document(doc, extract_document_blocks(doc), stats, self.profile())
        rows = {row["id"]: row for row in stats["university_checklist"]}
        self.assertEqual(rows["abstract_fa"]["status"], "not_detected")
        self.assertEqual(rows["chapters"]["status"], "not_detected")
        self.assertTrue(any(issue.rule_id == "KASHAN_SECTION_REVIEW" for issue in issues))

    def test_actual_toc_heading_is_detected_but_toc_entry_is_not(self):
        doc = Document()
        doc.add_paragraph("فهرست مطالب", style="TOC Heading")
        doc.add_paragraph("چکیده ........ ۱")
        stats = {}
        kashan_checker.check_document(doc, extract_document_blocks(doc), stats, self.profile())
        rows = {row["id"]: row for row in stats["university_checklist"]}
        self.assertEqual(rows["toc"]["status"], "detected")
        self.assertEqual(rows["abstract_fa"]["status"], "not_detected")

    def test_checker_integration_keeps_default_and_tums_rules_isolated(self):
        import checker
        defaults_before = checker.load_rules()
        tums_before = checker.load_rules(mode="tums")
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "نمونه.docx"
            doc = Document()
            doc.add_paragraph("چکیده")
            doc.add_paragraph("متن چکیدهٔ کوتاه دانشگاهی با اعداد ۱۲۳ و فرمول E=mc² بدون تغییر داده‌ها.")
            doc.add_paragraph("فصل اول")
            doc.add_paragraph("این متن پژوهش دارای تعداد فصل ثابت نیست و به بررسی نویسنده نیاز دارد.")
            doc.save(source)
            rules = checker.load_rules(mode="kashan_master")
            rules["document_scope"] = "fragment"
            _, issues, stats = checker.analyze(source, mode="kashan_master", rules=rules)
        forbidden = {"page_size", "page_margin", "size_fa", "size_en", "font_fa", "font_en", "line_spacing", "abstract_length", "keywords_count", "chapters"}
        self.assertFalse(any(issue.rule_id in forbidden or issue.rule_id.startswith("TUMS_") for issue in issues))
        self.assertEqual(stats["university_profile"]["id"], "kashan")
        self.assertEqual(checker.load_rules(), defaults_before)
        self.assertEqual(checker.load_rules(mode="tums"), tums_before)

    def test_duplicate_and_alphabetic_reference_checks_use_actual_entries(self):
        texts = ["References", "Zulu, A. (2020). First research title. Journal.", "Alpha, B. (2021). Another research title. Journal.", "Alpha, B. (2021). Another research title. Journal."]
        _, blocks, issues, stats = self.review(texts, scope="fragment")
        reference_issues = [issue for issue in issues if issue.rule_id.startswith("KASHAN_REFERENCE")]
        self.assertEqual({issue.rule_id for issue in reference_issues}, {"KASHAN_REFERENCE_ALPHABETICAL", "KASHAN_REFERENCE_DUPLICATE"})
        self.assertEqual(len(reference_issues), 2)
        self.assertEqual(stats["university_reference_checks"]["recognized_entries"], 3)
        self.assertIsNone(stats["university_reference_checks"]["edition"])
        self.assertFalse(stats["university_reference_checks"]["online_verified"])
        by_id = {block.block_id: block for block in blocks}
        for issue in reference_issues:
            self.assertEqual(issue.confidence, "B")
            self.assertTrue(issue.validate_evidence(by_id[issue.block_id].text))
            self.assertIn("بخش ۵.۴", issue.message)

    def test_reference_wrapping_and_unknown_author_keep_sort_check_manual(self):
        _, _, issues, stats = self.review(["References", "Zulu, A. (2020). Study title. Publisher.", "continuation paragraph from the same source", "Alpha, B. (2021). Study title. Publisher."], scope="fragment")
        self.assertFalse(any(issue.rule_id == "KASHAN_REFERENCE_ALPHABETICAL" for issue in issues))
        self.assertFalse(stats["university_reference_checks"]["alphabetical_checked"])

    def test_duplicate_reference_evidence_does_not_make_scientific_numbers_editable(self):
        source = "Alpha, B. (2021). Study title. DOI:10.1234/example."
        _, _, issues, _ = self.review(["References", source, source], scope="fragment")
        duplicate = next(issue for issue in issues if issue.rule_id == "KASHAN_REFERENCE_DUPLICATE")
        self.assertTrue(duplicate.protected_overlap)
        self.assertFalse(duplicate.auto_applied)

    def test_persian_reference_sorting_uses_alphabet_not_codepoints(self):
        _, _, issues, stats = self.review(["منابع", "بابایی، علی. (۱۳۹۹). عنوان کتاب فارسی نمونه. ناشر.", "پارسایی، محمد. (۱۴۰۰). عنوان کتاب فارسی دیگر. ناشر.", "ترابی، رضا. (۱۴۰۱). عنوان کتاب پژوهش. ناشر."], scope="fragment")
        self.assertTrue(stats["university_reference_checks"]["alphabetical_checked"])
        self.assertFalse(any(issue.rule_id == "KASHAN_REFERENCE_ALPHABETICAL" for issue in issues))

    def test_source_a_is_excluded_and_metadata_does_not_claim_official_verification(self):
        _, _, _, stats = self.review(["نمونه"], scope="fragment")
        profile = stats["university_profile"]
        self.assertEqual(profile["source_status"], "user_provided_summary")
        self.assertEqual(profile["applied_sources"], ["B"])
        self.assertEqual(profile["excluded_sources"][0]["id"], "A")
        self.assertIn("تأیید نشده", profile["source_label"])
        self.assertTrue(all("source" in row and row["source"] for row in stats["university_checklist"]))

    def test_unicode_docx_bytes_and_scientific_text_are_unchanged(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "پژوهش دانشگاه کاشان.docx"
            document = Document()
            document.add_paragraph("دانشگاه کاشان")
            document.add_paragraph("رساله دکتری")
            document.add_paragraph("University of Kashan")
            document.add_paragraph("Thesis")
            document.add_paragraph("Faculty of Science")
            science = "ΔE = mc²; p < 0.05؛ دوز ۱۲٫۵ mg و مقدار میکرو µ محفوظ است."
            document.add_paragraph(science)
            document.save(source)
            before_hash = sha256_file(source)
            loaded = Document(source)
            before_xml = loaded.element.xml
            stats = {}
            issues = kashan_checker.check_document(loaded, extract_document_blocks(loaded), stats, self.profile("phd", "fragment"))
            self.assertTrue(any(issue.rule_id == "KASHAN_ENGLISH_DEGREE" for issue in issues))
            self.assertEqual(sha256_file(source), before_hash)
            self.assertEqual(loaded.element.xml, before_xml)
            self.assertEqual(loaded.paragraphs[-1].text, science)
            self.assertTrue(all(not issue.auto_applied for issue in issues))


if __name__ == "__main__":
    unittest.main()
