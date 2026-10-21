"""Student-facing reports retain coverage gaps, limits and annotation decisions."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement

from analysis_model import Issue
import checker
from report_artifacts import html_report, location_label, report_data


class StudentReportTests(unittest.TestCase):
    def setUp(self):
        self.rules = checker.load_rules()

    def test_strict_suppression_is_counted_without_promoting_ambiguity(self):
        candidate = Issue("TEST", message="manual", target_text="missing", char_start=0, char_end=7)
        audit = {"ambiguous_candidates": 0}
        with patch("checker._check_text", return_value=[candidate]):
            findings = checker.check_text(0, "original", "پاراگراف ۱", self.rules, audit=audit)
        self.assertEqual([], findings)
        self.assertEqual(1, audit["ambiguous_candidates"])
        summary = checker.build_summary(findings, audit, self.rules)
        self.assertIn("C، مبهم: 1", summary)
        self.assertIn("فقط برای بررسی دستی", summary)
        self.assertEqual(1, report_data(findings, audit, self.rules)["summary"]["by_confidence"]["C"])

    def test_unknown_suppression_is_not_displayed_as_zero(self):
        summary = checker.build_summary([], {}, self.rules)
        self.assertIn("C، مبهم: نامشخص", summary)
        data = report_data([], {}, self.rules)
        self.assertIsNone(data["summary"]["by_confidence"]["C"])
        self.assertIn("صفر فرض نشده", html_report(data))

    def test_full_text_report_has_no_hidden_thousand_item_cap(self):
        findings = [Issue("TEST", message=f"unique-item-{index}", confidence="B") for index in range(1001)]
        full = checker.build_report(findings, {"ambiguous_candidates": 0}, self.rules)
        self.assertIn("unique-item-1000", full)
        self.assertEqual(1001, full.count("📝 شرح:"))
        preview = checker.build_report(findings, {}, self.rules, examples=2)
        self.assertIn("999 مورد", preview)
        self.assertEqual(2, preview.count("📝 شرح:"))

    def test_manual_candidates_stay_in_report_and_default_html_filter_stays_safe(self):
        findings = [Issue("TEST", message="manual-only", confidence="C")]
        text = checker.build_report(findings, {"ambiguous_candidates": 2}, self.rules)
        self.assertIn("manual-only", text)
        self.assertIn("این مورد مبهم است", text)
        self.assertIn("C، مبهم: 3", text)
        html = html_report(report_data(findings, {"ambiguous_candidates": 2}, self.rules))
        self.assertIn('data-confidence="C"', html)
        self.assertIn("confidence!=='C'", html)

    def test_actual_word_coverage_distinguishes_absent_partial_and_analyzed(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "coverage.docx"
            doc = Document()
            p = doc.add_paragraph("متن اصلی برای بررسی.")
            p._p.append(OxmlElement("m:oMath"))
            doc.sections[0].header.paragraphs[0].text = "سرصفحه"
            doc.save(source)
            _, _, stats = checker.analyze(source, rules=self.rules)
        self.assertEqual("partial", stats["coverage"]["body"]["status"])
        self.assertEqual("analyzed", stats["coverage"]["header"]["status"])
        self.assertEqual("absent", stats["coverage"]["footnote"]["status"])
        self.assertEqual("partial", stats["analysis_status"])
        self.assertIn("بخشی از سند", checker.build_summary([], stats, self.rules))

    def test_budget_skips_are_visible_per_story(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "limited.docx"
            doc = Document()
            doc.add_paragraph("بخاطر نمونه ها")
            doc.add_paragraph("بخاطر نمونه ها")
            doc.sections[0].header.paragraphs[0].text = "بخاطر نمونه ها"
            doc.save(source)
            rules = checker.load_rules()
            rules["max_findings"] = 1
            _, _, stats = checker.analyze(source, rules=rules)
        self.assertEqual("partial", stats["analysis_status"])
        self.assertEqual("partial", stats["coverage"]["body"]["status"])
        self.assertEqual("present", stats["coverage"]["header"]["status"])
        self.assertGreater(stats["skipped_blocks"], 0)

    def test_annotation_audit_and_untrusted_text_are_preserved_safely(self):
        self.assertEqual('زیر «body»، پاراگراف 1', location_label('زیر «body»، پاراگراف 1'))
        findings = [Issue("TEST", message="<script>source</script>", loc="body / table 1 / row 2 / cell 1",
                          category="citation_consistency", confidence="B", block_id="b1"),
                    Issue("OTHER", message="second", confidence="C")]
        stats = {"ambiguous_candidates": 0,
                 "annotation_summary": {"requested": 2, "applied": 1, "unapplied": 1, "output_created": True,
                                        "private_path": "C:/private/document.docx",
                                        "reasons": {"ambiguous_excluded": 1},
                                        "records": [{"index": 0, "rule_id": "TEST", "block_id": "b1", "status": "applied", "reason": "annotated"},
                                                    {"index": 1, "rule_id": "OTHER", "status": "unapplied", "reason": "ambiguous_excluded", "private_path": "secret"}]},
                 "coverage": {"body": {"status": "analyzed", "present_blocks": 1, "analyzed_blocks": 1,
                                       "private_path": "secret"}}}
        data = report_data(findings, stats, self.rules)
        self.assertNotIn("private_path", repr(data))
        self.assertEqual("applied", data["findings"][0]["annotation"]["status"])
        html = html_report(data)
        self.assertIn("&lt;script&gt;source&lt;/script&gt;", html)
        self.assertNotIn("<script>source", html)
        self.assertIn("متن اصلی / جدول 1 / سطر 2 / سلول 1", html)
        self.assertIn("ارجاع‌دهی و فهرست منابع", html)
        self.assertIn("از 2 مورد، 1 مورد در Word کامنت شد", html)
        self.assertIn("مورد مبهم؛ فقط برای بررسی دستی", html)
        text = checker.build_report(findings, stats, self.rules)
        self.assertIn("وضعیت Word: در Word کامنت شد", text)
        summary = checker.build_summary(findings, stats, self.rules)
        self.assertIn("از 2 مورد گزارش‌شده، 1 مورد", summary)
        self.assertLess(len(summary), 650)

    def test_pdf_annotation_audit_uses_pdf_and_discloses_experimental_status(self):
        findings = [Issue("TEST", message="review", confidence="B", page_index=0)]
        stats = {"ambiguous_candidates": 0, "annotation_format": "pdf",
                 "annotation_summary": {"requested": 1, "applied": 0, "unapplied": 1,
                                        "output_created": False, "experimental": True,
                                        "reasons": {"extraction_mapping_mismatch": 1},
                                        "records": [{"index": 0, "rule_id": "TEST", "page_index": 0,
                                                     "status": "unapplied", "reason": "extraction_mapping_mismatch"}]}}
        text = checker.build_report(findings, stats, self.rules)
        summary = checker.build_summary(findings, stats, self.rules)
        html = html_report(report_data(findings, stats, self.rules))
        for result in (text, summary, html):
            self.assertNotIn("در Word", result)
            self.assertIn("در PDF", result)
        self.assertIn("آزمایشی", html)
        self.assertIn("نگاشت متن به نویسه‌های صفحه", html)


if __name__ == "__main__":
    unittest.main()
