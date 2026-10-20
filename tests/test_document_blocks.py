"""Source integrity and coverage regressions for structured Word extraction."""
import hashlib
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from document_blocks import DocumentExtractionLimit, extract_document_blocks, W, MC


class DocumentBlocksTests(unittest.TestCase):
    def add_note_part(self, doc, story, body):
        name = story + "s"
        xml = ('<w:' + name + ' xmlns:w="' + W + '">' + body + '</w:' + name + '>').encode()
        part = Part(PackURI('/word/' + name + '.xml'),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.' + name + '+xml',
                    xml, doc.part.package)
        doc.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + name)

    def add_textbox(self, parent, text):
        box = OxmlElement('w:txbxContent')
        p, r, t = OxmlElement('w:p'), OxmlElement('w:r'), OxmlElement('w:t')
        t.text = text
        r.append(t)
        p.append(r)
        box.append(p)
        parent.append(box)
        return box

    def test_body_indices_whitespace_and_xml_locations(self):
        doc = Document()
        doc.add_paragraph('  alpha\tbeta\n gamma  ')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'table text'
        doc.add_paragraph('last body')
        before = doc.element.xml
        blocks = extract_document_blocks(doc)
        body = [b for b in blocks if b.story == 'body']
        self.assertEqual([0, 1], [b.paragraph_index for b in body])
        self.assertEqual('  alpha\tbeta\n gamma  ', body[0].original_text)
        self.assertIs(body[0].paragraph_element, doc.paragraphs[0]._p)
        self.assertTrue(all(b.source_part == '/word/document.xml' for b in blocks))
        for block in blocks:
            found = doc.element.getroottree().xpath(block.xml_path, namespaces=doc.element.nsmap)
            self.assertEqual(1, len(found))
        self.assertEqual(before, doc.element.xml)
        self.assertEqual([b.block_id for b in blocks], [b.block_id for b in extract_document_blocks(doc)])

    def test_nested_tables_and_merged_cells_are_not_duplicated(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'horizontal merged'
        table.cell(1, 0).text = 'outer text'
        table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'nested text'
        table.cell(1, 1).text = 'other cell'
        blocks = extract_document_blocks(doc)
        texts = [b.text for b in blocks if b.text]
        self.assertEqual(['horizontal merged', 'outer text', 'nested text', 'other cell'], texts)
        self.assertTrue(all(b.story == 'table' and b.paragraph_index == -1 for b in blocks))
        self.assertIn('/ table 1 / row 2 / cell 1 / table 1', next(b.location for b in blocks if b.text == 'nested text'))
        self.assertEqual(len(blocks), len({b.block_id for b in blocks}))

    def test_vertical_merge_content_is_analyzed_once(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).merge(table.cell(1, 0)).text = 'vertical merged'
        blocks = extract_document_blocks(doc)
        self.assertEqual(1, sum(b.text == 'vertical merged' for b in blocks))

    def test_existing_linked_headers_and_footers_once_no_new_parts(self):
        doc = Document()
        doc.add_paragraph('body')
        doc.sections[0].header.paragraphs[0].text = 'header words'
        doc.sections[0].footer.paragraphs[0].text = 'footer words'
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        before_parts = [str(p.partname) for p in doc.part.package.parts]
        before_xml = doc.element.xml
        blocks = extract_document_blocks(doc)
        self.assertEqual(1, sum(b.text == 'header words' for b in blocks))
        self.assertEqual(1, sum(b.text == 'footer words' for b in blocks))
        self.assertEqual(before_parts, [str(p.partname) for p in doc.part.package.parts])
        self.assertEqual(before_xml, doc.element.xml)
        self.assertTrue(all(b.paragraph_index == -1 for b in blocks if b.story in {'header', 'footer'}))

    def test_missing_headers_are_not_created(self):
        doc = Document()
        doc.add_paragraph('body')
        before = [str(p.partname) for p in doc.part.package.parts]
        extract_document_blocks(doc)
        self.assertEqual(before, [str(p.partname) for p in doc.part.package.parts])
        self.assertFalse(any('header' in name or 'footer' in name for name in before))

    def test_textbox_paragraphs_are_separate_and_not_double_counted(self):
        doc = Document()
        host = doc.add_paragraph('host words')
        self.add_textbox(host.add_run()._r, 'textbox words')
        blocks = extract_document_blocks(doc)
        self.assertEqual(['host words', 'textbox words'], [b.text for b in blocks])
        self.assertEqual(['body', 'textbox'], [b.story for b in blocks])
        self.assertEqual(-1, blocks[1].paragraph_index)
        self.assertNotEqual(blocks[0].block_id, blocks[1].block_id)

    def test_alternate_textbox_fallback_does_not_double_count(self):
        doc = Document()
        host = doc.add_paragraph('host')
        alternate = etree.Element('{' + MC + '}AlternateContent')
        choice = etree.SubElement(alternate, '{' + MC + '}Choice', Requires='wps')
        fallback = etree.SubElement(alternate, '{' + MC + '}Fallback')
        self.add_textbox(choice, 'selected box')
        self.add_textbox(fallback, 'fallback duplicate')
        host.add_run()._r.append(alternate)
        blocks = extract_document_blocks(doc)
        self.assertEqual(['host', 'selected box'], [b.text for b in blocks])
        self.assertTrue(all('alternate_content' in b.extraction_status for b in blocks))

    def test_note_parts_exclude_separators_and_include_nested_tables(self):
        doc = Document()
        doc.add_paragraph('body')
        self.add_note_part(doc, 'footnote',
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t>separator</w:t></w:r></w:p></w:footnote>'
            '<w:footnote w:id="1"><w:p><w:r><w:t>footnote text</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>note table</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:footnote>')
        self.add_note_part(doc, 'endnote',
            '<w:endnote w:id="2"><w:p><w:r><w:t>endnote text</w:t></w:r></w:p></w:endnote>')
        blocks = extract_document_blocks(doc)
        self.assertNotIn('separator', [b.text for b in blocks])
        self.assertEqual({'footnote text', 'note table'}, {b.text for b in blocks if b.story == 'footnote'})
        self.assertEqual(['endnote text'], [b.text for b in blocks if b.story == 'endnote'])
        self.assertTrue(all(b.paragraph_index == -1 for b in blocks if b.story in {'footnote', 'endnote'}))

    def test_tracked_deletions_fields_and_math_are_explicitly_partial(self):
        doc = Document()
        p = doc.add_paragraph('visible ')
        deletion = OxmlElement('w:del')
        dr, dt = OxmlElement('w:r'), OxmlElement('w:delText')
        dt.text = 'deleted words'
        dr.append(dt)
        deletion.append(dr)
        p._p.append(deletion)
        inserted = OxmlElement('w:ins')
        ir, it = OxmlElement('w:r'), OxmlElement('w:t')
        it.text = 'inserted '
        ir.append(it)
        inserted.append(ir)
        p._p.append(inserted)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'private instruction')
        fr, ft = OxmlElement('w:r'), OxmlElement('w:t')
        ft.text = 'field result'
        fr.append(ft)
        field.append(fr)
        p._p.append(field)
        p._p.append(OxmlElement('m:oMath'))
        block = extract_document_blocks(doc)[0]
        self.assertEqual('visible inserted field result', block.text)
        self.assertEqual('partial:fields,math,tracked_changes', block.extraction_status)

    def test_deleted_textboxes_are_excluded(self):
        doc = Document()
        p = doc.add_paragraph('visible')
        deletion = OxmlElement('w:del')
        self.add_textbox(deletion, 'deleted textbox')
        p._p.append(deletion)
        self.assertEqual(['visible'], [b.text for b in extract_document_blocks(doc)])

    def test_source_package_bytes_and_caller_xml_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'source.docx'
            doc = Document()
            doc.add_paragraph('body')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'table'
            doc.save(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            loaded = Document(path)
            xml = loaded.element.xml
            extract_document_blocks(loaded)
            self.assertEqual(xml, loaded.element.xml)
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_budgets_raise_instead_of_silent_truncation(self):
        doc = Document()
        doc.add_paragraph('alpha')
        doc.add_paragraph('beta')
        for kwargs in ({'max_blocks': 1}, {'max_total_chars': 4}, {'max_part_bytes': 1},
                       {'max_nodes': 1}, {'max_depth': 1}):
            with self.subTest(kwargs=kwargs), self.assertRaises(DocumentExtractionLimit):
                extract_document_blocks(doc, **kwargs)
        with self.assertRaises(ValueError):
            extract_document_blocks(doc, max_blocks=0)


if __name__ == '__main__':
    unittest.main()
