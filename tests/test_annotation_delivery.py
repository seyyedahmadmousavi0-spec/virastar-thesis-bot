"""Review-copy delivery: selection order, complete accounting and source integrity."""
import copy
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from analysis_model import Issue
from doc_guard import sha256_file
from safe_annotator import annotate_document


def finding(target, paragraph=0, confidence="A", rule_id="mi_prefix", **values):
    return Issue(rule_id, message="در این عبارت، نیم‌فاصله نیازمند بازبینی است.",
                 suggestion="با توجه به بافت متن، استفاده از نیم‌فاصله را بررسی کنید.",
                 par_index=paragraph, target_text=target, confidence=confidence, **values)


class AnnotationDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="annotation_delivery_")
        self.root = Path(self.directory.name)
        self.output = self.root / "review.docx"

    def tearDown(self):
        self.directory.cleanup()

    def assert_accounted(self, document):
        summary = document._last_annotation_summary
        self.assertTrue(summary["output_created"])
        self.assertEqual(summary["requested"], summary["applied"] + summary["unapplied"])
        self.assertEqual(summary["unapplied"], sum(summary["reasons"].values()))
        self.assertEqual(summary["applied"], sum(r["status"] == "applied" for r in summary["records"]))
        self.assertEqual(list(range(summary["requested"])), [r["index"] for r in summary["records"]])
        return summary

    def test_later_a_outranks_earlier_b_across_paragraphs(self):
        document = Document()
        document.add_paragraph("این پیشنهاد اول است.")
        document.add_paragraph("عبارت دوم مهم‌تر است.")
        issues = [finding("پیشنهاد", confidence="B"), finding("عبارت دوم", paragraph=1)]
        original_xml = document.element.xml
        self.assertTrue(annotate_document(document, issues, self.output, max_comments=1))
        self.assertEqual(document.element.xml, original_xml)
        result = Document(self.output)
        self.assertEqual(len(result.comments), 1)
        self.assertFalse(any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in result.paragraphs[0].runs))
        self.assertEqual("عبارت دوم", "".join(r.text for r in result.paragraphs[1].runs
                                                if r.font.highlight_color == WD_COLOR_INDEX.YELLOW))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["records"][0]["reason"], "comment_limit")
        self.assertEqual(summary["records"][1]["status"], "applied")

    def test_priority_is_respected_for_overlapping_witnesses(self):
        document = Document()
        document.add_paragraph("پژوهش حاضر نیازمند بررسی است.")
        issues = [finding("پژوهش حاضر", confidence="B"), finding("پژوهش")]
        self.assertTrue(annotate_document(document, issues, self.output))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["records"][0]["reason"], "overlapping_target")
        self.assertEqual(summary["applied"], 1)

    def test_paragraph_cap_and_disabled_cap_account_for_every_finding(self):
        document = Document()
        words = ["الف", "ب", "پ", "ت", "ث", "ج"]
        document.add_paragraph(" ".join(words))
        issues = [finding(word) for word in words]
        self.assertTrue(annotate_document(document, issues, self.output, per_par=2))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["reasons"], {"paragraph_limit": 4})
        self.assertEqual(summary["applied"], 2)
        self.assertTrue(annotate_document(document, issues, self.output, per_par=None))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["applied"], 6)
        self.assertEqual(len(Document(self.output).comments), 6)

    def test_every_skip_reason_is_visible_without_mutating_findings(self):
        document = Document()
        document.add_paragraph("پژوهش انجام شد.")
        issues = [finding("پژوهش"), finding("پژوهش"),
                  finding("انجام", confidence="C"),
                  finding("انجام", protected_overlap=True),
                  finding("انجام", source_part="word/footnotes.xml", story="footnote"),
                  finding("ناموجود"), finding("پژوهش", paragraph=8)]
        before = [issue.to_dict() for issue in issues]
        self.assertTrue(annotate_document(document, issues, self.output))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["reasons"], {"duplicate": 1, "ambiguous_excluded": 1,
                         "protected_content": 1, "unsupported_story": 1,
                         "uncertain_location": 1, "location_unavailable": 1})
        self.assertEqual([issue.to_dict() for issue in issues], before)

    def test_non_body_legacy_location_never_targets_body(self):
        document = Document()
        document.add_paragraph("متن مشابه در بدنه")
        self.assertFalse(annotate_document(document, [finding("متن", story="header")], self.output))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["reasons"], {"unsupported_story": 1})
        self.assertEqual(len(Document(self.output).comments), 0)

    def test_c_requires_explicit_opt_in_and_remains_qualified(self):
        document = Document()
        document.add_paragraph("پژوهش حاضر")
        issues = [finding("پژوهش", confidence="C")]
        self.assertFalse(annotate_document(document, issues, self.output))
        self.assertEqual(self.assert_accounted(document)["applied"], 0)
        self.assertTrue(annotate_document(document, issues, self.output, include_ambiguous=True))
        comment = list(Document(self.output).comments)[0].text
        self.assertIn("C — مبهم؛ فقط بررسی دستی", comment)
        self.assertEqual(issues[0].confidence, "C")

    def test_comments_are_three_short_lines_and_preserve_qualifiers(self):
        document = Document()
        document.add_paragraph("در این پژوهش می شود بازبینی کرد.")
        issue = finding("می شود", confidence="B")
        issue.message = "ممکن است این صورت در بافت خاص پذیرفتنی باشد.\n بازبینی دستی لازم است."
        issue.suggestion = "پیشنهاد را بررسی کنید؛\n تصمیم نهایی با نویسنده است."
        self.assertTrue(annotate_document(document, [issue], self.output))
        comment = list(Document(self.output).comments)[0].text
        self.assertEqual(len(comment.splitlines()), 3)
        self.assertIn("ممکن است", comment)
        self.assertIn("تصمیم نهایی با نویسنده است.", comment)
        self.assertIn("B — نیازمند بررسی", comment)
        self.assertLess(len(comment), 400)
        issue.message = "توضیح طولانی " * 40 + "اما این مورد قطعی نیست."
        issue.suggestion = "توضیح تکمیلی " * 40 + "فقط پس از بررسی انسانی."
        annotate_document(document, [issue], self.output)
        comment = list(Document(self.output).comments)[0].text
        self.assertEqual(len(comment.splitlines()), 3)
        self.assertIn("دلیل کامل در گزارش", comment)
        self.assertIn("پیشنهاد کامل گزارش", comment)
        self.assertNotIn("توضیح طولانی", comment)
        self.assertLess(len(comment), 400)

    def source_document(self):
        source = self.root / "original.docx"
        document = Document()
        document.add_paragraph("دادهٔ علمی µ ² و شماره ۱۴۰۵؛ پژوهش حاضر")
        document.sections[0].header.paragraphs[0].text = "سرصفحهٔ دست‌نخورده"
        document.save(source)
        with zipfile.ZipFile(source, "a") as package:
            package.writestr("customXml/opaque.bin", b"\x00original opaque package bytes\xff")
        document = Document(source)
        document._review_source_path = str(source)
        document._review_source_hash = sha256_file(source)
        return source, document

    def test_zero_comment_copy_is_independent_valid_and_byte_identical(self):
        source, document = self.source_document()
        original_bytes = source.read_bytes()
        self.assertFalse(annotate_document(document, [], self.output))
        summary = self.assert_accounted(document)
        self.assertEqual(summary["requested"], 0)
        self.assertEqual(source.read_bytes(), original_bytes)
        self.assertEqual(self.output.read_bytes(), original_bytes)
        self.assertEqual(Document(self.output).paragraphs[0].text, document.paragraphs[0].text)
        self.output.write_bytes(b"independent copy")
        self.assertEqual(source.read_bytes(), original_bytes)

    def test_real_annotation_preserves_original_and_opaque_package_parts(self):
        source, document = self.source_document()
        original_bytes = source.read_bytes()
        self.assertTrue(annotate_document(document, [finding("پژوهش")], self.output))
        self.assert_accounted(document)
        self.assertEqual(source.read_bytes(), original_bytes)
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(self.output) as reviewed:
            mutable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
            for name in original.namelist():
                if name not in mutable:
                    self.assertEqual(original.read(name), reviewed.read(name), name)
        result = Document(self.output)
        self.assertEqual(result.paragraphs[0].text, document.paragraphs[0].text)
        self.assertEqual(result.sections[0].header.paragraphs[0].text, "سرصفحهٔ دست‌نخورده")
        self.assertEqual(len(result.comments), 1)

    def test_rich_run_and_paragraph_properties_survive_cross_run_highlight(self):
        def formatting_signature(document):
            signature = []
            for paragraph in document.paragraphs:
                properties = paragraph._p.pPr
                paragraph_properties = (etree.tostring(properties, method="c14n")
                                        if properties is not None else b"")
                character_properties = []
                for run in paragraph.runs:
                    properties = copy.deepcopy(run._r.rPr)
                    if properties is not None:
                        for child in list(properties):
                            if child.tag == qn("w:highlight"):
                                properties.remove(child)
                    # A newly introduced highlight-only rPr is otherwise empty.
                    run_properties = (etree.tostring(properties, method="c14n")
                                      if properties is not None and len(properties) else b"")
                    character_properties.extend((character, run_properties) for character in run.text)
                signature.append((paragraph_properties, character_properties))
            return signature

        document = Document()
        paragraph = document.add_paragraph(style="Heading 2")
        paragraph.paragraph_format.space_after = Pt(19)
        first = paragraph.add_run("در این پژوهش می ")
        first.bold = True
        first.italic = False
        first.font.name = "B Nazanin"
        first.font.size = Pt(15)
        first.font.color.rgb = RGBColor(20, 50, 90)
        first.font.underline = WD_UNDERLINE.DOUBLE
        language = OxmlElement("w:lang")
        language.set(qn("w:bidi"), "fa-IR")
        first._r.rPr.append(language)
        first._r.rPr.rFonts.set(qn("w:cs"), "IRLotus")
        second = paragraph.add_run("شود بازبینی کرد.")
        second.font.name = "Calibri"
        second.font.size = Pt(11)
        second.font.superscript = True
        second.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
        before = formatting_signature(document)
        self.assertTrue(annotate_document(document, [finding("می شود")], self.output))
        self.assertEqual(before, formatting_signature(document))
        self.assertEqual(before, formatting_signature(Document(self.output)))

    def test_zero_limit_and_invalid_limit_do_not_discard_accounting(self):
        document = Document()
        document.add_paragraph("پژوهش")
        self.assertFalse(annotate_document(document, [finding("پژوهش")], self.output, max_comments=0))
        self.assertEqual(self.assert_accounted(document)["reasons"], {"comment_limit": 1})
        for value in (True, -1, 1.5):
            with self.assertRaises(ValueError):
                annotate_document(document, [], self.output, max_comments=value)


if __name__ == "__main__":
    unittest.main()
