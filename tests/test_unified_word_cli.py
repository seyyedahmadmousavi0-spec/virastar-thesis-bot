"""The default CLI review Word covers every finding in that one output."""
import contextlib
import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from analysis_model import Issue
from doc_guard import sha256_file
import checker
import review_cli


class UnifiedWordCliTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="unified_word_cli_")
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.source = self.root / "source.docx"
        document = Document()
        for _ in range(28):
            document.add_paragraph("متن پژوهش برای بررسی است.")
        document.save(self.source)
        self.before = self.source.read_bytes()
        self.issues = [
            Issue("mi_prefix", message="فاصله‌گذاری را بررسی کنید.",
                  suggestion="با توجه به بافت بازبینی کنید.", confidence="A",
                  par_index=index, target_text="متن", char_start=0, char_end=3,
                  location_status="exact", loc=f"پاراگراف {index + 1}")
            for index in range(28)
        ]
        self.issues.extend([
            Issue("missing_section", message="بخش لازم در متن استخراج‌شده یافت نشد.",
                  suggestion="وجود بخش را دستی بررسی کنید.", confidence="B", loc="کل سند"),
            Issue("mi_prefix", message="ساختار نیازمند بازبینی است.", confidence="C",
                  par_index=1, target_text="پژوهش", char_start=4, char_end=9,
                  location_status="exact", loc="پاراگراف ۲"),
            Issue("mi_prefix", message="محتوای حفاظت‌شده نیازمند بررسی دستی است.",
                  protected_overlap=True, confidence="B", par_index=2,
                  target_text="متن", char_start=0, char_end=3,
                  location_status="exact", loc="پاراگراف ۳"),
        ])
        self.engine = SimpleNamespace(
            load_rules=checker.load_rules, annotate=checker.annotate,
            build_report=checker.build_report,
            analyze=self.analyze_fixture,
        )

    def analyze_fixture(self, *args):
        document = Document(self.source)
        # Mirror checker.analyze provenance so the annotator can preserve the
        # actual input ZIP, including timestamps, rather than re-serialize it.
        document._review_source_path = str(self.source.resolve())
        document._review_source_hash = sha256_file(self.source)
        return document, self.issues, {"mode": "thesis"}

    def run_cli(self, extra=()):
        reviewed = self.root / "reviewed.docx"
        output, errors = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(errors):
            result = review_cli.main(
                [str(self.source), "--annotate", str(reviewed), "--format", "json", *extra],
                engine=self.engine,
            )
        self.assertEqual(0, result, errors.getvalue())
        self.assertEqual(self.before, self.source.read_bytes())
        return reviewed, json.loads(output.getvalue())

    def assert_complete_in_same_word(self, reviewed, report, expected_source):
        result = Document(reviewed)
        summary = report["stats"]["unified_word_review"]
        self.assertEqual(len(self.issues), len(result.comments))
        self.assertEqual(len(self.issues), summary["requested"])
        self.assertEqual(len(self.issues), summary["comments"])
        self.assertEqual(len(self.issues), summary["highlighted"])
        self.assertEqual(expected_source, summary["source_comments"])
        self.assertEqual(len(self.issues) - expected_source, summary["context_comments"])
        self.assertEqual(0, summary.get("appendix_comments", 0))
        original = Document(self.source)
        self.assertEqual([p.text for p in original.paragraphs],
                         [p.text for p in result.paragraphs])
        self.assertTrue(any(run.font.highlight_color == WD_COLOR_INDEX.YELLOW
                            for p in result.paragraphs for run in p.runs))
        structural_comments = [comment.text for comment in result.comments
                               if "missing_section" in comment.text]
        self.assertEqual(1, len(structural_comments))
        self.assertIn("سطح سند", structural_comments[0])
        self.assertIn("نیست", structural_comments[0])
        manual_comments = [comment.text for comment in result.comments if "C" in comment.text]
        self.assertEqual(1, len(manual_comments))
        self.assertIn("بررسی دستی", manual_comments[0])
        self.assertNotIn("review_document", report["stats"])
        self.assertEqual({"source.docx", "reviewed.docx"}, {p.name for p in self.root.glob("*.docx")})

    def test_default_output_contains_source_and_all_manual_findings(self):
        reviewed, report = self.run_cli()
        self.assert_complete_in_same_word(reviewed, report, expected_source=28)

    def test_explicit_source_cap_cannot_drop_findings_from_same_word(self):
        reviewed, report = self.run_cli(["--max-comments", "25"])
        self.assert_complete_in_same_word(reviewed, report, expected_source=25)

    def test_zero_source_anchors_still_puts_all_findings_in_same_word(self):
        reviewed, report = self.run_cli(["--max-comments", "0"])
        self.assert_complete_in_same_word(reviewed, report, expected_source=0)

    def test_no_findings_keeps_copy_identical_and_adds_no_text(self):
        self.issues.clear()
        reviewed, report = self.run_cli()
        self.assertEqual(self.before, reviewed.read_bytes())
        summary = report["stats"]["unified_word_review"]
        self.assertTrue(summary["output_created"])
        self.assertEqual(0, summary["comments"])
        self.assertEqual(0, summary["context_comments"])
        self.assertEqual(0, summary.get("appendix_comments", 0))
        self.assertEqual(28, len(Document(reviewed).paragraphs))

    def test_failed_completion_returns_error_before_publishing_report_artifacts(self):
        reviewed = self.root / "reviewed.docx"
        artifacts = self.root / "reports"
        output, errors = io.StringIO(), io.StringIO()
        with patch("unified_word_review.finalize_word_review", side_effect=ValueError("Incomplete review")), \
                contextlib.redirect_stdout(output), contextlib.redirect_stderr(errors):
            result = review_cli.main([
                str(self.source), "--annotate", str(reviewed),
                "--output-dir", str(artifacts), "--format", "json",
            ], engine=self.engine)
        self.assertEqual(2, result)
        self.assertEqual("", output.getvalue())
        self.assertIn("Incomplete review", errors.getvalue())
        self.assertFalse(artifacts.exists())
        self.assertEqual(self.before, self.source.read_bytes())


if __name__ == "__main__":
    unittest.main()
