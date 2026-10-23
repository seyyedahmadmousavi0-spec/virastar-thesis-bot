"""All findings reach Word without using an invented source selection."""
import io
import os
os.environ["BALE_SKIP_DOTENV"] = "1"
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

import bot
from analysis_model import Issue
from report_store import ReportStore
from word_review_report import WordReviewReportError


class CompleteWordDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "source.docx"
        document = Document()
        self.words = [f"witness{index:03d}" for index in range(30)]
        document.add_paragraph(" ".join(self.words))
        document.save(self.source)
        self.before = self.source.read_bytes()
        self.sent = []
        self.messages = []
        self.findings = [Issue(f"RULE{index}", message=f"دلیل {index}", suggestion="بازبینی شود",
                               target_text=word, par_index=0) for index, word in enumerate(self.words)]
        self.stats = {"analysis_status": "complete", "ambiguous_candidates": 0}
        for name, value in (("reports", ReportStore()), ("MAX_COMMENTS", None)):
            target = patch.object(bot, name, value)
            target.start(); self.addCleanup(target.stop)

    def run_document(self, report_error=False):
        def analyze(*args):
            return Document(self.source), self.findings, self.stats

        def capture(chat, path, caption):
            self.sent.append((Path(path).name, Path(path).read_bytes(), caption))

        with patch.object(bot, "download", side_effect=lambda fid, dest: Path(dest).write_bytes(self.before)), \
                patch.object(bot.checker, "analyze", side_effect=analyze), \
                patch.object(bot, "send_file", side_effect=capture), \
                patch.object(bot, "send", side_effect=lambda chat, text, **kw: self.messages.append(text) or True), \
                patch.object(bot, "call", side_effect=AssertionError("Real transport forbidden")):
            if report_error:
                with patch.object(bot, "finalize_word_review", side_effect=WordReviewReportError("bounded")):
                    bot.process_document(7, {"file_name": "study.docx", "file_id": "synthetic"})
            else:
                bot.process_document(7, {"file_name": "study.docx", "file_id": "synthetic"})
        self.assertEqual(self.before, self.source.read_bytes())

    def yellow(self, document):
        return [run for paragraph in document.paragraphs for run in paragraph.runs
                if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]

    def test_thirty_exact_findings_in_one_paragraph_all_in_source_copy(self):
        self.run_document()
        self.assertEqual(1, len(self.sent))
        reviewed = Document(io.BytesIO(self.sent[0][1]))
        self.assertEqual(30, len(reviewed.comments))
        self.assertEqual(" ".join(self.words), reviewed.paragraphs[0].text)
        self.assertEqual(self.words, [run.text for run in self.yellow(reviewed)])
        self.assertEqual(30, self.stats["annotation_summary"]["applied"])
        self.assertEqual(30, self.stats["unified_word_review"]["comments"])
        self.assertNotIn("review_document", self.stats)

    def test_every_manual_structural_and_unsupported_finding_in_complete_word(self):
        self.findings += [
            Issue("MISSING", confidence="B", message="بخش مورد انتظار در متن استخراج‌شده دیده نشد"),
            Issue("FOOTNOTE", message="پاورقی بازبینی شود", story="footnote", source_part="word/footnotes.xml"),
            Issue("MANUAL", confidence="C", message="نیاز به بررسی دستی", par_index=0, target_text=self.words[0]),
            Issue("PROTECTED", message="مقدار علمی حفظ شود", protected_overlap=True, par_index=0, target_text=self.words[1]),
        ]
        self.run_document()
        self.assertEqual(1, len(self.sent))
        complete = Document(io.BytesIO(self.sent[0][1]))
        self.assertEqual(" ".join(self.words), complete.paragraphs[0].text)
        self.assertEqual(34, len(complete.comments))
        self.assertEqual(34, len(self.yellow(complete)))
        self.assertTrue(self.sent[0][0].startswith("نسخه-بازبینی-"))
        self.assertIn("همین فایل", self.sent[0][2])
        self.assertIn("یافته‌های تکمیلی بازبینی", "\n".join(p.text for p in complete.paragraphs))
        self.assertIn("34", self.messages[-1])
        self.assertEqual({"output_created": True, "requested": 34, "comments": 34, "highlighted": 34,
                          "source_comments": 30, "appendix_comments": 4,
                          "exact_comments": 30, "paragraph_comments": 0}, self.stats["unified_word_review"])
        self.assertNotIn("review_document", self.stats)
        stored = bot.reports.get(7, 7)
        self.assertEqual({"html", "txt"}, set(stored))
        for data in stored.values():
            text = data.decode("utf-8")
            for rule in ("MISSING", "FOOTNOTE", "MANUAL", "PROTECTED"):
                self.assertIn(rule, text)

    def test_full_review_failure_sends_reports_without_a_partial_word(self):
        self.findings.append(Issue("MISSING", message="بخشی یافت نشد", confidence="B"))
        self.run_document(report_error=True)
        self.assertEqual([".html", ".txt"], [Path(name).suffix for name, _, _ in self.sent])
        self.assertFalse(self.stats["unified_word_review"]["output_created"])
        self.assertIn("آماده نشد", self.messages[-1])
        self.assertIn("MISSING", self.sent[-1][1].decode("utf-8"))
        self.assertNotIn("نسخهٔ بازبینی آماده است", self.messages[-1])


if __name__ == "__main__":
    unittest.main()
