"""Synthetic integration and adversarial regressions for the professional upgrade."""
import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

import academic_checker
from analysis_model import Issue, OffsetMapper
import checker
from doc_guard import DocumentIntegrityError, sha256_file, validate_document
import en_checker
import pdf_handler
from protected_content import ProtectedContentDetector
import review_cli


class ProfessionalIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rules = checker.load_rules()

    def document(self, text="بخاطر نمونه ها بررسی شد."):
        path = self.root / "source.docx"
        doc = Document()
        doc.add_paragraph(text)
        doc.save(path)
        return path

    def test_presentation_ligature_mapping_rejects_half_a_character(self):
        mapper = OffsetMapper.create_normalized("xﻻy")
        self.assertEqual("xلاy", mapper.normalized_text)
        self.assertEqual((1, 2, True), mapper.map_slice_to_original(1, 3))
        self.assertFalse(mapper.map_slice_to_original(1, 2)[2])
        self.assertEqual("xﻻy", mapper.original_text)

    def test_combining_cluster_and_scientific_compatibility_are_preserved(self):
        source = "كِ CO₂ x² µm μ K"
        mapper = OffsetMapper.create_normalized(source)
        self.assertIn("CO₂ x² µm μ K", mapper.normalized_text)
        self.assertFalse(mapper.map_slice_to_original(0, 1)[2])
        self.assertEqual((0, 2, True), mapper.map_slice_to_original(0, 2))

    def test_issue_evidence_does_not_search_another_occurrence(self):
        issue = Issue("TEST", target_text="alpha", char_start=2, char_end=7, auto_applied=True)
        self.assertFalse(issue.validate_evidence("alpha beta"))
        self.assertEqual("C", issue.confidence)
        self.assertIsNone(issue.to_dict()["original_text"])
        self.assertFalse(issue.auto_applied)

    def test_scientific_compound_units_and_arabic_digits_are_whole_spans(self):
        for text in ("15 mg/dL", "5 m/s", "−١٢٫٥ µm", "25 °C", "4.2e-3 mol/L"):
            with self.subTest(text=text):
                self.assertTrue(any(span.start == 0 and span.end == len(text)
                                    for span in ProtectedContentDetector.find_spans(text)))

    def test_latin_abbreviation_and_persian_citation_punctuation(self):
        for token in ("e.g.", "et al.", "(احمدی، ١٤٠٠)"):
            text = "prefix " + token + " suffix"
            left = text.index(token)
            self.assertTrue(ProtectedContentDetector.is_protected(left, left+len(token), ProtectedContentDetector.find_spans(text)))

    def test_code_and_email_are_not_rewritten_by_prose_rules(self):
        text = "نشانی name@example.org و کد `بخاطر نمونه ها` ثبت شد."
        issues = checker.check_text(0, text, "synthetic", self.rules)
        self.assertNotIn("common_mistake", {issue.rule_id for issue in issues})
        self.assertNotIn("plural_space", {issue.rule_id for issue in issues})

    def test_bibliography_titles_and_author_names_are_not_prose_errors(self):
        text = "Smith, A. (2020). A honest mistake; تحقيق بخاطر منبع."
        self.assertEqual([], checker.check_text(0, text, "reference", self.rules, in_refs=True))

    def test_english_noun_modifiers_and_valid_repetition(self):
        for text in ("These study designs were selected.", "Many research studies are available.",
                     "The sample had had sufficient exposure.", "We found that that result was correct."):
            with self.subTest(text=text):
                self.assertFalse(any(i.rule_id in {"EN_SVA", "EN_PUNCT"} for i in en_checker.check_english_text(text)))

    def test_english_explicit_subject_verb_error(self):
        issues = en_checker.check_english_text("These results is inconsistent.")
        self.assertTrue(any(i.rule_id == "EN_SVA" and i.target_text == "is" and i.suggestion == "are" for i in issues))

    def test_scientific_interval_does_not_trigger_bracket_error(self):
        self.assertFalse(any(i.rule_id == "EN_BRACKET" for i in en_checker.check_english_text("The interval [0, 1) was used.")))
        self.assertTrue(any(i.rule_id == "EN_BRACKET" for i in en_checker.check_english_text("The result )was unexpected(.")))

    def test_first_person_phrase_is_not_counted_twice(self):
        issues = academic_checker.check_academic_style("به نظر من این نتیجه درست است.")
        self.assertEqual(1, len([i for i in issues if i.rule_id == "ACAD_FIRST_PERSON"]))

    def test_full_analysis_does_not_create_header_or_footer_parts(self):
        path = self.document()
        original = Document(path)
        xml = original.element.xml
        parts = {str(part.partname) for part in original.part.package.parts}
        analyzed, _, _ = checker.analyze(path)
        self.assertEqual(xml, analyzed.element.xml)
        self.assertEqual(parts, {str(part.partname) for part in analyzed.part.package.parts})

    def test_nested_table_and_header_issues_keep_source_identity(self):
        path = self.root / "nested.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "بخاطر آزمون"
        doc.sections[0].header.paragraphs[0].text = "بخاطر سربرگ"
        doc.save(path)
        _, issues, stats = checker.analyze(path)
        found = [i for i in issues if i.rule_id == "common_mistake"]
        self.assertEqual({"table", "header"}, {i.story for i in found})
        self.assertEqual(2, len({i.block_id for i in found}))
        self.assertTrue(all(i.xml_path and i.location_status == "exact" for i in found))
        self.assertGreaterEqual(stats["protected_content_audit"]["count"], 0)

    def test_table_annotation_uses_xml_source_and_preserves_text(self):
        path = self.root / "table.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "بخاطر آزمون"
        doc.save(path)
        before = sha256_file(path)
        analyzed, issues, _ = checker.analyze(path)
        out = self.root / "reviewed.docx"
        selected = [i for i in issues if i.rule_id == "common_mistake"]
        self.assertTrue(checker.annotate(analyzed, selected, out))
        reopened = Document(out)
        self.assertEqual("بخاطر آزمون", reopened.tables[0].cell(0, 0).text)
        self.assertEqual(1, len(reopened.comments))
        self.assertEqual(before, sha256_file(path))

    def test_changed_in_memory_document_is_rejected(self):
        path = self.document()
        doc, issues, _ = checker.analyze(path)
        doc.paragraphs[0].add_run("changed")
        with self.assertRaises(DocumentIntegrityError):
            checker.annotate(doc, issues, self.root / "out.docx")

    def test_custom_named_existing_comments_part_is_updated(self):
        path = self.root / "custom.docx"
        doc = Document()
        paragraph = doc.add_paragraph("original")
        doc.add_comment(paragraph.runs[0], text="existing")
        doc.add_paragraph("alpha beta")
        doc.save(path)
        buffer = io.BytesIO()
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as output:
            for item in source.infolist():
                name, data = item.filename, source.read(item.filename)
                if name == "word/comments.xml":
                    name = "word/comments99.xml"
                if item.filename in {"word/_rels/document.xml.rels", "[Content_Types].xml"}:
                    data = data.replace(b"comments.xml", b"comments99.xml")
                output.writestr(name, data)
        path.write_bytes(buffer.getvalue())
        analyzed, _, _ = checker.analyze(path)
        issue = Issue("TEST", par_index=1, target_text="beta", char_start=6, char_end=10)
        out = self.root / "out.docx"
        self.assertTrue(checker.annotate(analyzed, [issue], out))
        self.assertEqual(2, len(Document(out).comments))
        with zipfile.ZipFile(out) as archive:
            self.assertIn("word/comments99.xml", archive.namelist())

    def test_dtd_and_utf16_dtd_are_rejected_before_docx_parser(self):
        for encoding in ("utf-8", "utf-16"):
            path = self.document()
            payload = io.BytesIO()
            with zipfile.ZipFile(path) as original, zipfile.ZipFile(payload, "w") as changed:
                for item in original.infolist():
                    data = original.read(item.filename)
                    if item.filename == "word/document.xml":
                        xml = data.decode("utf-8")
                        xml = xml[xml.index("?>")+2:]
                        data = ('<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE document [<!ENTITY x "test">]>' + xml).encode(encoding)
                    changed.writestr(item, data)
            path.write_bytes(payload.getvalue())
            with self.assertRaises(ValueError):
                validate_document(path, "docx")

    def test_finding_budget_is_explicit_and_bounded(self):
        rules = dict(self.rules, max_findings_per_block=10)
        issues = checker.check_text(0, "ي " * 1000, "synthetic", rules)
        self.assertLessEqual(len(issues), 11)
        self.assertTrue(any(i.rule_id == "findings_limit" for i in issues))

    def test_mixed_pdf_keeps_short_text_and_marks_empty_page(self):
        import pymupdf
        path = self.root / "mixed.pdf"
        with pymupdf.open() as doc:
            doc.new_page().insert_text((40, 40), "a honest mistake")
            doc.new_page()
            doc.save(path)
        _, issues, stats = checker.analyze_pdf(path, mode="english")
        self.assertTrue(any(i.rule_id == "EN_ARTICLE" and i.page_index == 0 for i in issues))
        self.assertTrue(any(i.rule_id == "pdf_page_unreadable" and i.page_index == 1 for i in issues))
        self.assertEqual(2, stats["total_pages"])

    def test_pdf_page_budget_failure_is_explicit(self):
        import pymupdf
        path = self.root / "two.pdf"
        with pymupdf.open() as doc:
            doc.new_page(); doc.new_page(); doc.save(path)
        self.assertEqual("budget", pdf_handler.extract_pdf_content(path, max_pages=1)["error_code"])

    def test_hyphenated_hour_article_and_uppercase_suggestion(self):
        self.assertFalse(any(i.rule_id == "EN_ARTICLE" for i in en_checker.check_english_text("An hour-long session was held.")))
        self.assertTrue(any(i.rule_id == "EN_ARTICLE" and i.suggestion == "An" for i in en_checker.check_english_text("A honest mistake occurred.")))

    def test_document_finding_limit_and_disabled_rules_apply_to_raw_text(self):
        rules = dict(self.rules, max_findings=2, disabled_rules=["page_numbers"])
        _, issues, _ = checker.analyze_text("ي " * 20, rules=rules)
        self.assertLessEqual(len(issues), 3)
        self.assertTrue(any(i.rule_id == "findings_limit" for i in issues))
        doc, issues, _ = checker.analyze(self.document(), rules=rules)
        self.assertNotIn("page_numbers", {i.rule_id for i in issues})

    def test_cli_corrupt_package_returns_error_without_traceback(self):
        path = self.root / "corrupt.docx"
        path.write_bytes(b"PKnot-a-zip")
        stream = io.StringIO()
        with contextlib.redirect_stderr(stream):
            self.assertEqual(2, review_cli.main([str(path)], engine=checker))
        self.assertNotIn("Traceback", stream.getvalue())

    def test_cli_json_and_complete_artifacts(self):
        path = self.document()
        out = self.root / "reports"
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            code = review_cli.main([str(path), "--format", "json", "--output-dir", str(out)], engine=checker)
        self.assertEqual(0, code)
        data = json.loads(stream.getvalue())
        self.assertEqual("2.0", data["schema_version"])
        self.assertEqual(4, len(list(out.iterdir())))

    def test_cli_rejects_annotation_to_input(self):
        path = self.document()
        before = sha256_file(path)
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(2, review_cli.main([str(path), "--annotate", str(path)], engine=checker))
        self.assertEqual(before, sha256_file(path))

    def test_cli_lists_rules_without_a_document(self):
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            self.assertEqual(0, review_cli.main(["--list-rules"], engine=checker))
        self.assertTrue(any(row["rule_id"] == "EN_ARTICLE" for row in json.loads(stream.getvalue())))
