"""Student-facing Kashan selection and complete delivery with mocked transport."""
import contextlib
import io
import json
import os
os.environ["BALE_SKIP_DOTENV"] = "1"
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from docx import Document
import pymupdf

import bot
import checker
import review_cli
from bot_runtime import JobGate
from report_store import ReportStore
from test_bot_runtime import RecordingPool


class KashanDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.pool = RecordingPool()
        self.gate = JobGate(2, 0)
        for name, value in (("modes", {}), ("document_scopes", {}), ("reports", ReportStore()),
                            ("jobs", self.gate), ("busy", self.gate.busy), ("lock", self.gate.lock),
                            ("_stop_event", threading.Event())):
            p = patch.object(bot, name, value); p.start(); self.addCleanup(p.stop)
        self.messages = []
        p = patch.object(bot, "send", side_effect=lambda chat, text, **kw: self.messages.append((text, kw)) or True)
        p.start(); self.addCleanup(p.stop)
        p = patch.object(bot, "call", side_effect=AssertionError("Real Bale transport is forbidden"))
        p.start(); self.addCleanup(p.stop)

    def message(self, text=None, document=None, caption=None):
        msg = {"chat": {"id": 7, "type": "private"}, "from": {"id": 7}}
        if text is not None: msg["text"] = text
        if document is not None: msg["document"] = document
        if caption is not None: msg["caption"] = caption
        return {"message": msg}

    def test_university_button_selects_degree_and_scope_without_guessing(self):
        self.assertIn("🏛 دانشگاه کاشان", str(bot.MAIN_KEYBOARD))
        bot.handle(self.message("🏛 دانشگاه کاشان"), self.pool)
        self.assertEqual("kashan", bot.modes[7])
        self.assertEqual("fragment", bot.document_scopes[7])
        self.assertEqual(bot.KASHAN_KEYBOARD, self.messages[-1][1]["reply_markup"])
        bot.handle(self.message("🎓 کاشان — دکتری"), self.pool)
        self.assertEqual("kashan_phd", bot.modes[7])
        self.assertIn("بخشی", self.messages[-1][0])
        bot.handle(self.message("/scope full"), self.pool)
        self.assertEqual("full", bot.document_scopes[7])
        bot.handle(self.message("/settings"), self.pool)
        self.assertIn("دکتری", self.messages[-1][0])
        self.assertIn("سند کامل", self.messages[-1][0])
        self.assertFalse(self.pool.calls)

    def test_mode_and_one_file_fragment_scope_are_captured_in_queue(self):
        bot.handle(self.message("/kashan master"), self.pool)
        bot.handle(self.message("/scope full"), self.pool)
        document = {"file_name": "chapter.docx", "file_id": "synthetic"}
        with patch.object(bot, "process_document") as process:
            bot.handle(self.message(document=document, caption="/part"), self.pool)
            bot.handle(self.message("/kashan_phd"), self.pool)
            self.pool.run()
            self.assertEqual("kashan_master", process.call_args.kwargs["mode"])
            self.assertEqual("fragment", process.call_args.kwargs["document_scope"])
        self.assertEqual("full", bot.document_scopes[7])
        self.assertFalse(self.gate.busy)

    def test_invalid_selection_and_scope_do_not_mutate_general_mode(self):
        bot.handle(self.message("/scope full"), self.pool)
        self.assertNotIn(7, bot.modes)
        bot.handle(self.message("/kashan article"), self.pool)
        self.assertNotIn(7, bot.modes)
        bot.handle(self.message("/kashan_master"), self.pool)
        bot.handle(self.message("/scope unknown"), self.pool)
        self.assertEqual("fragment", bot.document_scopes[7])
        bot.handle(self.message("/article"), self.pool)
        self.assertEqual("article", bot.modes[7])
        bot.handle(self.message("⬅️ منوی اصلی"), self.pool)
        self.assertEqual("article", bot.modes[7])

    def test_complete_word_job_delivers_copy_and_profile_reports(self):
        source = self.root / "original.docx"
        doc = Document(); doc.add_paragraph("بخاطر نمونه ها بررسی شد؛ مقدار 15 mg/dL و CO₂ حفظ شود.")
        doc.save(source); before = source.read_bytes()
        delivered = []
        def capture(chat, path, *args):
            delivered.append((Path(path).suffix, Path(path).read_bytes()))
            return {"message_id": 1}
        with patch.object(bot, "download", side_effect=lambda fid, path: Path(path).write_bytes(before)), \
                patch.object(bot, "send_file", side_effect=capture):
            bot.process_document(7, {"file_name": "پژوهش.docx", "file_id": "synthetic"},
                                 mode="kashan_master", sender_id=7, document_scope="fragment")
            self.assertEqual([".docx"], [item[0] for item in delivered])
            bot.deliver_reports(7, sender_id=7)
        self.assertEqual(before, source.read_bytes())
        reviewed = Document(io.BytesIO(delivered[0][1]))
        self.assertEqual(doc.paragraphs[0].text, reviewed.paragraphs[0].text)
        self.assertGreater(len(reviewed.comments), 0)
        self.assertEqual([".docx", ".html", ".txt"], [item[0] for item in delivered])
        for extension, content in delivered[1:]:
            text = content.decode("utf-8")
            self.assertIn("دانشگاه کاشان", text)
            self.assertIn("کارشناسی ارشد", text)
            self.assertIn("اعتبار", text)
            self.assertNotIn("KASHAN_SECTION_REVIEW", text)
        self.assertTrue(any("دانشگاه کاشان" in message for message, _ in self.messages))

    def test_pdf_and_raw_text_keep_profile_without_inventing_layout_findings(self):
        source = self.root / "sample.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((30, 50), "Abstract\nThis is a synthetic research summary.")
            pdf.save(source)
        _, issues, stats = checker.analyze_pdf(source, mode="kashan_master")
        self.assertEqual("kashan", stats["university_profile"]["id"])
        self.assertFalse(any(issue.rule_id in {"page_margin", "font_fa", "KASHAN_SECTION_REVIEW"} for issue in issues))
        _, issues, stats = checker.analyze_text("این متن یک چکیده است و فایل کامل نیست.", mode="kashan_phd")
        self.assertEqual("fragment", stats["university_profile"]["document_scope"])
        self.assertEqual("phd", stats["university_profile"]["degree"])
        self.assertFalse(any(issue.rule_id == "KASHAN_SECTION_REVIEW" for issue in issues))

    def test_cli_explicit_full_scope_exports_selected_degree_and_evidence(self):
        source = self.root / "source.docx"
        doc = Document(); doc.add_paragraph("چکیده"); doc.add_paragraph("متن کوتاه برای نمونه."); doc.save(source)
        output = self.root / "reports"
        with contextlib.redirect_stdout(io.StringIO()):
            result = review_cli.main([str(source), "--mode", "kashan_master", "--scope", "full",
                                      "--format", "json", "--output-dir", str(output)])
        self.assertEqual(0, result)
        data = json.loads((output / "report.json").read_text(encoding="utf-8"))
        self.assertEqual("master", data["stats"]["university_profile"]["degree"])
        self.assertEqual("full", data["stats"]["university_profile"]["document_scope"])
        self.assertTrue(any(item["rule_id"] == "KASHAN_SECTION_REVIEW" for item in data["findings"]))
        self.assertTrue(all(item["confidence"] != "A" for item in data["findings"] if item["rule_id"].startswith("KASHAN_")))


if __name__ == "__main__":
    unittest.main()
