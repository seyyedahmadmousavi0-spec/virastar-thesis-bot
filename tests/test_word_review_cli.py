"""CLI completeness, independent output paths and annotation-limit semantics."""
import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
import pymupdf

import review_cli


class WordReviewCliTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="word_review_cli_")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "source.docx"
        document = Document()
        for _ in range(35):
            document.add_paragraph("بخاطر پژوهش انجام شد.")
        document.save(self.source)
        self.before = self.source.read_bytes()

    def run_cli(self, arguments):
        output, errors = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(errors):
            result = review_cli.main([str(self.source), *map(str, arguments)])
        self.assertEqual(self.before, self.source.read_bytes())
        return result, output.getvalue(), errors.getvalue()

    def test_unlimited_default_and_complete_word_report_precede_json_export(self):
        reviewed = self.root / "reviewed.docx"
        complete = self.root / "all-findings.docx"
        artifacts = self.root / "reports"
        result, output, errors = self.run_cli([
            "--annotate", reviewed, "--word-report", complete,
            "--output-dir", artifacts, "--format", "json",
        ])
        self.assertEqual(0, result, errors)
        report = json.loads(output)
        annotation = report["stats"]["annotation_summary"]
        self.assertIsNone(annotation["max_comments"])
        self.assertGreaterEqual(annotation["applied"], 35)
        self.assertEqual(annotation["applied"], len(Document(reviewed).comments))
        source_document = Document(self.source)
        self.assertEqual([p.text for p in source_document.paragraphs],
                         [p.text for p in Document(reviewed).paragraphs])
        review = report["stats"]["review_document"]
        expected = len(report["findings"])
        self.assertEqual({"created": True, "findings": expected, "comments": expected,
                          "highlighted": expected, "filetype": "docx"}, review)
        complete_document = Document(complete)
        self.assertEqual(expected, len(complete_document.comments))
        self.assertEqual(expected, sum(run.font.highlight_color == WD_COLOR_INDEX.YELLOW
                                      for p in complete_document.paragraphs for run in p.runs))
        saved = json.loads((artifacts / "report.json").read_text(encoding="utf-8"))
        self.assertEqual(review, saved["stats"]["review_document"])
        self.assertNotIn(str(complete), output)

    def test_explicit_zero_keeps_source_copy_unmarked_but_review_complete(self):
        reviewed = self.root / "reviewed.docx"
        complete = self.root / "all-findings.docx"
        result, output, errors = self.run_cli([
            "--annotate", reviewed, "--word-report", complete,
            "--max-comments", "0", "--format", "json",
        ])
        self.assertEqual(0, result, errors)
        report = json.loads(output)
        self.assertEqual(0, report["stats"]["annotation_summary"]["applied"])
        self.assertEqual(self.before, reviewed.read_bytes())
        self.assertEqual(0, len(Document(reviewed).comments))
        self.assertEqual(len(report["findings"]), len(Document(complete).comments))

    def test_positive_limit_does_not_truncate_complete_word_report(self):
        reviewed = self.root / "reviewed.docx"
        complete = self.root / "all-findings.docx"
        result, output, errors = self.run_cli([
            "--annotate", reviewed, "--word-report", complete,
            "--max-comments", "25", "--format", "json",
        ])
        self.assertEqual(0, result, errors)
        report = json.loads(output)
        self.assertEqual(25, len(Document(reviewed).comments))
        self.assertGreater(len(Document(complete).comments), 25)
        self.assertEqual(len(report["findings"]), len(Document(complete).comments))

    def test_colliding_outputs_fail_before_any_output_is_replaced(self):
        existing = self.root / "existing.docx"
        existing.write_bytes(b"existing output must be retained")
        cases = [
            ["--word-report", self.source],
            ["--annotate", existing, "--word-report", existing],
            ["--word-report", self.root / "report.txt", "--output-dir", self.root],
        ]
        for arguments in cases:
            with self.subTest(arguments=arguments):
                result, output, errors = self.run_cli(arguments)
                self.assertEqual(2, result)
                self.assertEqual("", output)
                self.assertIn("independent", errors)
                self.assertEqual(b"existing output must be retained", existing.read_bytes())

    def test_all_findings_report_alone_accepts_text_input(self):
        source = self.root / "paragraph.txt"
        source.write_text("بخاطر پژوهش انجام شد.", encoding="utf-8")
        complete = self.root / "text-review.docx"
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            result = review_cli.main([str(source), "--text", "--word-report", str(complete),
                                      "--format", "json"])
        self.assertEqual(0, result)
        report = json.loads(output.getvalue())
        self.assertGreater(len(report["findings"]), 0)
        self.assertEqual(len(report["findings"]), len(Document(complete).comments))

    def test_existing_artifacts_allow_new_independent_word_output(self):
        artifacts = self.root / "reports"
        artifacts.mkdir()
        for suffix in ("html", "json", "csv", "txt"):
            (artifacts / f"report.{suffix}").write_text("old report", encoding="utf-8")
        reviewed = self.root / "new-reviewed.docx"
        complete = self.root / "new-complete.docx"
        result, output, errors = self.run_cli([
            "--annotate", reviewed, "--word-report", complete,
            "--output-dir", artifacts, "--format", "json",
        ])
        self.assertEqual(0, result, errors)
        self.assertTrue(reviewed.is_file())
        self.assertTrue(complete.is_file())
        report = json.loads((artifacts / "report.json").read_text(encoding="utf-8"))
        self.assertEqual(len(report["findings"]), len(Document(complete).comments))

    def test_pdf_default_and_large_explicit_limit_remain_within_pdf_budget(self):
        source = self.root / "source.pdf"
        with pymupdf.open() as document:
            document.new_page().insert_text((50, 50), "This is a honest observation in a research document.")
            document.save(source)
        before = source.read_bytes()
        from pdf_annotator import annotate_pdf
        for limit in (None, 10000):
            reviewed = self.root / f"reviewed-{limit}.pdf"
            arguments = [str(source), "--mode", "english", "--annotate-pdf", str(reviewed), "--format", "json"]
            if limit is not None:
                arguments += ["--max-comments", str(limit)]
            with self.subTest(limit=limit), patch("pdf_annotator.annotate_pdf", wraps=annotate_pdf) as annotate:
                output, errors = io.StringIO(), io.StringIO()
                with contextlib.redirect_stdout(output), contextlib.redirect_stderr(errors):
                    result = review_cli.main(arguments)
                self.assertEqual(0, result, errors.getvalue())
                report = json.loads(output.getvalue())
                expected_limit = min(5000, len(report["findings"]) if limit is None else limit)
                self.assertEqual(expected_limit, annotate.call_args.kwargs["max_comments"])
                self.assertTrue(reviewed.is_file())
                self.assertEqual(before, source.read_bytes())


if __name__ == "__main__":
    unittest.main()
