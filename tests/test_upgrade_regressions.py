"""Synthetic regressions for original bytes, exact locations and false alerts."""
import io
import os
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.enum.text import WD_COLOR_INDEX
import checker
import safe_annotator
import en_checker
import academic_checker
from analysis_model import Issue, OffsetMapper
from doc_guard import DocumentGuard, DocumentIntegrityError, sha256_file, validate_document
from protected_content import ProtectedContentDetector

class UpgradeRegressions(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rules = checker.load_rules()

    def source(self, text="alpha beta alpha"):
        p = self.root / "source.docx"
        d = Document()
        d.add_paragraph(text)
        d.save(p)
        return p

    def test_reject_same_input_output(self):
        p = self.source()
        d, _, _ = checker.analyze(p)
        before = sha256_file(p)
        with self.assertRaises(DocumentIntegrityError):
            checker.annotate(d, [], p)
        self.assertEqual(before, sha256_file(p))

    def test_reject_hardlink_output(self):
        p = self.source()
        alias = self.root / "alias.docx"
        os.link(p, alias)
        d, _, _ = checker.analyze(p)
        with self.assertRaises(DocumentIntegrityError):
            checker.annotate(d, [], alias)

    def test_annotation_does_not_mutate_caller(self):
        p = self.source()
        d, _, _ = checker.analyze(p)
        xml = d.element.xml
        issue = Issue("TEST", par_index=0, target_text="beta", char_start=6, char_end=10)
        checker.annotate(d, [issue], self.root / "review.docx")
        self.assertEqual(xml, d.element.xml)

    def test_repeated_target_respects_second_offset(self):
        p = self.source()
        d, _, _ = checker.analyze(p)
        issue = Issue("TEST", par_index=0, target_text="alpha", char_start=11, char_end=16)
        out = self.root / "out.docx"
        self.assertTrue(checker.annotate(d, [issue], out))
        runs = Document(out).paragraphs[0].runs
        self.assertEqual("alpha beta ", "".join(r.text for r in runs if r.font.highlight_color != WD_COLOR_INDEX.YELLOW))
        self.assertEqual("alpha", "".join(r.text for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW))

    def test_ambiguous_target_is_not_annotated(self):
        p = self.source()
        d, _, _ = checker.analyze(p)
        out = self.root / "out.docx"
        self.assertFalse(checker.annotate(d, [Issue("TEST", par_index=0, target_text="alpha")], out))
        self.assertEqual(len(Document(out).comments), 0)

    def test_invalid_offset_does_not_fallback(self):
        d = Document()
        d.add_paragraph("alpha beta")
        out = self.root / "out.docx"
        self.assertFalse(checker.annotate(d, [Issue("TEST", par_index=0, target_text="beta", char_start=0, char_end=4)], out))

    def test_target_really_spans_two_runs(self):
        d = Document()
        p = d.add_paragraph()
        p.add_run("alpha be").bold = True
        p.add_run("ta gamma").italic = True
        out = self.root / "out.docx"
        self.assertTrue(checker.annotate(d, [Issue("TEST", par_index=0, target_text="beta", char_start=6, char_end=10)], out))
        p = Document(out).paragraphs[0]
        highlighted = [r for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual("beta", "".join(r.text for r in highlighted))
        self.assertTrue(highlighted[0].bold)
        self.assertTrue(highlighted[-1].italic)

    def test_complex_field_is_preserved_and_not_annotated(self):
        d = Document()
        p = d.add_paragraph("alpha")
        r = p.add_run()
        r._r.append(OxmlElement("w:fldChar"))
        xml = p._p.xml
        out = self.root / "out.docx"
        self.assertFalse(checker.annotate(d, [Issue("TEST", par_index=0, target_text="alpha")], out))
        self.assertEqual(xml, Document(out).paragraphs[0]._p.xml)

    def test_unknown_package_part_preserved(self):
        p = self.source()
        with zipfile.ZipFile(p, "a") as z:
            z.writestr("customXml/synthetic.bin", b"unchanged")
        d, _, _ = checker.analyze(p)
        out = self.root / "out.docx"
        checker.annotate(d, [Issue("TEST", par_index=0, target_text="beta")], out)
        with zipfile.ZipFile(p) as original, zipfile.ZipFile(out) as reviewed:
            for name in original.namelist():
                if name not in {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "word/comments.xml"}:
                    self.assertEqual(original.read(name), reviewed.read(name), name)

    def test_source_change_after_analysis_rejected(self):
        p = self.source()
        d, _, _ = checker.analyze(p)
        p.write_bytes(p.read_bytes()+b"changed")
        with self.assertRaises(DocumentIntegrityError):
            checker.annotate(d, [], self.root/"out.docx")

    def test_guard_work_directory_cannot_collide(self):
        p = self.root/"work_original.txt"
        p.write_text("synthetic")
        with DocumentGuard(p, work_dir=self.root) as guard:
            self.assertNotEqual(p, Path(guard.isolated_working_copy))
        self.assertEqual("synthetic",p.read_text())

    def test_guard_cleans_up_on_error(self):
        p = self.root/"original.txt"
        p.write_text("synthetic")
        with self.assertRaises(RuntimeError):
            with DocumentGuard(p) as guard:
                temporary = guard.temp_dir
                raise RuntimeError("synthetic")
        self.assertFalse(os.path.exists(temporary))

    def test_offset_out_of_bounds_is_unreliable(self):
        m = OffsetMapper.create_normalized("abc")
        self.assertFalse(m.map_slice_to_original(-1,2)[2])
        self.assertFalse(m.map_slice_to_original(1,9)[2])

    def test_raw_text_keeps_leading_space_offsets(self):
        text = "  بخاطر آزمون"
        _, issues, _ = checker.analyze_text(text)
        issue = next(i for i in issues if i.rule_id=="common_mistake")
        self.assertEqual(issue.target_text,text[issue.char_start:issue.char_end])
        self.assertEqual(2,issue.char_start)

    def test_original_evidence_after_arabic_letters(self):
        text = "تحقيق بخاطر نمونه"
        for i in checker.check_text(0,text,"test",self.rules):
            if i.target_text:
                self.assertEqual(i.target_text,text[i.char_start:i.char_end])
                self.assertIn(i.target_text,i.snippet)

    def test_numbers_greek_and_merged_evidence(self):
        text = "E = 2.5 µm; CO₂; −١٢٫٥; α + β"
        spans=ProtectedContentDetector.find_spans(text)
        for span in spans:
            self.assertEqual(text[span.start:span.end],span.text)
        for token in ["2.5","CO₂","١٢٫٥","α","β"]:
            start=text.index(token)
            self.assertTrue(ProtectedContentDetector.is_protected(start,start+len(token),spans))

    def test_scientific_unicode_stays_distinct(self):
        text = "CO₂ x² µm α"
        self.assertEqual(text,OffsetMapper.create_normalized(text).normalized_text)

    def test_quoted_style_is_not_authors_voice(self):
        self.assertFalse(academic_checker.check_academic_style('نویسنده گفت: «این شاهکار است».'))

    def test_style_positive_is_review_not_certain(self):
        found=academic_checker.check_academic_style("این شاهکار است.")
        self.assertTrue(found)
        self.assertTrue(all(i.confidence=="B" for i in found))

    def test_cohesion_abstains_by_default(self):
        self.assertEqual([],academic_checker.evaluate_cohesion(["word "*40]*4))

    def test_article_silent_h_positive(self):
        self.assertIn("EN_ARTICLE",{i.rule_id for i in en_checker.check_english_text("This is a honest mistake.")})

    def test_articles_negative_and_ambiguous(self):
        for text in ["a university", "an MRI scan", "a DNA sequence", "an hour", "a European study", "an SQL query"]:
            with self.subTest(text=text):
                self.assertNotIn("EN_ARTICLE",{i.rule_id for i in en_checker.check_english_text(text)})

    def test_valid_s_ending_nouns_and_data(self):
        for text in ["Each analysis was complete.", "Every species was sampled.", "This data is useful."]:
            self.assertNotIn("EN_SVA",{i.rule_id for i in en_checker.check_english_text(text)})

    def test_valid_prepositions(self):
        for text in ["The advantage may consist in precision.", "It is different than expected."]:
            self.assertNotIn("EN_PREP",{i.rule_id for i in en_checker.check_english_text(text)})

    def test_configured_terms_and_disabled_rules(self):
        rules=dict(self.rules,protected_terms=["بخاطر"],disabled_rules=["arabic_letters"])
        self.assertEqual([],checker.check_text(0,"بخاطر تحقيق","test",rules))

    def test_persian_full_stop_positive(self):
        self.assertIn("no_space_after_punct",{i.rule_id for i in checker.check_text(0,"نتایج اعلام شد.هیچ تفاوتی دیده نشد.","test",self.rules)})

    def test_fake_docx_and_pdf_rejected(self):
        p=self.root/"fake.docx";p.write_bytes(b"not a document")
        with self.assertRaises(ValueError): checker.analyze(p)
        with self.assertRaises(ValueError): checker.analyze_pdf(p)

    def test_zip_expansion_budget(self):
        p=self.root/"bomb.docx"
        with zipfile.ZipFile(p,"w",zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml","x")
            z.writestr("word/document.xml","0"*1000000)
        with self.assertRaises(ValueError): validate_document(p,"docx")

    def test_zip_duplicate_members_rejected(self):
        p=self.root/"duplicate.docx"
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with zipfile.ZipFile(p,"w") as z:
                z.writestr("[Content_Types].xml","x")
                z.writestr("word/document.xml","x")
                z.writestr("word/document.xml","y")
        with self.assertRaises(ValueError): validate_document(p,"docx")

    def test_pdf_preserves_raw_text_and_downgrades_locations(self):
        try:
            import pymupdf
        except ImportError:
            self.skipTest("PyMuPDF unavailable")
        p=self.root/"synthetic.pdf"
        d=pymupdf.open()
        page=d.new_page()
        page.insert_text((40,40),"This is a honest mistake in a synthetic experimental report. "*3)
        d.save(p);d.close()
        before=sha256_file(p)
        _,issues,stats=checker.analyze_pdf(p,mode="english")
        self.assertTrue(any(i.rule_id=="EN_ARTICLE" for i in issues))
        self.assertTrue(all(i.confidence!="A" for i in issues))
        self.assertEqual(before,sha256_file(p))

    def test_empty_pdf_yields_extraction_warning(self):
        try:
            import pymupdf
        except ImportError:
            self.skipTest("PyMuPDF unavailable")
        p=self.root/"empty.pdf"
        d=pymupdf.open();d.new_page();d.save(p);d.close()
        _,issues,_=checker.analyze_pdf(p)
        self.assertEqual(["pdf_scanned_warning"],[i.rule_id for i in issues])

    def test_serialized_issue_never_says_auto_applied(self):
        self.assertFalse(Issue("TEST",auto_applied=True).to_dict()["auto_applied"])

    def test_article_before_acronym_keeps_scientific_name(self):
        issues=checker.check_text(0,"This is a MRI scan.","test",self.rules)
        issue=next(i for i in issues if i.rule_id=="EN_ARTICLE")
        self.assertEqual("a",issue.target_text)
        self.assertEqual("an",issue.suggestion)

    def test_multiple_nonoverlapping_comments(self):
        d=Document();d.add_paragraph("alpha beta gamma")
        out=self.root/"out.docx"
        issues=[Issue("ONE",par_index=0,target_text="alpha",char_start=0,char_end=5),
                Issue("TWO",par_index=0,target_text="gamma",char_start=11,char_end=16)]
        self.assertTrue(checker.annotate(d,issues,out))
        self.assertEqual(2,len(Document(out).comments))

    def test_health_payload_contains_no_message_or_identity(self):
        import bot
        import json
        with patch.dict(bot._bot_status,{"last_message":{"text":"SYNTHETIC_PRIVATE_TEXT","chat_id":123}}):
            payload=json.dumps(bot._health_payload())
        self.assertNotIn("SYNTHETIC_PRIVATE_TEXT",payload)
        self.assertNotIn("chat_id",payload)

    def test_bot_annotations_require_explicit_opt_in(self):
        import bot
        def download(file_id,destination):
            d=Document();d.add_paragraph("alpha beta");d.save(destination)
        for requested in (False,True):
            with self.subTest(requested=requested), patch.object(bot,"download",side_effect=download), patch.object(bot,"send"), patch.object(bot,"send_file"), patch.object(checker,"annotate",return_value=False) as annotate:
                bot.process_document(98765,{"file_name":"synthetic.docx","file_id":"synthetic"},annotate_requested=requested)
                self.assertEqual(int(requested),annotate.call_count)

    def test_per_rule_benchmark_rejects_unrelated_alert(self):
        import contextlib,json,run_benchmark
        out=self.root/"metrics.json"
        with patch.object(checker,"check_text",return_value=[Issue("UNRELATED")]),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(1,run_benchmark.run_benchmark(out))
        result=json.loads(out.read_text())
        self.assertTrue(all(row["tp"]==0 and row["fn"]==1 for row in result["per_rule"].values()))

    def test_tums_ratio_threshold_is_configurable(self):
        import tums_checker
        with patch.object(tums_checker,"load_tums_profile",return_value={"tums_special_rules":{"max_intro_review_pct":60}}):
            self.assertEqual([],tums_checker.check_tums_chapter_proportions({1:400},1000))
        with patch.object(tums_checker,"load_tums_profile",return_value={"tums_special_rules":{"max_intro_review_pct":20}}):
            self.assertTrue(tums_checker.check_tums_chapter_proportions({1:400},1000))

    def test_huge_block_is_reviewed_without_rule_execution(self):
        issues=checker.check_text(0,"a"*200001,"test",self.rules)
        self.assertEqual(["analysis_limit"],[i.rule_id for i in issues])
        self.assertEqual("B",issues[0].confidence)

    def test_merged_cells_are_analyzed_once(self):
        p=self.root/"table.docx"
        d=Document();table=d.add_table(rows=1,cols=2)
        table.cell(0,0).merge(table.cell(0,1)).text="بخاطر آزمون"
        d.save(p)
        _,issues,stats=checker.analyze(p)
        self.assertEqual(1,stats["table_cells_analyzed"])
        self.assertEqual(1,len([i for i in issues if i.rule_id=="common_mistake"]))

    def test_integrity_checked_even_when_parser_raises(self):
        from doc_guard import immutable_input
        p=self.root/"synthetic.txt";p.write_text("before")
        @immutable_input
        def bad_parser(path):
            path.write_text("after")
            raise ValueError("synthetic error")
        with self.assertRaises(DocumentIntegrityError):
            bad_parser(p)

    def citation_document(self,body,refs):
        d=Document();d.add_paragraph(body);d.add_heading("References",level=1)
        for ref in refs: d.add_paragraph(ref)
        return d

    def test_citation_surname_substring_is_not_a_match(self):
        d=self.citation_document("(Smith, 2020)",["Goldsmith, A. (2020). Synthetic study."])
        issues,_=checker.check_citations(d,self.rules)
        self.assertEqual(["cite_missing_ref"],[i.rule_id for i in issues])

    def test_citation_suffix_and_multiple_citations(self):
        d=self.citation_document("(Smith, 2020a; Jones, 2021)",["Smith, A. (2020a). Synthetic study.","Jones, B. (2021). Another study."])
        issues,_=checker.check_citations(d,self.rules)
        self.assertEqual([],issues)

    def test_citation_narrative_and_ambiguous_candidates(self):
        d=self.citation_document("Smith (2020) reported results.",["Smith, A. (2020). One study.","Smith, B. (2020). Other study."])
        issues,_=checker.check_citations(d,self.rules)
        self.assertEqual(["cite_ambiguous_ref"],[i.rule_id for i in issues])
        self.assertEqual("C",issues[0].confidence)

    def test_numeric_citation_uses_labels_not_reference_count(self):
        d=self.citation_document("See [3,5].",["[3] Smith. Synthetic study.","[5] Jones. Another study."])
        issues,_=checker.check_citations(d,self.rules)
        self.assertEqual([],issues)

    def test_numeric_citation_range_detects_missing_middle(self):
        d=self.citation_document("See [3-5].",["[3] Smith. Synthetic study.","[5] Jones. Another study."])
        issues,_=checker.check_citations(d,self.rules)
        self.assertEqual(["numeric_cite"],[i.rule_id for i in issues])
        self.assertIn("[4]",issues[0].message)

    def test_deduplication_keeps_distinct_reasons(self):
        issues=[Issue("TEST",message="one"),Issue("TEST",message="two")]
        self.assertEqual(2,len(checker.deduplicate_issues(issues)))
