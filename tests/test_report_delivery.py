"""Report ownership, retention and complete mocked messenger delivery paths."""
import io
import contextlib
import os
os.environ["BALE_SKIP_DOTENV"] = "1"
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from docx import Document
import pymupdf

import bot
import review_cli
from bot_runtime import JobGate
from report_store import ReportStore
from test_bot_runtime import RecordingPool


class StoreTests(unittest.TestCase):
    def test_owner_chat_and_expiry_are_all_enforced(self):
        now = [1.0]
        store = ReportStore(ttl_seconds=10, clock=lambda: now[0])
        payload = {"html": b"html", "txt": b"text"}
        self.assertTrue(store.put(-100, 1, payload))
        self.assertIsNone(store.get(-100, 2))
        self.assertIsNone(store.get(100, 1))
        self.assertEqual(payload, store.get(-100, 1))
        now[0] = 11
        self.assertIsNone(store.get(-100, 1))
        self.assertEqual(0, store._bytes)

    def test_budgets_evict_oldest_and_reject_oversize_without_stale_results(self):
        store = ReportStore(max_bytes=10, max_entry_bytes=8, max_entries=2)
        payload = {"html": b"aaa", "txt": b"bbb"}
        store.put(1, 1, payload)
        store.put(2, 2, payload)
        self.assertIsNone(store.get(1, 1))
        self.assertEqual(6, store._bytes)
        self.assertFalse(store.put(2, 2, {"html": b"a" * 9, "txt": b"b"}))
        self.assertIsNone(store.get(2, 2))
        self.assertEqual(0, store._bytes)

    def test_delete_and_shutdown_release_data_and_do_not_retain_documents(self):
        store = ReportStore()
        with self.assertRaises(ValueError):
            store.put(1, 1, {"docx": b"PRIVATE_DOCUMENT"})
        store.put(1, 1, {"html": b"a", "txt": b"b"})
        self.assertTrue(store.delete(1, 1))
        self.assertFalse(store.delete(1, 1))
        store.put(1, 1, {"html": b"a", "txt": b"b"})
        store.close()
        self.assertIsNone(store.get(1, 1))
        self.assertFalse(store.put(1, 1, {"html": b"a", "txt": b"b"}))


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.store = ReportStore()
        for name, value in (("reports", self.store), ("_stop_event", threading.Event())):
            p = patch.object(bot, name, value)
            p.start()
            self.addCleanup(p.stop)
        self.sent = []
        self.messages = []
        p = patch.object(bot, "send", side_effect=lambda chat, text, **kw: self.messages.append(text) or True)
        p.start(); self.addCleanup(p.stop)
        p = patch.object(bot, "call", side_effect=AssertionError("Real messenger transport forbidden"))
        p.start(); self.addCleanup(p.stop)
        p = patch.object(bot, "send_file", side_effect=self.capture)
        p.start(); self.addCleanup(p.stop)

    def capture(self, chat, path, *args):
        self.sent.append((chat, Path(path).suffix, Path(path).read_bytes()))
        return {"message_id": 1}

    def test_default_word_short_summary_and_optional_complete_report(self):
        source = self.root / "source.docx"
        doc = Document()
        doc.add_paragraph("بخاطر نمونه ها بررسی شد.")
        doc.save(source)
        before = source.read_bytes()
        with patch.object(bot, "download", side_effect=lambda fid, dest: Path(dest).write_bytes(before)):
            bot.process_document(-100, {"file_name": "پایان‌نامه.docx", "file_id": "synthetic"}, sender_id=7)
        self.assertEqual(before, source.read_bytes())
        self.assertEqual([".docx"], [row[1] for row in self.sent])
        result = Document(io.BytesIO(self.sent[0][2]))
        self.assertEqual(doc.paragraphs[0].text, result.paragraphs[0].text)
        self.assertGreater(len(result.comments), 0)
        full_review = result
        from docx.enum.text import WD_COLOR_INDEX
        yellow = [run for paragraph in full_review.paragraphs for run in paragraph.runs
                  if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertGreater(len(yellow), 0)
        from checker import analyze
        _, findings, _ = analyze(source)
        self.assertEqual(len(findings), len(full_review.comments))
        self.assertIn("یافته‌های تکمیلی بازبینی", "\n".join(p.text for p in result.paragraphs))
        self.assertIn("کامنت", self.messages[-1])
        self.assertLess(len(self.messages[-1]), 800)
        self.assertIn("/report", self.messages[-1])
        bot.deliver_reports(-100, sender_id=8)
        self.assertEqual(1, len(self.sent))
        bot.deliver_reports(-100, sender_id=7)
        self.assertEqual([".docx", ".html", ".txt"], [row[1] for row in self.sent])
        self.assertIn("common_mistake", self.sent[-1][2].decode("utf-8"))

    def test_cache_rejection_sends_complete_reports_immediately(self):
        source = self.root / "source.docx"
        doc = Document(); doc.add_paragraph("بخاطر پژوهش"); doc.save(source)
        with patch.object(bot, "reports", ReportStore(max_bytes=1)), patch.object(
                bot, "download", side_effect=lambda fid, dest: Path(dest).write_bytes(source.read_bytes())):
            bot.process_document(1, {"file_name": "a.docx", "file_id": "synthetic"})
        self.assertEqual({".docx", ".html", ".txt"}, {row[1] for row in self.sent})
        self.assertIn("ظرفیت", self.messages[-1])

    def test_pdf_default_retains_page_numbered_text_fallback(self):
        source = self.root / "source.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((50, 50), "This is a honest observation in a synthetic research document.")
            pdf.save(source)
        with patch.object(bot, "PDF_ANNOTATION_EXPERIMENTAL", False), patch.object(
                bot, "download", side_effect=lambda fid, dest: Path(dest).write_bytes(source.read_bytes())):
            bot.process_document(1, {"file_name": "a.pdf", "file_id": "synthetic"})
        self.assertEqual([".txt"], [row[1] for row in self.sent])
        self.assertIn("صفحه", self.sent[0][2].decode("utf-8"))
        self.assertIsNotNone(self.store.get(1, 1))

    def test_commands_use_sender_scope_and_reject_unattributed_groups(self):
        gate = JobGate(2, 0)
        self.store.put(-100, 7, {"html": b"private html", "txt": b"private text"})
        pool = RecordingPool()
        with patch.object(bot, "jobs", gate):
            bot.handle({"message": {"chat": {"id": -100, "type": "group"}, "from": {"id": 8}, "text": "/report"}}, pool)
            pool.run()
            self.assertFalse(self.sent)
            bot.handle({"message": {"chat": {"id": -100, "type": "group"}, "text": "/report"}}, pool)
            self.assertEqual(1, len(pool.calls))
            bot.handle({"message": {"chat": {"id": -100, "type": "group"}, "from": {"id": 7}, "text": "/forget"}}, pool)
        self.assertIsNone(self.store.get(-100, 7))

    def test_report_does_not_inherit_ten_second_analysis_cooldown(self):
        now = [100]
        gate = JobGate(1, 10, clock=lambda: now[0])
        token, _, _ = gate.reserve(1); gate.release(token)
        now[0] += 3
        pool = RecordingPool()
        with patch.object(bot, "jobs", gate):
            bot.handle({"message": {"chat": {"id": 1}, "text": "/report"}}, pool)
        self.assertEqual(1, len(pool.calls))
        pool.run()


class ProgressTests(unittest.TestCase):
    def test_status_is_edited_and_failed_final_edit_falls_back(self):
        with patch.object(bot, "call", return_value={"message_id": 42}) as api:
            progress = bot.Progress(1)
            progress.update("در حال بررسی")
            progress.update("آماده", final=True)
            self.assertEqual(["sendMessage", "editMessageText", "editMessageText"],
                             [item.args[0] for item in api.call_args_list])
            api.side_effect = [bot.BaleAPIError("edit unavailable"), {"message_id": 43}]
            progress.update("نتیجه نهایی", final=True)
            self.assertEqual("sendMessage", api.call_args.args[0])


class CliDeliveryTests(unittest.TestCase):
    def test_cli_reports_actual_annotation_cap_after_generation(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            reviewed = Path(tmp) / "reviewed.docx"
            output = Path(tmp) / "reports"
            doc = Document(); doc.add_paragraph("بخاطر نمونه ها بررسی شد."); doc.save(source)
            before = source.read_bytes()
            with contextlib.redirect_stdout(io.StringIO()):
                result = review_cli.main([str(source), "--annotate", str(reviewed),
                                          "--max-comments", "1", "--output-dir", str(output)])
            self.assertEqual(0, result)
            self.assertEqual(before, source.read_bytes())
            self.assertGreater(len(Document(reviewed).comments), 1)
            import json
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(1, report["stats"]["annotation_summary"]["applied"])
            self.assertEqual(1, report["stats"]["annotation_summary"]["max_comments"])
            self.assertGreater(report["stats"]["annotation_summary"]["unapplied"], 0)
            self.assertEqual(len(report["findings"]), len(Document(reviewed).comments))
            self.assertTrue(report["stats"]["unified_word_review"]["output_created"])


if __name__ == "__main__":
    unittest.main()
