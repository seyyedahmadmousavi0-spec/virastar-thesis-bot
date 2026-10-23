"""All safely located findings retain independent comments and yellow witnesses."""
import copy
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from analysis_model import Issue
from doc_guard import sha256_file
from safe_annotator import annotate_document


def issue(target="", paragraph=0, rule="TEST", **values):
    return Issue(rule, message="این مورد با توجه به متن اصلی بازبینی شود.",
                 suggestion="پیشنهاد را پیش از تغییر با متن اصلی مقایسه کنید.",
                 target_text=target, par_index=paragraph, **values)


def ranges(paragraph):
    """Read actual OOXML comment selections, including crossing/nested ranges."""
    active = set()
    selected = {}
    for child in paragraph._p:
        if child.tag == qn("w:commentRangeStart"):
            key = int(child.get(qn("w:id")))
            active.add(key)
            selected[key] = ""
        elif child.tag == qn("w:commentRangeEnd"):
            active.discard(int(child.get(qn("w:id"))))
        elif child.tag == qn("w:r"):
            from docx.text.run import Run
            text = Run(child, paragraph).text
            for key in active:
                selected[key] += text
    return selected


class CompleteWordAnnotationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="all_word_annotations_")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.output = self.root / "review.docx"

    def annotate(self, document, findings, **options):
        original = document.element.xml
        result = annotate_document(document, findings, self.output, **options)
        self.assertEqual(original, document.element.xml)
        summary = document._last_annotation_summary
        self.assertTrue(summary["output_created"])
        self.assertEqual(len(findings), summary["requested"])
        self.assertEqual(summary["requested"], summary["applied"] + summary["unapplied"])
        self.assertEqual(summary["unapplied"], sum(summary["reasons"].values()))
        self.assertEqual(summary["applied"], summary["highlighted"])
        self.assertEqual(summary["applied"], summary["exact_applied"] + summary["paragraph_applied"])
        return result, Document(self.output), summary

    def test_default_does_not_stop_at_250_comments(self):
        document = Document()
        findings = []
        for index in range(300):
            document.add_paragraph(f"شاهد مستقل {index:04d} و متن پژوهش")
            findings.append(issue(f"شاهد مستقل {index:04d}", paragraph=index))
        result, reviewed, summary = self.annotate(document, findings)
        self.assertTrue(result)
        self.assertEqual(300, len(reviewed.comments))
        self.assertEqual(300, summary["exact_applied"])
        self.assertEqual({}, summary["reasons"])
        self.assertIsNone(summary["max_comments"])
        self.assertIsNone(summary["per_par"])
        for original, paragraph in zip(document.paragraphs, reviewed.paragraphs):
            self.assertEqual(original.text, paragraph.text)
            witness = " ".join(original.text.split()[:3])
            self.assertEqual(witness, "".join(run.text for run in paragraph.runs
                                             if run.font.highlight_color == WD_COLOR_INDEX.YELLOW))
            self.assertEqual([witness], list(ranges(paragraph).values()))

    def test_30_comments_in_one_paragraph_all_keep_their_own_witnesses(self):
        words = [f"مورد{index:03d}" for index in range(30)]
        document = Document()
        document.add_paragraph(" ".join(words))
        _, reviewed, summary = self.annotate(document, [issue(word) for word in words])
        self.assertEqual(30, summary["applied"])
        self.assertEqual(30, len(reviewed.comments))
        self.assertEqual(words, list(ranges(reviewed.paragraphs[0]).values()))

    def test_dense_single_run_keeps_all_600_exact_comment_ranges(self):
        # Unlike a separate-paragraph fixture, this exercises splitting a long
        # single run at every boundary and resolving hundreds of frozen anchors.
        words = [f"witness{index:04d}" for index in range(600)]
        text = " ".join(words)
        document = Document()
        document.add_paragraph(text)
        findings = [issue(word, char_start=index * 12, char_end=index * 12 + 11)
                    for index, word in enumerate(words)]
        _, reviewed, summary = self.annotate(document, findings)
        self.assertEqual(text, reviewed.paragraphs[0].text)
        self.assertEqual(600, summary["exact_applied"])
        self.assertEqual(600, len(reviewed.comments))
        self.assertEqual(words, list(ranges(reviewed.paragraphs[0]).values()))

    def test_nested_overlapping_and_identical_spans_keep_distinct_comments(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("alpha be").bold = True
        paragraph.add_run("ta gamma delta").italic = True
        findings = [issue("alpha beta", rule="FIRST", char_start=0, char_end=10),
                    issue("beta gamma", rule="CROSS", char_start=6, char_end=16),
                    issue("beta", rule="NESTED", char_start=6, char_end=10),
                    issue("beta", rule="SAME", char_start=6, char_end=10)]
        _, reviewed, summary = self.annotate(document, findings)
        self.assertEqual(4, summary["applied"])
        self.assertEqual({0: "alpha beta", 1: "beta gamma", 2: "beta", 3: "beta"},
                         ranges(reviewed.paragraphs[0]))
        self.assertEqual("alpha beta gamma", "".join(run.text for run in reviewed.paragraphs[0].runs
                                                     if run.font.highlight_color == WD_COLOR_INDEX.YELLOW))
        for original, output in zip(("FIRST", "CROSS", "NESTED", "SAME"), reviewed.comments):
            self.assertTrue(output.text.startswith(original + " |"))

    def test_distinct_reasons_for_same_rule_and_span_are_not_discarded(self):
        document = Document()
        document.add_paragraph("پژوهش حاضر")
        findings = [issue("پژوهش"), issue("پژوهش")]
        findings[1].message = "دلیل متفاوت در همین محل"
        _, reviewed, summary = self.annotate(document, findings)
        self.assertEqual(2, len(reviewed.comments))
        self.assertEqual(2, summary["applied"])

    def test_format_finding_uses_explicit_paragraph_scope_not_invented_phrase(self):
        document = Document()
        document.add_paragraph("متن این پاراگراف از نظر جهت و فاصله بررسی می‌شود.")
        _, reviewed, summary = self.annotate(document, [issue(rule="alignment", confidence="B")])
        self.assertEqual(0, summary["exact_applied"])
        self.assertEqual(1, summary["paragraph_applied"])
        self.assertEqual("paragraph", summary["records"][0]["anchor_scope"])
        self.assertIn("[در سطح پاراگراف]", list(reviewed.comments)[0].text)
        self.assertEqual({0: document.paragraphs[0].text}, ranges(reviewed.paragraphs[0]))
        self.assertTrue(all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW
                            for run in reviewed.paragraphs[0].runs if run.text))

    def test_invalid_target_or_offsets_never_fall_back_to_paragraph_scope(self):
        document = Document()
        document.add_paragraph("alpha beta")
        findings = [issue("missing"), issue("beta", char_start=0, char_end=4),
                    issue("", char_start=1, char_end=3), issue("", paragraph=-1),
                    issue("", char_start=-2, char_end=-2), issue("beta", char_end=10)]
        result, reviewed, summary = self.annotate(document, findings)
        self.assertFalse(result)
        self.assertEqual(0, len(reviewed.comments))
        self.assertEqual({"uncertain_location": 5, "location_unavailable": 1}, summary["reasons"])

    def test_existing_comments_bookmarks_and_tabs_are_preserved(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("alpha\tbeta\ngamma")
        document.add_comment(paragraph.runs, "یادداشت اصلی نویسنده", author="Original", initials="OA")
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), "42")
        bookmark_start.set(qn("w:name"), "OriginalBookmark")
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), "42")
        paragraph._p.insert(0, bookmark_start)
        paragraph._p.append(bookmark_end)
        source = self.root / "original.docx"
        document.save(source)
        document = Document(source)
        document._review_source_path = str(source)
        document._review_source_hash = sha256_file(source)
        before = source.read_bytes()
        _, reviewed, summary = self.annotate(document, [issue("beta", rule="EXACT"),
                                                      issue("alpha\tbeta", rule="WITH_TAB")])
        self.assertEqual(before, source.read_bytes())
        self.assertEqual(2, summary["applied"])
        self.assertEqual("alpha\tbeta\ngamma", reviewed.paragraphs[0].text)
        self.assertEqual(3, len(reviewed.comments))
        self.assertEqual("یادداشت اصلی نویسنده", list(reviewed.comments)[0].text)
        self.assertEqual("Original", list(reviewed.comments)[0].author)
        self.assertEqual({0: "alpha\tbeta\ngamma", 1: "beta", 2: "alpha\tbeta"},
                         ranges(reviewed.paragraphs[0]))
        self.assertEqual(1, len(reviewed.paragraphs[0]._p.findall(qn("w:bookmarkStart"))))
        self.assertEqual(1, len(reviewed.paragraphs[0]._p.findall(qn("w:bookmarkEnd"))))

    def test_overlapping_highlights_preserve_per_character_properties_and_run_attributes(self):
        def signature(paragraph):
            result = []
            for run in paragraph.runs:
                properties = copy.deepcopy(run._r.rPr)
                if properties is not None:
                    for highlight in list(properties.findall(qn("w:highlight"))):
                        properties.remove(highlight)
                props = etree.tostring(properties, method="c14n") if properties is not None and len(properties) else b""
                result.extend((character, dict(run._r.attrib), props) for character in run.text)
            return result
        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run("alpha beta gamma")
        run.bold = True
        run.font.size = Pt(15)
        run._r.set(qn("w:rsidR"), "01234567")
        before = signature(paragraph)
        _, reviewed, _ = self.annotate(document, [issue("alpha beta"), issue("beta gamma", rule="OTHER")])
        self.assertEqual(before, signature(reviewed.paragraphs[0]))

    def test_explicit_caps_remain_opt_in_and_account_for_every_skip(self):
        document = Document()
        document.add_paragraph("alpha beta gamma")
        findings = [issue("alpha", rule="ONE"), issue("beta", rule="TWO"), issue("gamma", rule="THREE")]
        _, reviewed, summary = self.annotate(document, findings, max_comments=1)
        self.assertEqual(1, len(reviewed.comments))
        self.assertEqual({"comment_limit": 2}, summary["reasons"])
        _, reviewed, summary = self.annotate(document, findings, per_par=0)
        self.assertEqual(0, len(reviewed.comments))
        self.assertEqual({"paragraph_limit": 3}, summary["reasons"])

    def test_protected_ambiguous_and_complex_stories_remain_explicitly_unapplied(self):
        document = Document()
        paragraph = document.add_paragraph("alpha beta")
        run = paragraph.add_run()
        run._r.append(OxmlElement("w:fldChar"))
        findings = [issue("alpha", protected_overlap=True), issue("alpha", confidence="C"),
                    issue("alpha", source_part="word/header1.xml", story="header"), issue("beta")]
        result, reviewed, summary = self.annotate(document, findings)
        self.assertFalse(result)
        self.assertEqual(0, len(reviewed.comments))
        self.assertEqual({"protected_content": 1, "ambiguous_excluded": 1,
                          "unsupported_story": 1, "unsafe_structure": 1}, summary["reasons"])


if __name__ == "__main__":
    unittest.main()
