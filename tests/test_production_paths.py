"""Exercise complete document jobs; only the Bale transport is replaced."""
import contextlib
import io
import json
import os
os.environ["BALE_SKIP_DOTENV"] = "1"
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import pymupdf
import requests

import bot
import checker
from doc_guard import safe_filename, sha256_file
from report_artifacts import report_data, html_report
import review_cli


class ProductionPathTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_unicode_filenames_leave_room_for_work_and_annotation_prefixes(self):
        for stem in ("پژوهش" * 60, "📄" * 200, "a" * 300):
            name = safe_filename(stem + ".docx")
            self.assertTrue(name.endswith(".docx"))
            self.assertLessEqual(len(name.encode("utf-8")), 180)
            self.assertLessEqual(len(("نسخه-بررسی‌شده-" + name).encode("utf-8")), 255)
            self.assertEqual(name, safe_filename(name))

    def run_document(self, *, name="study.docx", fail_delivery=False):
        source = self.root / "original.docx"
        document = Document()
        document.add_paragraph("بخاطر نمونه ها بررسی شد.")
        document.save(source)
        before = sha256_file(source)
        downloads, delivered = [], {}

        def download(_file_id, destination):
            # Linux limits a filename component by bytes, even on a Windows test host.
            if len(Path(destination).name.encode("utf-8")) > 255:
                raise OSError("Filename too long")
            downloads.append(Path(destination))
            Path(destination).write_bytes(source.read_bytes())

        def send_file(_chat, path, *args):
            if fail_delivery:
                raise bot.BaleAPIError("sendDocument: HTTP 503")
            delivered[Path(path).suffix] = Path(path).read_bytes()
            return {"message_id": 1}

        with patch.object(bot, "download", side_effect=download), \
                patch.object(bot, "send", return_value=True) as send, \
                patch.object(bot, "send_file", side_effect=send_file):
            bot.process_document(123, {"file_name": name, "file_id": "synthetic"},
                                 annotate_requested=True)
        self.assertEqual(before, sha256_file(source))
        self.assertTrue(all(not path.parent.exists() for path in downloads))
        return delivered, send

    def test_complete_word_job_delivers_reports_and_preserves_annotated_text(self):
        delivered, _ = self.run_document(name="پژوهش" * 60 + ".docx")
        self.assertEqual({".html", ".txt", ".docx"}, set(delivered))
        self.assertIn("common_mistake", delivered[".html"].decode("utf-8"))
        result = Document(io.BytesIO(delivered[".docx"]))
        self.assertEqual("بخاطر نمونه ها بررسی شد.", result.paragraphs[0].text)
        self.assertGreater(len(result.comments), 0)

    def test_delivery_failure_does_not_blame_a_valid_document(self):
        _, send = self.run_document(fail_delivery=True)
        message = send.call_args.args[1]
        self.assertIn("ارتباط", message)
        self.assertNotIn("سلامت فایل", message)

    def test_download_transport_failure_is_sanitized(self):
        with patch.object(bot, "call", return_value={"file_path": "file.docx"}), \
                patch.object(bot, "http") as http:
            http.get.side_effect = requests.Timeout("PRIVATE_TOKEN_IN_URL")
            with self.assertRaises(bot.BaleAPIError) as raised:
                bot.download("synthetic", self.root / "download.docx")
        self.assertNotIn("PRIVATE_TOKEN", str(raised.exception))

    def test_unreadable_pdf_cannot_be_reported_as_success_by_disabling_a_rule(self):
        source = self.root / "encrypted.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 40), "Synthetic protected content")
            pdf.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                     owner_pw="owner", user_pw="reader")
        rules = checker.load_rules()
        rules["disabled_rules"] = ["pdf_read_error"]
        _, issues, stats = checker.analyze_pdf(source, rules=rules)
        self.assertEqual("failed", stats.get("analysis_status"))
        self.assertEqual("encrypted", stats.get("analysis_error_code"))
        self.assertIn("pdf_read_error", {issue.rule_id for issue in issues})
        self.assertNotIn("به پایان رسید", checker.build_summary(issues, stats, rules))
        self.assertIn("تکمیل نشد", checker.build_summary(issues, stats, rules))
        data = report_data(issues, stats, rules)
        self.assertEqual("failed", data["stats"]["analysis_status"])
        self.assertIn("تکمیل نشد", html_report(data))
        config = self.root / "rules.json"
        config.write_text(json.dumps(rules, ensure_ascii=False), encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(2, review_cli.main([str(source), "--rules", str(config)]))


if __name__ == "__main__":
    unittest.main()
