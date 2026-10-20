# -*- coding: utf-8 -*-
"""
تست تضمین تغییرناپذیری فایل اصلی کاربر (Immutability & Zero-Mutation Test).
راستی‌آزمایی اینکه فایل ورودی (DOCX / PDF / TXT) چه در شرایط عادی و چه در حین خطا،
کوچک‌ترین تغییر بایتی نداشته و هش SHA-256 آن دقیقاً یکسان باقی می‌ماند.
"""
import os
import shutil
import tempfile
import unittest
from docx import Document

from doc_guard import DocumentGuard, DocumentIntegrityError, sha256_file
import checker


class TestDocumentImmutability(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="test_immutability_")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_docx_analysis_and_annotation_leaves_original_untouched(self):
        """تست عدم تغییر فایل اصلی Word پس از تحلیل و کامنت‌گذاری."""
        docx_path = os.path.join(self.temp_dir, "sample.docx")
        doc = Document()
        doc.add_heading("فصل اول: مقدمه", level=1)
        doc.add_paragraph("این یک متن نمونه است که در سال 2020 نوشته شده و می رود تا بررسی شود.")
        doc.add_paragraph("غلظت محلول 15 mg/dL و دما 25 °C بود.")
        doc.save(docx_path)

        pre_hash = sha256_file(docx_path)
        pre_size = os.path.getsize(docx_path)

        # اجرای تحلیل
        d, issues, stats = checker.analyze(docx_path, mode="thesis")
        self.assertGreater(len(issues), 0)

        # ایجاد نسخه خروجی کامنت‌گذاری‌شده در مسیر دیگر
        out_docx = os.path.join(self.temp_dir, "annotated_sample.docx")
        checker.annotate(d, issues, out_docx)

        # فایل خروجی باید به درستی ایجاد شده باشد
        self.assertTrue(os.path.exists(out_docx))

        # بررسی هش و حجم فایل اصلی ورودی
        post_hash = sha256_file(docx_path)
        post_size = os.path.getsize(docx_path)

        self.assertEqual(pre_hash, post_hash, "هش فایل اصلی Word پس از بررسی تغییر کرده است!")
        self.assertEqual(pre_size, post_size, "حجم فایل اصلی Word تغییر کرده است!")

    def test_document_guard_detects_mutation(self):
        """تست اینکه اگر فایلی تغییر داده شود، DocumentGuard با خطا مانع می‌شود."""
        sample_path = os.path.join(self.temp_dir, "protected.txt")
        with open(sample_path, "w", encoding="utf-8") as f:
            f.write("محتوای اصلی و غیرقابل تغییر")

        with self.assertRaises(DocumentIntegrityError):
            with DocumentGuard(sample_path) as guard:
                # شبیه‌سازی دستکاری غیرمجاز در فایل اصلی
                with open(sample_path, "a", encoding="utf-8") as f:
                    f.write("تغییر غیرمجاز!")


if __name__ == "__main__":
    unittest.main()
