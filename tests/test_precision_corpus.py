# -*- coding: utf-8 -*-
"""Labelled precision/recall corpus and end-to-end placement of every finding.

"_" in the corpus stands for ZWNJ. CLEAN lines are correct academic Persian
and English and must produce no finding. Each ERRORS line must produce at least
its expected rule ids and nothing else.
"""
import os
os.environ["BALE_SKIP_DOTENV"] = "1"
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Cm

import checker
import fa_spell
import kashan_checker
import structure_checks
from safe_annotator import annotate_document
from unified_word_review import finalize_word_review

Z = "‌"


def z(text):
    return text.replace("_", Z)


CLEAN = [z(text) for text in [
    "این پژوهش با هدف بررسی تأثیر آموزش مجازی بر یادگیری دانشجویان کارشناسی انجام شده است.",
    "داده_ها با استفاده از پرسش_نامهٔ استاندارد گردآوری و با نرم_افزار SPSS نسخهٔ ۲۶ تحلیل شدند.",
    "نتایج نشان داد که میانگین نمرهٔ گروه آزمایش به_طور معناداری بیشتر از گروه کنترل بود (p < 0.05).",
    "یافته_های این مطالعه با نتایج پژوهش_های پیشین همسو است (احمدی، ۱۳۹۸؛ Smith, 2019).",
    "جامعهٔ آماری شامل ۳۵۰ دانشجوی دانشگاه کاشان در سال تحصیلی ۱۴۰۱–۱۴۰۲ بود.",
    "پایایی ابزار با ضریب آلفای کرونباخ ۰٫۸۷ تأیید شد.",
    "پژوهشگران دیگر نیز به نتایج مشابهی دست یافته_اند و این یافته را تأیید کرده_ایم.",
    "بزرگ_ترین محدودیت این پژوهش، دسترسی دشوار به داده_های بالینی بود.",
    "غلظت گلوکز ۱۲۰ mg/dL و دمای نگهداری ۴ °C در نظر گرفته شد.",
    "از ترکیب H2O و NaCl در دمای 25 °C استفاده شد.",
    "خانهٔ سالمندان، محیطی مناسب برای اجرای این مداخله بود.",
    "نامه_ای رسمی برای دریافت مجوز اخلاق به کمیتهٔ پژوهشی ارسال شد.",
    "آیا آموزش مجازی می_تواند جایگزین آموزش حضوری شود؟",
    "کم_کم و آهسته_آهسته مشخص شد که متغیر سن نیز بر نتایج اثرگذار است.",
    "داده_ها در ماه می ۲۰۲۱ گردآوری شدند.",
    "دانشجویان در خانه نبودند و پژوهشگران به آنجا رفتند.",
    "این کتابها و کتاب_ها در کتابخانهٔ مرکزی دانشگاه موجود است.",
    "ابهام در الهام_بخشی نتایج تا انتهای پژوهش باقی ماند.",
    "این مقاله در مجلهٔ علوم تربیتی، دورهٔ ۱۲، شمارهٔ ۳، صص. ۴۵–۶۰ منتشر شده است.",
    "تاریخ دفاع: ۱۴۰۲/۰۶/۱۵",
    "به گفتهٔ یکی از شرکت_کنندگان: «این دوره برای من بسیار مفید بود».",
    "اطلاعات بیشتر در https://doi.org/10.1016/j.cell.2021.05.001 آمده است.",
    "ضریب تعیین مدل (R² = 0.71) نشان داد که متغیرهای مستقل بخش عمدهٔ واریانس را تبیین می_کنند.",
    "این یافته با پژوهش Johnson و همکاران (2018) همخوانی دارد.",
    "در سال ۲۰۲۰ میلادی، همه_گیری کووید-۱۹ آموزش را دگرگون کرد.",
    "مسئولیت صحت مطالب این پایان_نامه بر عهدهٔ اینجانب است.",
    "The results were consistent with previous studies in this field.",
    "Data were analyzed using a mixed-effects model.",
]]

ERRORS = [(z(text), set(rules)) for text, rules in [
    ("این پژوهش بر روی دانشجويان انجام شد.", {"arabic_letters"}),
    ("نتایج نشان میدهد که روش جدید بهتر است.", {"mi_prefix"}),
    ("پژوهشگران می توانند از این روش استفاده کنند.", {"mi_prefix"}),
    ("این روش بزرگ تر و موثر تر است.", {"comparative_space", "common_mistake"}),
    ("بعنوان مثال، این متغیر مهم است.", {"common_mistake"}),
    ("این موضوع خیلی مهم است.", {"informal_word"}),
    ("داده ها جمع آوری شد .", {"plural_space", "space_before_punct"}),
    ("نتایج،نشان داد که فرضیه تأیید شد.", {"no_space_after_punct"}),
    ("آیا این روش موثر است?", {"latin_punct", "common_mistake"}),
    ('اصطلاح "یادگیری" در این پژوهش تعریف شد.', {"straight_quotes"}),
    ("این  پژوهش در دو مرحله انجام شد.", {"double_space"}),
    ("نتایج این این پژوهش مثبت بود.", {"repeated_word"}),
    ("آهسته آهسته میزان مشارکت افزایش یافت.", {"repeated_word"}),
    ("در سال 1400 این پژوهش انجام شد.", {"latin_digits"}),
    ("تعداد ٣٥ نفر در مطالعه شرکت کردند.", {"arabic_digits"}),
    ("به نظر من این روش بهتر است.", {"ACAD_FIRST_PERSON"}),
    ("خوشبختانه نتایج مثبت بود.", {"ACAD_EMOTIONAL"}),
    ("سیر صعودی رو به بالا در داده ها دیده شد.", {"ACAD_PLEONASM", "plural_space"}),
    ("این یافته ( که مهم است ) تأیید شد.", {"paren_space"}),
    ("خانه ی سالمندان محل پژوهش بود.", {"ye_space"}),
    ("پژوهشگران در این زمینه مطالعه کرده اند.", {"verb_suffix_space"}),
    ("این مقالات در سال گذشته منتشر شده اند.", {"verb_suffix_space"}),
    ("دادههای این پژوهش کافی بود.", {"he_suffix_join"}),
    ("پژوهشگران این روش را بررسی کردهاند.", {"he_suffix_join"}),
    ("These result is significant.", {"EN_SVA"}),
    ("This is a important finding in the study.", {"EN_ARTICLE"}),
]]


class PrecisionCorpusTests(unittest.TestCase):
    rules = checker.load_rules(mode="thesis")

    def test_correct_text_has_no_findings(self):
        for text in CLEAN:
            with self.subTest(text=text.replace(Z, "_")):
                found = checker.check_text(0, text, "", self.rules)
                self.assertEqual([], [(i.rule_id, i.target_text) for i in found])

    def test_every_seeded_error_is_found_without_extras(self):
        for text, expected in ERRORS:
            with self.subTest(text=text.replace(Z, "_")):
                found = {i.rule_id for i in checker.check_text(0, text, "", self.rules)}
                self.assertEqual(expected, found)

    def test_concrete_suggestions(self):
        cases = {"یافته های پژوهش آمده است.": z("یافته_های"), "نتایج نشان میدهد.": z("می_ده"),
                 "این کار را کرده اند.": z("کرده_اند"), "دادههای پژوهش": z("داده_های"),
                 "در سال 1400 انجام شد.": "ل ۱۴۰۰", "آیا موثر است?": "؟"}
        for text, suggestion in cases.items():
            with self.subTest(text=text):
                found = checker.check_text(0, text, "", self.rules)
                self.assertIn(suggestion, [i.suggestion for i in found])

    def test_reduplication_is_a_join_suggestion_not_a_redundant_word(self):
        found = checker.check_text(0, "دانشجویان یکی یکی وارد شدند.", "", self.rules)
        self.assertEqual(["repeated_word"], [i.rule_id for i in found])
        self.assertIn("نیم", found[0].message)
        self.assertNotIn("اضافی", found[0].message)

    def test_citation_messages_are_persian(self):
        for issue in checker.check_text(0, "These result is significant.", "", self.rules):
            self.assertRegex(issue.message, "[؀-ۿ]")

    def test_kashan_heading_matching_ignores_ezafe_marks(self):
        for heading in ("تعهدنامهٔ اصالت اثر", "تعهدنامۀ اصالت اثر", z("تعهدنامه_ی اصالت اثر")):
            self.assertEqual("تعهدنامه اصالت اثر", kashan_checker._norm(heading))


class SpellingTests(unittest.TestCase):
    rules = checker.load_rules(mode="thesis")

    def found(self, text):
        return {(i.target_text, i.suggestion) for i in checker.check_text(0, text, "", self.rules)
                if i.rule_id == "spelling"}

    def test_typos_get_a_single_correction(self):
        text = "نتایج نشان داد که دانشجیان از آموزش رضایت داشتند و نمرات افزیش یافت."
        self.assertEqual({("دانشجیان", "دانشجویان"), ("افزیش", "افزایش")}, self.found(text))

    def test_two_forms_of_one_word_are_offered_together(self):
        self.assertEqual({("پژوهیش", "پژوهش / پژوهشی")}, self.found("نتایج این پژوهیش مهم بود."))

    def test_inflections_compounds_names_and_terms_are_not_flagged(self):
        for text in CLEAN + [z("دانشجویان و شرکت_کنندگان بیدارشدن و آزادکردن داده_ها را برخط پیگیری کردند."),
                             "مقیاس لیکرت و آلفای کرونباخ در پرسشنامه به کار رفت."]:
            with self.subTest(text=text.replace(Z, "_")):
                self.assertEqual(set(), self.found(text))

    def test_repeated_unknown_word_is_treated_as_a_term(self):
        counts = fa_spell.document_counts(["افزیش"] * 3)
        found = checker.check_text(0, "میزان افزیش بود.", "", self.rules, spelling_counts=counts)
        self.assertNotIn("spelling", {i.rule_id for i in found})

    def test_specific_rule_wins_over_spelling(self):
        found = checker.check_text(0, "بعنوان نمونه", "", self.rules)
        self.assertEqual(["common_mistake"], [i.rule_id for i in found])

    def test_protected_terms_and_quotes_are_skipped(self):
        rules = dict(self.rules, protected_terms=["افزیش"])
        self.assertEqual([], [i for i in checker.check_text(0, "میزان افزیش بود.", "", rules) if i.rule_id == "spelling"])
        self.assertEqual(set(), self.found("پاسخ داد: «دانشجیان خسته بودند»."))


class ConsistencyTests(unittest.TestCase):
    def paragraphs(self, texts):
        document = Document()
        for text in texts:
            document.add_paragraph(z(text))
        return document.paragraphs

    def test_uncited_reference_is_reported_on_its_entry(self):
        paragraphs = self.paragraphs([
            "این یافته با پژوهش احمدی (۱۳۹۸) و (Smith, 2019) همسو است.",
            "منابع",
            "احمدی، علی. (۱۳۹۸). آموزش مجازی. تهران.",
            "Smith, J. (2019). Online learning. Journal, 1, 1-2.",
            "Taylor, M. (2017). Blended learning. London.",
        ])
        issues = structure_checks.uncited_references(paragraphs, 1, len(paragraphs))
        self.assertEqual([4], [i.par_index for i in issues])
        self.assertIn("Taylor", issues[0].message)

    def test_unrecognised_citation_style_abstains(self):
        paragraphs = self.paragraphs(["متن بدون هیچ ارجاعی.", "منابع"] +
                                     [f"Author{n}, A. (20{10 + n}). Title." for n in range(4)])
        self.assertEqual([], structure_checks.uncited_references(paragraphs, 1, len(paragraphs)))

    def test_numbered_references(self):
        paragraphs = self.paragraphs(["نتایج همسو بود [1، 2].", "منابع", "[1] A.", "[2] B.", "[3] C."])
        issues = structure_checks.uncited_references(paragraphs, 1, len(paragraphs), [1, 2, 3])
        self.assertEqual([4], [i.par_index for i in issues])

    def test_mixed_spelling_of_one_term(self):
        paragraphs = self.paragraphs(["از پرسشنامه استفاده شد.", "پرسش_نامهٔ دوم نیز توزیع شد."])
        issues = structure_checks.inconsistent_terms(paragraphs)
        self.assertEqual(1, len(issues))
        self.assertEqual(z("پرسش_نامه"), issues[0].target_text)
        self.assertEqual([], structure_checks.inconsistent_terms(self.paragraphs(["کتابها و کتاب_ها"])))

    def test_caption_never_mentioned_in_text(self):
        paragraphs = self.paragraphs(["نتایج در جدول ۴–۱ آمده است.", "جدول ۴–۱: آمار توصیفی",
                                      "شکل ۴–۲: نمودار میانگین"])
        issues = structure_checks.unreferenced_captions(paragraphs)
        self.assertEqual([2], [i.par_index for i in issues])
        self.assertEqual([], structure_checks.unreferenced_captions(self.paragraphs(["جدول ۱: الف"])))


class DocumentPlacementTests(unittest.TestCase):
    """The bot's default Word path puts every finding near its own location."""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="precision_doc_")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def review(self, document, mode="thesis"):
        source, output = self.root / "thesis.docx", self.root / "reviewed.docx"
        document.save(source)
        rules = checker.load_rules(mode=mode)
        parsed, issues, stats = checker.analyze(str(source), mode, rules)
        annotate_document(parsed, issues, output)
        stats["annotation_summary"] = parsed._last_annotation_summary
        result = finalize_word_review(source, output, issues, stats, rules)
        reviewed = Document(output)
        self.assertEqual(len(issues), len(reviewed.comments))
        self.assertEqual(len(issues), result["comments"])
        return issues, result

    def test_uncaptioned_table_and_section_margin_have_their_own_anchors(self):
        document = Document()
        document.add_paragraph("عنوان پایان_نامه".replace("_", Z), style="Title")
        document.add_paragraph("جدول ۱: نمونهٔ دارای عنوان")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "الف"
        document.add_paragraph("متن پیش از جدول بدون عنوان در این بخش آمده است.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "ب"
        document.add_section(WD_SECTION.NEW_PAGE).top_margin = Cm(1)
        document.add_paragraph("فصل اول: کلیات", style="Heading 1")
        issues, result = self.review(document)
        caption = [i for i in issues if i.rule_id == "caption_missing"]
        self.assertEqual(1, len(caption))
        self.assertEqual(2, caption[0].par_index)
        margins = [record for record in result["context_records"]
                   if issues[record["index"]].rule_id == "page_margin"
                   and issues[record["index"]].loc == "بخش 2"]
        self.assertTrue(margins)
        self.assertTrue(all(record["scope"] == "section" for record in margins))
        self.assertTrue(all(record["anchor_par_index"] == 4 for record in margins))

    def test_style_wide_formatting_problem_is_one_comment(self):
        document = Document()
        for index in range(12):
            paragraph = document.add_paragraph(z(f"این پاراگراف شمارهٔ {index} متنی طولانی برای بررسی قالب_بندی متن اصلی است."))
            paragraph.paragraph_format.line_spacing = 1.0
        issues, _ = self.review(document)
        spacing = [i for i in issues if i.rule_id == "line_spacing"]
        self.assertEqual(1, len(spacing))
        self.assertIn("12 پاراگراف", spacing[0].message)
        self.assertIn("12", spacing[0].snippet)


if __name__ == "__main__":
    unittest.main()
