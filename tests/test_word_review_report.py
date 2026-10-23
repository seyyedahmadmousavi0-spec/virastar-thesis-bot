"""Every returned finding has a yellow, commented entry in a separate Word report."""
import copy
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from analysis_model import Issue
from doc_guard import DocumentIntegrityError, sha256_file
from word_review_report import TITLE, WordReviewReportError, write_word_review


class WordReviewReportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="word_review_")
        self.root = Path(self.directory.name)
        self.output = self.root / "full_review.docx"

    def tearDown(self):
        self.directory.cleanup()

    def text(self, document):
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def yellow(self, document):
        return [run.text for paragraph in document.paragraphs for run in paragraph.runs
                if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]

    def test_every_finding_has_its_own_comment_beyond_250_including_manual_cases(self):
        issues = [Issue("TEST", message=f"دلیل مستقل {index}", suggestion=f"پیشنهاد مستقل {index}",
                        loc=f"پاراگراف {index + 1}", confidence=("A", "B", "C")[index % 3])
                  for index in range(263)]
        issues[0].protected_overlap = True
        issues[1].story = "footnote"
        issues[2].loc = ""
        before = copy.deepcopy(issues)
        stats = {"annotation_summary": {
            "requested": 263, "applied": 25, "unapplied": 238, "output_created": True,
            "records": [{"index": 0, "rule_id": "TEST", "status": "unapplied", "reason": "protected_content"},
                        {"index": 1, "rule_id": "TEST", "status": "unapplied", "reason": "unsupported_story"},
                        {"index": 2, "rule_id": "TEST", "status": "unapplied", "reason": "uncertain_location"}]}}
        original_stats = copy.deepcopy(stats)
        result = write_word_review(self.output, issues, stats, {})
        document = Document(self.output)
        self.assertEqual(263, result["comments"])
        self.assertEqual(263, result["highlighted_entries"])
        self.assertEqual(263, len(document.comments))
        self.assertEqual(263, len(self.yellow(document)))
        self.assertEqual(263, len(document.element.xpath("//w:commentRangeStart")))
        self.assertIn("مورد 263", list(document.comments)[-1].text)
        full = self.text(document)
        self.assertIn("دلیل مستقل 262", full)
        self.assertIn("پیشنهاد مستقل 262", full)
        self.assertIn("این مورد مبهم است؛ فقط بررسی دستی", full)
        self.assertIn("تداخل با محتوای حفاظت‌شده", full)
        self.assertIn("درج امن کامنت در این بخش پشتیبانی نمی‌شود", full)
        self.assertIn("محل دقیق قابل تأیید نیست", full)
        self.assertIn("بررسی در سطح سند", full)
        self.assertEqual(before, issues)
        self.assertEqual(original_stats, stats)

    def test_scientific_unicode_is_quoted_exactly_and_original_file_is_untouched(self):
        source = self.root / "thesis.docx"
        scientific = "x₁ = 10⁻³ µm; E = mc²; α + β; [۱۲]؛ می‌رود؛ ك ي؛ اَ؛ 😀"
        original = Document()
        original.add_paragraph(scientific)
        original.add_table(rows=1, cols=1).cell(0, 0).text = "مقدار ۲٫۵"
        original.save(source)
        before_bytes = source.read_bytes()
        before_hash = sha256_file(source)
        issue = Issue("TEST", message="این عبارت ممکن است نیازمند بازبینی باشد.",
                      suggestion="تنها پس از بررسی انسانی تغییر دهید.", confidence="B",
                      target_text=scientific, location_status="exact", char_start=0,
                      char_end=len(scientific), loc="body / paragraph 1")
        result = write_word_review(self.output, [issue], {"file_sha256": before_hash}, {}, source_path=source)
        document = Document(self.output)
        full = self.text(document)
        self.assertIn("شاهد دقیق از متن اصلی: «" + scientific + "»", full)
        self.assertIn("متن اصلی / پاراگراف 1", full)
        self.assertEqual(before_bytes, source.read_bytes())
        self.assertEqual(before_hash, sha256_file(source))
        self.assertEqual(TITLE, document.core_properties.title)
        self.assertIn("گزارش مستقل بازبینی", full)
        self.assertIn("نسخهٔ پایان‌نامه یا متن اصلاح‌شدهٔ آن نیست", full)
        self.assertEqual(str(self.output), result["path"])
        self.assertEqual(1, len(document.comments))
        self.assertIn("ممکن است", list(document.comments)[0].text)
        for paragraph in document.paragraphs:
            self.assertIsNotNone(paragraph._p.pPr.find(qn("w:bidi")))
        self.assertIsNotNone(list(document.comments)[0].paragraphs[0]._p.pPr.find(qn("w:bidi")))

    def test_unverified_target_is_not_presented_as_exact_evidence(self):
        issue = Issue("missing_section", message="بخش مورد انتظار در متن استخراج‌شده پیدا نشد.",
                      confidence="B", target_text="شاهد نامطمئن", snippet="متن پیرامونی")
        write_word_review(self.output, [issue], {}, {})
        full = self.text(Document(self.output))
        self.assertNotIn("شاهد دقیق از متن اصلی", full)
        self.assertNotIn("شاهد نامطمئن", full)
        self.assertIn("متن پیرامونی ثبت‌شده؛ محل دقیق تأیید نشده", full)

    def test_paragraph_anchor_and_university_context_remain_explicit(self):
        issue = Issue("KASHAN", confidence="B", block_id="b1", message="قالب پاراگراف را بررسی کنید.")
        stats = {"annotation_summary": {"requested": 1, "applied": 1,
                 "records": [{"index": 0, "rule_id": "KASHAN", "block_id": "b1",
                              "status": "applied", "anchor_scope": "paragraph"}]},
                 "university_profile": {"id": "kashan", "title": "دانشگاه کاشان",
                    "degree": "phd", "degree_label": "دکتری", "document_scope": "fragment",
                    "reference_style": "APA", "limitations": ["جزئیات قالب نیازمند بررسی دستی است."]}}
        write_word_review(self.output, [issue], stats, {})
        full = self.text(Document(self.output))
        self.assertIn("کل پاراگراف کامنت و زرد شد؛ این نشانه محل دقیق یک عبارت نیست", full)
        self.assertNotIn("در محل تأییدشدهٔ نسخهٔ بازبینی سند اصلی کامنت شد", full)
        self.assertIn("پروفایل دانشگاهی: دانشگاه کاشان", full)
        self.assertIn("مقطع: دکتری", full)
        self.assertIn("اعتبار رسمی و به‌روز بودن مستقلاً تأیید نشده", full)
        self.assertIn("بخش یا متن محدود", full)

    def test_xml_markup_and_forbidden_control_characters_remain_visible_data(self):
        issue = Issue("<RULE>&", message="<w:r>متن & داده</w:r>\x00\x0b\ud800",
                      suggestion="اَ µ ² ₁ می‌رود \ufffe", confidence="C")
        write_word_review(self.output, [issue], {}, {})
        document = Document(self.output)
        full = self.text(document)
        self.assertIn("<w:r>متن & داده</w:r>⟦U+0000⟧⟦U+000B⟧⟦U+D800⟧", full)
        self.assertIn("اَ µ ² ₁ می‌رود ⟦U+FFFE⟧", full)
        self.assertIn("<RULE>&", list(document.comments)[0].text)
        with zipfile.ZipFile(self.output) as archive:
            xml = archive.read("word/document.xml")
            self.assertIn(b"&lt;w:r&gt;", xml)
            self.assertNotIn(b"<!DOCTYPE", xml)
            self.assertNotIn(b"<!ENTITY", xml)

    def test_empty_report_is_valid_with_no_false_clean_bill_for_partial_analysis(self):
        result = write_word_review(self.output, [], {
            "analysis_status": "partial", "ambiguous_candidates": 2,
            "coverage": {"footnote": {"status": "present"}},
            "private_path": "C:/secret/customer/file.docx",
            "token": "private-token", "arbitrary": "private-metadata"}, {})
        document = Document(self.output)
        self.assertEqual(0, result["comments"])
        self.assertEqual(0, len(document.comments))
        self.assertEqual([], self.yellow(document))
        text = self.text(document)
        self.assertIn("در بخش‌های بررسی‌شده", text)
        self.assertIn("بررسی همهٔ محتوای سند تکمیل نشده", text)
        self.assertIn("2 نامزد مبهم", text)
        self.assertIn("پوشش پاورقی: وجود دارد؛ بررسی نشد", text)
        for secret in ("C:/secret", "private-token", "private-metadata"):
            self.assertNotIn(secret, text)

    def test_source_output_aliases_are_rejected_before_publication(self):
        source = self.root / "input.docx"
        Document().save(source)
        original = source.read_bytes()
        with self.assertRaises(DocumentIntegrityError):
            write_word_review(source, [Issue("TEST")], {}, {}, source_path=source)
        self.assertEqual(original, source.read_bytes())
        hardlink = self.root / "hardlink.docx"
        try:
            os.link(source, hardlink)
        except OSError:
            return
        with self.assertRaises(DocumentIntegrityError):
            write_word_review(hardlink, [Issue("TEST")], {}, {}, source_path=source)
        self.assertEqual(original, source.read_bytes())

    def test_explicit_budgets_and_validation_failures_preserve_existing_output(self):
        sentinel = b"previous report remains until complete replacement"
        self.output.write_bytes(sentinel)
        with patch("word_review_report.MAX_REVIEW_FINDINGS", 2):
            with self.assertRaises(WordReviewReportError):
                write_word_review(self.output, (Issue("TEST") for _ in range(3)), {}, {})
        self.assertEqual(sentinel, self.output.read_bytes())
        with patch("word_review_report.MAX_REVIEW_CHARACTERS", 1):
            with self.assertRaises(WordReviewReportError):
                write_word_review(self.output, [Issue("TEST", message="full text")], {}, {})
        self.assertEqual(sentinel, self.output.read_bytes())
        with patch("word_review_report.MAX_REVIEW_BYTES", 1):
            with self.assertRaises(WordReviewReportError):
                write_word_review(self.output, [Issue("TEST")], {}, {})
        self.assertEqual(sentinel, self.output.read_bytes())
        with patch("word_review_report.Document", side_effect=[Document(), Document()]):
            with self.assertRaises(DocumentIntegrityError):
                write_word_review(self.output, [Issue("TEST")], {}, {})
        self.assertEqual(sentinel, self.output.read_bytes())

    def test_mismatched_comment_reference_is_not_published(self):
        issue = Issue("TEST", message="یک دلیل", suggestion="یک پیشنهاد")
        write_word_review(self.output, [issue], {}, {})
        broken = Document(self.output)
        broken.element.xpath("//w:commentReference")[0].set(qn("w:id"), "999")
        sentinel = b"previous complete report"
        self.output.write_bytes(sentinel)
        with patch("word_review_report.Document", side_effect=[Document(), broken]):
            with self.assertRaises(DocumentIntegrityError):
                write_word_review(self.output, [issue], {}, {})
        self.assertEqual(sentinel, self.output.read_bytes())

    def test_concurrent_source_change_aborts_before_replacing_output(self):
        import word_review_report

        source = self.root / "thesis.docx"
        Document().save(source)
        sentinel = b"previous complete report"
        self.output.write_bytes(sentinel)
        original_paragraph = word_review_report._paragraph
        mutated = False

        def mutate_source_after_hash(*args, **kwargs):
            nonlocal mutated
            paragraph = original_paragraph(*args, **kwargs)
            if not mutated:
                source.write_bytes(b"a concurrent writer changed the source")
                mutated = True
            return paragraph

        with patch("word_review_report._paragraph", side_effect=mutate_source_after_hash):
            with self.assertRaises(DocumentIntegrityError):
                write_word_review(self.output, [Issue("TEST")], {}, {}, source_path=source)
        self.assertTrue(mutated)
        self.assertEqual(sentinel, self.output.read_bytes())

    def test_long_comment_uses_complete_report_instead_of_cutting_qualifiers(self):
        reason = "ممکن است این عبارت نیازمند بازبینی باشد؛ " * 10
        suggestion = "پیش از هر تغییر نیازمند ارزیابی تخصصی پژوهشگر است. " * 10
        write_word_review(self.output, [Issue("TEST", message=reason, suggestion=suggestion, confidence="C")], {}, {})
        document = Document(self.output)
        self.assertIn(reason, self.text(document))
        self.assertIn(suggestion, self.text(document))
        self.assertIn("دلیل کامل", list(document.comments)[0].text)
        self.assertIn("پیشنهاد کامل", list(document.comments)[0].text)
        self.assertIn("C — مبهم", list(document.comments)[0].text)


if __name__ == "__main__":
    unittest.main()
