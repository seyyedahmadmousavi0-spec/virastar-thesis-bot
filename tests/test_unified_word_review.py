"""Every returned finding belongs to one reviewed source Word, without edits."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from lxml import etree

from analysis_model import Issue
from doc_guard import DocumentIntegrityError, sha256_file
from safe_annotator import annotate_document
from unified_word_review import APPENDIX_TITLE, finalize_word_review
from word_review_report import WordReviewReportError


def finding(rule="TEST", paragraph=-1, target="", **options):
    return Issue(rule, par_index=paragraph, target_text=target,
                 message="این مورد را با متن اصلی بررسی کنید.",
                 suggestion="پس از بررسی دستی دربارهٔ تغییر تصمیم بگیرید.", **options)


class UnifiedWordReviewTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="unified_word_")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "original.docx"
        self.review = self.root / "review.docx"

    def prepare(self, document, issues, *, limit=None, opaque=None):
        document.save(self.source)
        if opaque:
            with zipfile.ZipFile(self.source, "a") as package:
                package.writestr("customXml/opaque.bin", opaque)
        loaded = Document(self.source)
        loaded._review_source_path = str(self.source)
        loaded._review_source_hash = sha256_file(self.source)
        annotate_document(loaded, issues, self.review, max_comments=limit)
        return {"annotation_summary": loaded._last_annotation_summary}

    def finalize(self, issues, stats):
        original = self.source.read_bytes()
        result = finalize_word_review(self.source, self.review, issues, stats)
        self.assertEqual(original, self.source.read_bytes())
        self.assertEqual(len(issues), result["comments"])
        self.assertEqual(len(issues), result["highlighted"])
        self.assertEqual(result["requested"], result["source_comments"] + result["appendix_comments"])
        return Document(self.review), result

    def test_over_25_findings_all_in_one_word_with_global_numbers(self):
        document = Document()
        issues = []
        for index in range(36):
            text = f"شاهد مستقل {index:03d} و محتوای پژوهش"
            document.add_paragraph(text)
            issues.append(finding(f"EXACT{index}", index, text[:15]))
        issues.insert(3, finding("MISSING", confidence="B", loc="کل سند"))
        issues.insert(19, finding("HEADER", source_part="word/header1.xml", story="header", target="عنوان"))
        issues.append(finding("AMBIGUOUS", 0, issues[0].target_text, confidence="C"))
        issues.append(finding("PROTECTED", 1, issues[1].target_text, protected_overlap=True))
        stats = self.prepare(document, issues)
        stats["ambiguous_candidates"] = 7
        reviewed, result = self.finalize(issues, stats)
        self.assertEqual(40, len(reviewed.comments))
        self.assertEqual(36, result["source_comments"])
        self.assertEqual(4, result["appendix_comments"])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs],
                         [paragraph.text for paragraph in reviewed.paragraphs[:36]])
        headings = [paragraph for paragraph in reviewed.paragraphs[36:]
                    if any(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in paragraph.runs)]
        self.assertEqual(["مورد 4 | MISSING | B", "مورد 20 | HEADER | A",
                          "مورد 39 | AMBIGUOUS | C", "مورد 40 | PROTECTED | A"],
                         [paragraph.text for paragraph in headings])
        titles = [paragraph for paragraph in reviewed.paragraphs if paragraph.text == APPENDIX_TITLE]
        self.assertEqual(1, len(titles))
        self.assertTrue(titles[0].paragraph_format.page_break_before)
        self.assertEqual("Normal", titles[0].style.name)
        self.assertIn("فقط بررسی دستی", list(reviewed.comments)[-2].text)
        self.assertIn("مورد 40", list(reviewed.comments)[-1].text)
        self.assertTrue(any("7 نامزد مبهم" in paragraph.text and "جزو 40 یافته" in paragraph.text
                            for paragraph in reviewed.paragraphs))

    def test_all_comments_in_original_needs_no_appendix_and_no_rewrite(self):
        document = Document()
        document.add_paragraph("alpha beta gamma")
        issues = [finding("EXACT", 0, "alpha"), finding("PARAGRAPH", 0)]
        stats = self.prepare(document, issues)
        before = self.review.read_bytes()
        reviewed, result = self.finalize(issues, stats)
        self.assertEqual(before, self.review.read_bytes())
        self.assertEqual(2, len(reviewed.comments))
        self.assertEqual(0, result["appendix_comments"])
        self.assertEqual(1, result["exact_comments"])
        self.assertEqual(1, result["paragraph_comments"])
        self.assertEqual(1, len(reviewed.paragraphs))

    def test_existing_comments_headers_opaque_parts_and_sections_survive(self):
        document = Document()
        paragraph = document.add_paragraph("متن علمی ۱۲۳ x=2")
        document.add_comment(paragraph.runs[0], text="نظر استاد", author="Supervisor")
        document.sections[0].header.paragraphs[0].text = "عنوان اصلی"
        issues = [finding("FIRST", 0, "متن علمی"), finding("MISSING", confidence="B")]
        opaque = b"opaque: \x00\xff untouched bytes"
        stats = self.prepare(document, issues, opaque=opaque)
        with zipfile.ZipFile(self.review) as package:
            original_parts = {name: package.read(name) for name in package.namelist()}
            old_document = etree.fromstring(package.read("word/document.xml"))
            section = etree.tostring(old_document.find(".//" + qn("w:sectPr")), method="c14n")
        reviewed, result = self.finalize(issues, stats)
        self.assertEqual(3, len(reviewed.comments))
        self.assertEqual("نظر استاد", list(reviewed.comments)[0].text)
        self.assertEqual("Supervisor", list(reviewed.comments)[0].author)
        self.assertEqual("متن علمی ۱۲۳ x=2", reviewed.paragraphs[0].text)
        with zipfile.ZipFile(self.review) as package:
            for name, payload in original_parts.items():
                if name not in {"word/document.xml", "word/comments.xml"}:
                    self.assertEqual(payload, package.read(name), name)
            final_document = etree.fromstring(package.read("word/document.xml"))
            self.assertEqual(section, etree.tostring(final_document.find(".//" + qn("w:sectPr")), method="c14n"))
            self.assertEqual(opaque, package.read("customXml/opaque.bin"))
            self.assertEqual(qn("w:sectPr"), list(final_document.find(qn("w:body")))[-1].tag)
        self.assertEqual(1, result["appendix_comments"])

    def test_zero_source_comment_cap_places_every_finding_in_same_appendix(self):
        document = Document()
        document.add_paragraph("alpha beta")
        issues = [finding("FIRST", 0, "alpha"), finding("SECOND", 0, "beta")]
        stats = self.prepare(document, issues, limit=0)
        reviewed, result = self.finalize(issues, stats)
        self.assertEqual(0, result["source_comments"])
        self.assertEqual(2, result["appendix_comments"])
        self.assertEqual(2, len(reviewed.comments))
        self.assertEqual("alpha beta", reviewed.paragraphs[0].text)
        self.assertTrue(all(run.font.highlight_color is None for run in reviewed.paragraphs[0].runs))

    def test_no_findings_keeps_byte_identical_review(self):
        document = Document()
        document.add_paragraph("متن اصلی")
        stats = self.prepare(document, [])
        before = self.review.read_bytes()
        reviewed, result = self.finalize([], stats)
        self.assertEqual(before, self.review.read_bytes())
        self.assertEqual(0, result["comments"])
        self.assertEqual("متن اصلی", reviewed.paragraphs[0].text)

    def test_invalid_audit_cannot_publish_incomplete_review(self):
        document = Document()
        document.add_paragraph("alpha")
        issues = [finding("FIRST", 0, "alpha"), finding("SECOND", confidence="B")]
        stats = self.prepare(document, issues)
        original, review = self.source.read_bytes(), self.review.read_bytes()
        stats["annotation_summary"]["records"][1]["index"] = 0
        with self.assertRaises(DocumentIntegrityError):
            finalize_word_review(self.source, self.review, issues, stats)
        self.assertEqual(original, self.source.read_bytes())
        self.assertEqual(review, self.review.read_bytes())

    def test_tampered_scientific_copy_is_rejected_before_publication(self):
        document = Document()
        document.add_paragraph("alpha beta")
        issues = [finding("FIRST", 0, "alpha"), finding("SECOND", confidence="B")]
        stats = self.prepare(document, issues)
        with zipfile.ZipFile(self.review) as package:
            parts = {item.filename: package.read(item.filename) for item in package.infolist()}
        parts["word/document.xml"] = parts["word/document.xml"].replace(b"beta", b"FAKE")
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            for name, value in parts.items():
                package.writestr(name, value)
        self.review.write_bytes(buffer.getvalue())
        tampered = self.review.read_bytes()
        with self.assertRaisesRegex(DocumentIntegrityError, "scientific document payload"):
            finalize_word_review(self.source, self.review, issues, stats)
        self.assertEqual(tampered, self.review.read_bytes())

    def test_source_alias_and_budget_failure_do_not_replace_inputs(self):
        document = Document()
        document.add_paragraph("alpha")
        issues = [finding("ONE", confidence="B"), finding("TWO", confidence="C")]
        stats = self.prepare(document, issues)
        original, review = self.source.read_bytes(), self.review.read_bytes()
        with self.assertRaises(DocumentIntegrityError):
            finalize_word_review(self.source, self.source, issues, stats)
        with patch("unified_word_review.MAX_REVIEW_FINDINGS", 1):
            with self.assertRaises(WordReviewReportError):
                finalize_word_review(self.source, self.review, issues, stats)
        with patch("unified_word_review.MAX_REVIEW_CHARACTERS", 1):
            with self.assertRaises(WordReviewReportError):
                finalize_word_review(self.source, self.review, issues, stats)
        self.assertEqual(original, self.source.read_bytes())
        self.assertEqual(review, self.review.read_bytes())

    def test_appendix_formatting_uses_word_schema_order_without_heading_style(self):
        document = Document()
        document.add_paragraph("متن اصلی")
        issues = [finding("MISSING", confidence="B")]
        stats = self.prepare(document, issues)
        reviewed, _ = self.finalize(issues, stats)
        title = next(paragraph for paragraph in reviewed.paragraphs if paragraph.text == APPENDIX_TITLE)
        self.assertEqual([qn("w:keepNext"), qn("w:pageBreakBefore"), qn("w:bidi")],
                         [child.tag for child in title._p.find(qn("w:pPr"))])
        entry = next(paragraph for paragraph in reviewed.paragraphs if paragraph.text.startswith("مورد 1 |"))
        self.assertEqual([qn("w:rFonts"), qn("w:b"), qn("w:bCs"), qn("w:sz"),
                          qn("w:szCs"), qn("w:highlight"), qn("w:rtl")],
                         [child.tag for child in entry.runs[0]._r.find(qn("w:rPr"))])
        self.assertIsNone(title._p.find(".//" + qn("w:pStyle")))

    def test_lost_source_yellow_witness_cannot_be_reported_as_complete(self):
        document = Document()
        document.add_paragraph("alpha beta")
        issues = [finding("FIRST", 0, "alpha"), finding("SECOND", confidence="B")]
        stats = self.prepare(document, issues)
        with zipfile.ZipFile(self.review) as package:
            parts = {item.filename: package.read(item.filename) for item in package.infolist()}
        parts["word/document.xml"] = parts["word/document.xml"].replace(b'w:val="yellow"', b'w:val="green"')
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            for name, value in parts.items():
                package.writestr(name, value)
        self.review.write_bytes(buffer.getvalue())
        before = self.review.read_bytes()
        with self.assertRaisesRegex(DocumentIntegrityError, "yellow witness"):
            finalize_word_review(self.source, self.review, issues, stats)
        self.assertEqual(before, self.review.read_bytes())


if __name__ == "__main__":
    unittest.main()
