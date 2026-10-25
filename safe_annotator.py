"""Conservative annotation on an independent package, with atomic publication."""
import copy
import hashlib
import unicodedata
import io
import os
import tempfile
import zipfile
from bisect import bisect_left, bisect_right
from collections import Counter, defaultdict
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from doc_guard import ensure_distinct_output, sha256_file, DocumentIntegrityError

def _simple_runs(p):
    # Existing comment/bookmark boundaries are zero-width and must survive.
    # Fields, drawings, hyperlinks, tracked edits and math need another mapper.
    allowed = {qn("w:pPr"), qn("w:r"), qn("w:commentRangeStart"),
               qn("w:commentRangeEnd"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}
    if any(child.tag not in allowed for child in p._p):
        return False
    for run in p.runs:
        for child in run._r:
            if child.tag in {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:cr")}:
                continue
            if child.tag == qn("w:br") and child.get(qn("w:type"), "textWrapping") == "textWrapping":
                continue
            if child.tag == qn("w:commentReference") and not run.text:
                continue
            return False
    return True


def _split_runs_once(paragraph, intervals):
    """Freeze run anchors after splitting every requested boundary just once.

    Every new run keeps the original run attributes, properties and non-text
    leaves. Comment ranges can then overlap without invalidating earlier run
    objects, which happened when each finding split a previous finding's runs.
    """
    original_text = paragraph.text
    boundaries = sorted({position for left, right in intervals for position in (left, right)})
    cursor = 0
    for run in list(paragraph.runs):
        text = run.text
        first = bisect_right(boundaries, cursor)
        last = bisect_left(boundaries, cursor + len(text))
        cuts = [position - cursor for position in boundaries[first:last]]
        cursor += len(text)
        if not cuts:
            continue
        parent = run._r.getparent()
        insertion = parent.index(run._r)
        units = []
        offset = 0
        for child in run._r:
            if child.tag == qn("w:rPr"):
                continue
            value = (child.text or "") if child.tag == qn("w:t") else ("\t" if child.tag == qn("w:tab") else "\n")
            child_template = copy.deepcopy(child)
            if child.tag == qn("w:t"):
                child_template.text = None
            units.append((offset, offset + len(value), child_template, value))
            offset += len(value)
        if offset != len(text):
            raise DocumentIntegrityError("Unsupported run content during annotation")
        # Clone a text-free template rather than copying a potentially long
        # original run once per finding. Attributes and run properties survive.
        template = copy.deepcopy(run._r)
        for child in list(template):
            if child.tag != qn("w:rPr"):
                template.remove(child)
        edges = [0, *cuts, len(text)]
        for part_index, (left, right) in enumerate(zip(edges, edges[1:])):
            new = copy.deepcopy(template)
            for start, end, child, value in units:
                if start >= right or end <= left:
                    continue
                piece = copy.deepcopy(child)
                if child.tag == qn("w:t"):
                    piece.text = value[max(0, left - start):min(end, right) - start]
                    if piece.text and (piece.text[0].isspace() or piece.text[-1].isspace()):
                        piece.set(qn("xml:space"), "preserve")
                new.append(piece)
            parent.insert(insertion + part_index, new)
        parent.remove(run._r)
    spans = []
    cursor = 0
    for run in paragraph.runs:
        end = cursor + len(run.text)
        if end > cursor:
            spans.append((cursor, end, run))
        cursor = end
    starts = [start for start, _, _ in spans]
    ends = [end for _, end, _ in spans]
    anchors = []
    for left, right in intervals:
        selected = [run for _, _, run in spans[bisect_right(ends, left):bisect_left(starts, right)]]
        if not selected or "".join(run.text for run in selected) != original_text[left:right]:
            raise DocumentIntegrityError("Selected witness changed during annotation")
        for run in selected:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        anchors.append(selected)
    return anchors

def _find_and_highlight_target(p, target_text, start=-1, end=-1):
    if not target_text or not _simple_runs(p):
        return None
    text = p.text
    if start >= 0:
        if end <= start or text[start:end] != target_text:
            return None
    else:
        if text.count(target_text) != 1:
            return None
        start = text.find(target_text)
        end = start + len(target_text)
    if (start < len(text) and unicodedata.category(text[start]).startswith("M")) or \
       (end < len(text) and unicodedata.category(text[end]).startswith("M")):
        return None
    return _split_runs_once(p, [(start, end)])[0]

def _payload(doc):
    # Text and every non-run-property XML leaf, in document order.
    ignored = {qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:commentReference")}
    result = []
    def visit(node):
        if node.tag == qn("w:rPr") or node.tag in ignored:
            return
        if node.tag != qn("w:r"):
            if len(node) == 0:
                result.append((node.tag, tuple(sorted(node.attrib.items())), node.text))
            elif node.tag not in {qn("w:t")}:
                result.append(("structure", node.tag, tuple(sorted(node.attrib.items()))))
        for child in node:
            visit(child)
    # Text runs may split; compare concatenated w:t separately.
    text = "".join(doc.element.xpath("//w:t/text()"))
    visit(doc.element)
    result = [r for r in result if r[0] != qn("w:t")]
    return text, result

def _comment_text(issue, scope="exact", repeats=0, fold_after=0):
    """Keep qualifiers intact; never cut a long sentence into a stronger claim."""
    from rule_registry import rule_spec
    labels = {"A": "اطمینان بالا", "B": "نیازمند بررسی", "C": "مبهم؛ فقط بررسی دستی"}
    reason = " ".join(str(issue.message or "").split())
    suggestion = " ".join(str(issue.suggestion or "").split())
    if not reason or len(reason) > 160:
        reason = rule_spec(issue.rule_id).title + "؛ دلیل کامل در گزارش آمده است."
    if len(suggestion) > 160:
        suggestion = "پیش از هر تغییر، پیشنهاد کامل گزارش را با متن اصلی مقایسه کنید."
    elif not suggestion:
        suggestion = "این مورد را با توجه به بافت متن بازبینی کنید."
    qualification = " [در سطح پاراگراف]" if scope == "paragraph" else ""
    text = (f"{issue.rule_id} | {issue.confidence} — {labels[issue.confidence]}{qualification}\n"
            f"دلیل: {reason}\nپیشنهاد: {suggestion}")
    if repeats:
        text += (f"\nتکرار: همین ایراد {repeats} بار در سند آمده و همهٔ موارد زرد شده‌اند؛ فقط {fold_after} مورد نخست کامنت دارند. "
                 "با Find & Replace (Ctrl+H) یک‌جا اصلاح کنید.")
    return text


# Word slows sharply with thousands of comment balloons; a thesis with 9000
# findings had only 58 distinct mistakes. Repeats of one mistake stay yellow
# where they occur, and only the first few carry their own comment.
FOLD_REPEATED_AFTER = 3


def _exact_target(paragraph, issue):
    """Resolve a unique literal witness without changing a single run."""
    target = issue.target_text
    left, right = issue.char_start, issue.char_end
    if not target:
        return None
    text = paragraph.text
    if left < 0:
        if left != -1 or right != -1 or text.count(target) != 1:
            return None
        left = text.find(target)
        right = left + len(target)
    if not 0 <= left < right <= len(text) or text[left:right] != target:
        return None
    if unicodedata.category(text[left]).startswith("M") or (
            right < len(text) and unicodedata.category(text[right]).startswith("M")):
        return None
    return left, right


def annotate_document(doc, issues, out_path, max_comments=None, per_par=None,
                      include_ambiguous=False, fold_after=FOLD_REPEATED_AFTER):
    """Annotate every safe witness on a copy, selecting A then B globally.

    The return value retains its historical meaning (at least one new comment).
    ``_last_annotation_summary.output_created`` describes successful publication
    separately, including the useful zero-comment review copy.
    """
    for name, value in (("max_comments", max_comments), ("per_par", per_par)):
        if value is None:
            continue
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise ValueError(f"{name} must be a nonnegative integer or None")
    if not isinstance(include_ambiguous, bool):
        raise ValueError("include_ambiguous must be a boolean")
    issues = list(issues)
    records = [dict(index=index, rule_id=issue.rule_id, confidence=issue.confidence,
                    block_id=issue.block_id, source_part=issue.source_part,
                    xml_path=issue.xml_path, story=issue.story, par_index=issue.par_index,
                    char_start=issue.char_start, char_end=issue.char_end,
                    status="unapplied", reason="location_unavailable", anchor_scope="")
               for index, issue in enumerate(issues)]
    summary = {"requested": len(issues), "applied": 0, "unapplied": len(issues),
               "max_comments": max_comments, "per_par": per_par,
               "include_ambiguous": include_ambiguous, "output_created": False,
               "exact_applied": 0, "paragraph_applied": 0, "highlighted": 0,
               "reasons": {}, "records": records}
    if doc is None or not hasattr(doc, "add_comment"):
        for record in records:
            record["reason"] = "annotation_unavailable"
        summary["reasons"] = dict(Counter(record["reason"] for record in records))
        if doc is not None:
            doc._last_annotation_summary = summary
        return False
    doc._last_annotation_summary = summary
    source = getattr(doc, "_review_source_path", None)
    initial_hash = getattr(doc, "_review_source_hash", None)
    if source:
        ensure_distinct_output(source, out_path)
        if sha256_file(source) != initial_hash:
            raise DocumentIntegrityError("Source changed since analysis")
        original_doc = Document(source)
        if doc.element.xml != original_doc.element.xml:
            raise DocumentIntegrityError("In-memory document changed since analysis")
        working = original_doc
    else:
        snapshot = io.BytesIO()
        doc.save(snapshot)
        snapshot.seek(0)
        working = Document(snapshot)
    before = _payload(working)
    grouped = defaultdict(list)
    paragraphs = {}
    selected = 0
    seen = defaultdict(set)
    # Resolve all original locations before creating any new comment markers.
    priority = {"A": 0, "B": 1, "C": 2}
    for index, issue in sorted(enumerate(issues), key=lambda pair: priority[pair[1].confidence]):
        record = records[index]
        if issue.confidence == "C" and not include_ambiguous:
            record["reason"] = "ambiguous_excluded"
            continue
        if issue.protected_overlap:
            record["reason"] = "protected_content"
            continue
        if issue.xml_path and issue.source_part.lstrip("/") == "word/document.xml" and issue.story in {"body", "table"}:
            key = ("xml", issue.xml_path)
            if key not in paragraphs:
                from lxml import etree
                from docx.text.paragraph import Paragraph
                try:
                    found = etree.XPath(issue.xml_path, namespaces={k: v for k, v in working.element.nsmap.items() if k})(working.element)
                except etree.XPathError:
                    record["reason"] = "location_unavailable"
                    continue
                if len(found) != 1 or found[0].tag != qn("w:p"):
                    record["reason"] = "location_unavailable"
                    continue
                paragraphs[key] = Paragraph(found[0], working)
        elif issue.par_index >= 0 and not issue.source_part and not issue.xml_path and issue.story in {"", "body"}:
            key = ("body", issue.par_index)
            if issue.par_index >= len(working.paragraphs):
                continue
            paragraphs[key] = working.paragraphs[issue.par_index]
        else:
            record["reason"] = "unsupported_story" if issue.source_part or issue.story else "location_unavailable"
            continue
        paragraph = paragraphs[key]
        if not _simple_runs(paragraph):
            record["reason"] = "unsafe_structure"
            continue
        target = _exact_target(paragraph, issue)
        scope = "exact"
        if not issue.target_text and issue.char_start == -1 and issue.char_end == -1 and paragraph.text:
            # Formatting/paragraph findings have a reliable paragraph location,
            # but no literal phrase. Clearly qualify their broader anchor.
            target = (0, len(paragraph.text))
            scope = "paragraph"
        if target is None:
            record["reason"] = "uncertain_location"
            continue
        left, right = target
        identity = (issue.rule_id, left, right, issue.target_text, issue.message,
                    issue.suggestion, issue.confidence)
        if identity in seen[key]:
            record["reason"] = "duplicate"
            continue
        seen[key].add(identity)
        if max_comments is not None and selected >= max_comments:
            record["reason"] = "comment_limit"
            continue
        if per_par is not None and len(grouped[key]) >= per_par:
            record["reason"] = "paragraph_limit"
            continue
        grouped[key].append((issue, index, left, right, scope))
        selected += 1
    # Identical exact mistakes (same rule, same text) beyond the first few are
    # highlighted only; the first comment states how often the mistake occurs.
    folded, totals = {}, {}
    if fold_after:
        repeats = defaultdict(list)
        for items in grouped.values():
            for issue, index, _, _, scope in items:
                if scope == "exact" and issue.target_text:
                    repeats[(issue.rule_id, issue.target_text)].append(index)
        for indexes in repeats.values():
            if len(indexes) > fold_after:
                indexes.sort()
                totals[indexes[0]] = len(indexes)
                folded.update((index, indexes[0]) for index in indexes[fold_after:])
    count = 0
    for key, items in grouped.items():
        paragraph = paragraphs[key]
        # Every selected witness, folded or not, is split out and highlighted.
        anchors = _split_runs_once(paragraph, [(left, right) for _, _, left, right, _ in items])
        for runs, (issue, index, left, right, scope) in zip(anchors, items):
            if index in folded:
                records[index].update(status="folded", reason="repeated_folded",
                                      anchor_scope=scope, folded_into=folded[index])
                continue
            working.add_comment(runs, text=_comment_text(issue, scope, totals.get(index, 0), fold_after),
                                author="Research Editor", initials="RE")
            records[index].update(status="applied", reason="annotated", anchor_scope=scope)
            count += 1
    summary.update(applied=count, folded=len(folded), unapplied=len(issues) - count - len(folded),
                   exact_applied=sum(r["status"] == "applied" and r["anchor_scope"] == "exact" for r in records),
                   paragraph_applied=sum(r["status"] == "applied" and r["anchor_scope"] == "paragraph" for r in records),
                   highlighted=count + len(folded),
                   reasons=dict(Counter(record["reason"] for record in records
                                        if record["status"] not in {"applied", "folded"})))
    if _payload(working) != before:
        raise DocumentIntegrityError("Annotation changed document payload")
    buffer = io.BytesIO()
    working.save(buffer)
    # Retain all original package bytes except main document and comment wiring.
    generated = zipfile.ZipFile(buffer)
    export = io.BytesIO()
    if source and count == 0:
        # With no comments, the independent review copy can retain every byte.
        with open(source, "rb") as original_stream:
            export.write(original_stream.read())
    elif source:
        # Deflate new parts too: an uncompressed comments part of thousands of
        # comments made reviewed files a hundred times larger than the source.
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(export, "w", zipfile.ZIP_DEFLATED) as output:
            names = set(original.namelist())
            replace = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
            # Existing valid packages may use a custom comments part name.
            for relationship in working.part.rels.values():
                if relationship.reltype.endswith("/comments") and not relationship.is_external:
                    replace.add(str(relationship.target_part.partname).lstrip("/"))
            for item in original.infolist():
                data = generated.read(item.filename) if item.filename in replace and item.filename in generated.namelist() else original.read(item.filename)
                if item.filename in {"word/_rels/document.xml.rels", "[Content_Types].xml"} and item.filename in generated.namelist():
                    from lxml import etree
                    old_tree = etree.fromstring(original.read(item.filename))
                    new_tree = etree.fromstring(data)
                    def entry_key(node):
                        return (node.tag, node.get("Id") or node.get("PartName") or node.get("Extension"))
                    known = {entry_key(node): node for node in old_tree}
                    for node in new_tree:
                        key = entry_key(node)
                        if key not in known:
                            old_tree.append(node)
                        elif dict(known[key].attrib) != dict(node.attrib):
                            raise DocumentIntegrityError("Conflicting package relationship")
                    data = etree.tostring(old_tree, xml_declaration=True, encoding="UTF-8", standalone=True)
                output.writestr(item, data)
            for name in generated.namelist():
                if name not in names:
                    output.writestr(name, generated.read(name), compress_type=zipfile.ZIP_DEFLATED)
    else:
        export.write(buffer.getvalue())
    generated.close()
    export.seek(0)
    reopened = Document(export)
    if _payload(reopened) != before:
        raise DocumentIntegrityError("Saved annotation failed content validation")
    if count and len(reopened.comments) != len(working.comments):
        raise DocumentIntegrityError("Saved annotation lost Word comments")
    destination = os.path.abspath(out_path)
    descriptor, temporary = tempfile.mkstemp(prefix=".review_", suffix=".docx", dir=os.path.dirname(destination))
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(export.getvalue())
        if source:
            ensure_distinct_output(source, destination)
            if sha256_file(source) != initial_hash:
                raise DocumentIntegrityError("Source changed during export")
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    summary["output_created"] = True
    return count > 0
