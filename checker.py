# -*- coding: utf-8 -*-
"""
موتور جامع ارزیابی و ویرایش علمی، نگارشی و ساختاری پایان‌نامه و مقاله (فایل Word، PDF و متن).
قاعده‌محور، دقیق و قابل گسترش؛ منطبق بر شیوه نامه‌های دانشگاهی و بین‌المللی.
شامل پشتیبانی از شیوه‌نامه دانشگاه علوم پزشکی تهران (TUMS).
"""
import argparse
import copy
import json
import os
import re
import time
import hashlib
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

import academic_checker
import en_checker
import pdf_handler
import tums_checker
from analysis_model import ConfidenceLevel, Issue, OffsetMapper
from doc_guard import DocumentGuard, safe_filename, sha256_file, validate_document, immutable_input
from protected_content import ProtectedContentDetector
import safe_annotator
from rule_registry import FindingBuffer, bounded_findings
from report_artifacts import (ANNOTATION_REASONS, COVERAGE_LABELS, STORY_LABELS,
                              ambiguous_count, annotation_format_label, category_label, location_label, safe_annotation_summary)

HERE = os.path.dirname(os.path.abspath(__file__))
ZWNJ = "\u200c"

# حروف فارسی/عربی
FA = "\u0621-\u063A\u0641-\u064A\u067E\u0686\u0698\u06A9\u06AF\u06CC\u06C0"
FA_RE = re.compile("[" + FA + "]")
LATIN_RE = re.compile("[A-Za-z]")
DIGIT_MAP = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

SEVERITY_ORDER = {"error": 0, "warn": 1, "info": 2}
SEVERITY_ICON = {"error": "🔴", "warn": "🟠", "info": "🔵"}

RULE_CATEGORIES = {
    # ۱. املایی و تایپوگرافی
    "arabic_letters": "۱. غلط املایی و تایپوگرافی",
    "arabic_digits": "۱. غلط املایی و تایپوگرافی",
    "latin_digits": "۱. غلط املایی و تایپوگرافی",
    "mi_prefix": "۱. غلط املایی و تایپوگرافی",
    "plural_space": "۱. غلط املایی و تایپوگرافی",
    "comparative_space": "۱. غلط املایی و تایپوگرافی",
    "ye_space": "۱. غلط املایی و تایپوگرافی",
    "common_mistake": "۱. غلط املایی و تایپوگرافی",
    "repeated_word": "۱. غلط املایی و تایپوگرافی",

    # ۲. دستور زبان و نگارش (فارسی و انگلیسی)
    "en_grammar_agreement": "۲. دستور زبان و ساختار جمله",
    "en_article": "۲. دستور زبان و ساختار جمله",
    "en_preposition": "۲. دستور زبان و ساختار جمله",

    # ۳. نشانه‌گذاری و فاصله‌گذاری
    "space_before_punct": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "no_space_after_punct": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "latin_punct": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "paren_space": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "straight_quotes": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "double_space": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "leading_space": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "blank_lines": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "en_punctuation": "۳. نشانه‌گذاری و فاصله‌گذاری",
    "en_unpaired_brackets": "۳. نشانه‌گذاری و فاصله‌گذاری",

    # ۴. سبک و لحن علمی
    "informal_word": "۴. سبک و لحن علمی",
    "first_person": "۴. سبک و لحن علمی",
    "academic_tone_emotional": "۴. سبک و لحن علمی",
    "academic_pleonasm": "۴. سبک و لحن علمی",
    "academic_first_person": "۴. سبک و لحن علمی",
    "en_academic_tone": "۴. سبک و لحن علمی",
    "en_wordiness": "۴. سبک و لحن علمی",

    # ۵. انسجام و ساختار فصول
    "long_paragraph": "۵. انسجام و ساختار فصول",
    "long_sentence": "۵. انسجام و ساختار فصول",
    "academic_cohesion": "۵. انسجام و ساختار فصول",
    "missing_section": "۵. انسجام و ساختار فصول",
    "recommended_section": "۵. انسجام و ساختار فصول",
    "chapters": "۵. انسجام و ساختار فصول",
    "heading_style": "۵. انسجام و ساختار فصول",
    "toc": "۵. انسجام و ساختار فصول",
    "abstract_length": "۵. انسجام و ساختار فصول",
    "keywords_count": "۵. انسجام و ساختار فصول",
    "tums_chapter_proportions": "۵. انسجام و ساختار فصول",

    # ۶. قالب‌بندی فایل و شیوه دانشگاه
    "page_size": "۶. قالب‌بندی و شیوه دانشگاه",
    "page_margin": "۶. قالب‌بندی و شیوه دانشگاه",
    "font_fa": "۶. قالب‌بندی و شیوه دانشگاه",
    "size_fa": "۶. قالب‌بندی و شیوه دانشگاه",
    "font_en": "۶. قالب‌بندی و شیوه دانشگاه",
    "size_en": "۶. قالب‌بندی و شیوه دانشگاه",
    "line_spacing": "۶. قالب‌بندی و شیوه دانشگاه",
    "alignment": "۶. قالب‌بندی و شیوه دانشگاه",
    "rtl": "۶. قالب‌بندی و شیوه دانشگاه",
    "page_numbers": "۶. قالب‌بندی و شیوه دانشگاه",
    "tums_page_margin": "۶. قالب‌بندی و شیوه دانشگاه",

    # ۷. جدول‌ها، شکل‌ها و فرمول‌ها
    "caption_missing": "۷. جدول‌ها، نمودارها و شکل‌ها",
    "caption_numbering": "۷. جدول‌ها، نمودارها و شکل‌ها",

    # ۸. ارجاع‌دهی و فهرست منابع
    "cite_missing_ref": "۸. ارجاع‌دهی و فهرست منابع",
    "ref_not_cited": "۸. ارجاع‌دهی و فهرست منابع",
    "numeric_cite": "۸. ارجاع‌دهی و فهرست منابع",
    "tums_et_al_forbidden": "۸. ارجاع‌دهی و فهرست منابع",
    "tums_ref_page_format": "۸. ارجاع‌دهی و فهرست منابع",

    # ۹. چک‌لیست دانشگاه علوم پزشکی تهران
    "tums_checklist_missing": "۹. چک‌لیست اختصاصی دانشگاه علوم پزشکی تهران",
}

RULE_TITLES = {
    "arabic_letters": "حرف عربی (ي / ك) به‌جای ی / ک فارسی",
    "arabic_digits": "ارقام عربی به‌جای ارقام فارسی",
    "latin_digits": "رقم لاتین در متن فارسی",
    "mi_prefix": "پیشوند «می / نمی» بدون نیم‌فاصله",
    "plural_space": "«ها / های» جدا از واژه (نیم‌فاصله لازم است)",
    "comparative_space": "«تر / ترین» جدا از واژه",
    "ye_space": "«ی» اضافه‌ی جدا از واژه",
    "space_before_punct": "فاصله قبل از علائم نگارشی (، ؛ : . ؟ !)",
    "no_space_after_punct": "نبودن فاصله بعد از علائم نگارشی",
    "latin_punct": "علائم لاتین (, ; ?) در متن فارسی",
    "paren_space": "فاصله‌گذاری نادرست پرانتز",
    "straight_quotes": "گیومه‌ی لاتین \" \" به‌جای « »",
    "double_space": "فاصله‌ی اضافی بین کلمات",
    "leading_space": "تورفتگی با Space/Tab به‌جای تنظیم پاراگراف",
    "repeated_word": "تکرار پشت‌سرهم یک واژه",
    "common_mistake": "غلط املایی/نگارشی رایج",
    "informal_word": "واژه‌ی محاوره‌ای در متن علمی",
    "first_person": "ضمیر اول‌شخص در نگارش علمی",
    "blank_lines": "پاراگراف خالی پشت‌سرهم (فاصله با Enter)",
    "long_paragraph": "پاراگراف بسیار طولانی",
    "long_sentence": "جمله‌ی بسیار طولانی",
    "page_size": "اندازه‌ی صفحه",
    "page_margin": "حاشیه‌های صفحه",
    "font_fa": "فونت متن فارسی",
    "size_fa": "اندازه‌ی قلم متن فارسی",
    "font_en": "فونت متن انگلیسی",
    "size_en": "اندازه‌ی قلم متن انگلیسی",
    "line_spacing": "فاصله‌ی خطوط",
    "alignment": "تراز پاراگراف (دوطرفه)",
    "rtl": "جهت پاراگراف (راست‌به‌چپ) تنظیم نشده",
    "missing_section": "بخش الزامی یافت نشد",
    "recommended_section": "بخش پیشنهادی یافت نشد",
    "chapters": "فصل‌ها",
    "heading_style": "عنوان فصل با استایل Heading نیست",
    "toc": "فهرست مطالب",
    "page_numbers": "شماره‌ی صفحه",
    "abstract_length": "طول چکیده",
    "keywords_count": "تعداد کلیدواژه‌ها",
    "caption_missing": "جدول/شکل بدون عنوان (کپشن)",
    "caption_numbering": "شماره‌گذاری جدول/شکل",
    "cite_missing_ref": "ارجاع درون‌متنی بدون مورد متناظر در فهرست منابع",
    "ref_not_cited": "منبع فهرست‌شده که در متن ارجاع نشده",
    "numeric_cite": "ارجاع عددی خارج از محدوده",
    # موارد افزوده شده جدید:
    "academic_tone_emotional": "لحن احساسی یا مبالغه‌آمیز در نگارش علمی",
    "academic_pleonasm": "حشو و زوائد بیانی (ساده‌نویسی علمی)",
    "academic_first_person": "کاربرد ضمیر اول‌شخص در نگارش دانشگاهی",
    "academic_cohesion": "انسجام و پیوستگی استدلال (نبود نشانگر گذار)",
    "en_punctuation": "نشانه‌گذاری نادرست در متن انگلیسی",
    "en_article": "کاربرد نادرست حروف تعریف (a / an) در انگلیسی",
    "en_academic_tone": "لحن و انتخاب واژگان آکادمیک در انگلیسی",
    "en_wordiness": "پرگویی و حشو در متن انگلیسی (Wordiness)",
    "en_preposition": "حرف اضافه نادرست در متن انگلیسی",
    "en_grammar_agreement": "تطابق دستوری فاعل/فعل یا شمارش اسامی انگلیسی",
    "en_unpaired_brackets": "عدم تقارن پرانتز یا گیومه در انگلیسی",
    "tums_checklist_missing": "بخش الزامی از چک‌لیست دانشگاه علوم پزشکی تهران",
    "tums_chapter_proportions": "توازن حجم فصول (سقف ۳۰٪ برای فصول ۱ و ۲)",
    "tums_et_al_forbidden": "ممنوعیت «همکاران» و «et al» در فهرست منابع TUMS",
    "tums_ref_page_format": "فرمت شماره صفحات در مراجع TUMS",
    "tums_page_margin": "انطباق حاشیه‌ها با ضوابط دانشگاه علوم پزشکی تهران",
}


# Issue is imported from analysis_model


# ----------------------------------------------------------------------------
# ابزارهای کمکی
# ----------------------------------------------------------------------------
def load_rules(path=None, mode="thesis"):
    """بارگذاری قوانین پایه و ادغام با پروفایل اختصاصی در صورت انتخاب TUMS."""
    base_path = path or os.path.join(HERE, "rules.json")
    with open(base_path, encoding="utf-8") as f:
        rules = json.load(f)
    rules.setdefault("strict", True)
    rules.setdefault("protected_terms", [])
    rules.setdefault("disabled_rules", [])

    if mode == "tums":
        tums_path = os.path.join(HERE, "profiles", "tums.json")
        if os.path.exists(tums_path):
            with open(tums_path, encoding="utf-8") as tf:
                tums_data = json.load(tf)
                # ادغام مقادیر
                rules["page"] = tums_data.get("page", rules["page"])
                rules["body"] = tums_data.get("body", rules["body"])
                rules["limits"] = tums_data.get("limits", rules["limits"])
                rules["tums_special_rules"] = tums_data.get("tums_special_rules", {})
                rules["checklist_items"] = tums_data.get("checklist_items", [])
                if "tums" not in rules["modes"]:
                    rules["modes"]["tums"] = {
                        "title": "پایان‌نامه دانشگاه علوم پزشکی تهران (TUMS)",
                        "required_sections": rules["modes"]["thesis"]["required_sections"],
                        "recommended_sections": rules["modes"]["thesis"].get("recommended_sections", [])
                    }
    elif mode == "english":
        if "english" not in rules["modes"]:
            rules["modes"]["english"] = {
                "title": "متن / مقاله انگلیسی",
                "required_sections": [],
                "recommended_sections": []
            }
    from rule_registry import validate_rules
    return validate_rules(rules, mode)


def _prepare_rules(rules, mode):
    from rule_registry import validate_rules
    return load_rules(mode=mode) if rules is None else validate_rules(copy.deepcopy(rules), mode)


def _finish_analysis(issues, stats, started, rules=None):
    from rule_registry import enrich_issue
    # A parser failure is an operational result, not an optional editing rule.
    issues = [issue for issue in issues if issue.rule_id == "pdf_read_error"
              or issue.rule_id not in (rules or {}).get("disabled_rules", [])]
    limit = (rules or {}).get("max_findings", 10000)
    if len(issues) > limit:
        issues = issues[:limit] + [Issue("findings_limit", "warn", "تعداد یافته‌های سند محدود شده است؛ بخش‌های باقیمانده نیازمند بازبینی‌اند.",
                                       confidence="B", extraction_status="analysis_limited")]
    stats["analysis_duration_seconds"] = round(time.monotonic() - started, 3)
    for issue in issues:
        enrich_issue(issue)
    if stats.get("analysis_status") != "failed":
        limited = (stats.get("skipped_blocks", 0) or stats.get("partial_blocks", 0)
                   or any(issue.extraction_status in {"analysis_limited", "analysis_skipped", "no_reliable_text"}
                          for issue in issues)
                   or any(info.get("status") in {"partial", "present"}
                          for info in stats.get("coverage", {}).values()))
        stats["analysis_status"] = "partial" if limited else stats.get("analysis_status", "complete")
    return deduplicate_issues(issues)


def _document_coverage(blocks, analyzed, block_limit):
    coverage = {}
    for story in ("body", "table", "header", "footer", "footnote", "endnote", "textbox"):
        selected = [block for block in blocks if block.story == story]
        complete = sum((block.block_id in analyzed and len(block.text) <= block_limit)
                       or (not block.text.strip() and block.extraction_status == "native") for block in selected)
        partial = sum(block.extraction_status != "native" or len(block.text) > block_limit for block in selected)
        status = ("absent" if not selected else "partial" if partial or 0 < complete < len(selected)
                  else "analyzed" if complete == len(selected) else "present")
        coverage[story] = {"status": status, "present_blocks": len(selected),
                           "analyzed_blocks": complete, "partial_blocks": partial}
    return coverage


def _attach_block_issues(issues, block):
    for issue in issues:
        issue.block_id = block.block_id
        issue.source_part = block.source_part
        issue.xml_path = block.xml_path
        issue.story = block.story
        issue.paragraph_index = block.paragraph_index
        issue.extraction_status = block.extraction_status
        issue.loc = issue.loc or block.location
        if issue.target_text:
            issue.validate_evidence(block.text)
        if block.extraction_status != "native" and issue.confidence == "A":
            issue.confidence = "B"
    return issues


def _protected_audit(blocks, block_limit=200000):
    digest = hashlib.sha256()
    counts = Counter()
    skipped = 0
    for block in blocks:
        text = block if isinstance(block, str) else block.text
        if len(text) > block_limit:
            skipped += 1
            continue
        for span in ProtectedContentDetector.find_spans(text):
            counts[span.kind] += 1
            encoded = span.text.encode("utf-8")
            digest.update(len(encoded).to_bytes(8, "big"))
            digest.update(encoded)
    return {"count": sum(counts.values()), "by_kind": dict(counts),
            "ordered_sha256": digest.hexdigest(), "skipped_blocks": skipped}


def norm(s):
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("ي", "ی").replace("ك", "ک").replace(ZWNJ, " ")
    return re.sub(r"\s+", " ", s).strip().lower()


def to_latin_digits(s):
    return s.translate(DIGIT_MAP)


def words(s):
    return re.findall(r"[\w\u200c]+", s)


def has_fa(s):
    if not s:
        return False
    s = unicodedata.normalize("NFKC", s)
    return bool(FA_RE.search(s))


def snippet_of(text, m, pad=18):
    a, b = max(0, m.start() - pad), min(len(text), m.end() + pad)
    out = text[a:b].replace("\n", " ")
    return ("…" if a else "") + out + ("…" if b < len(text) else "")


def is_heading(p):
    n = (p.style.name or "").lower() if p.style is not None else ""
    return n.startswith("heading") or n.startswith("عنوان") or n == "title"


def is_toc_style(p):
    n = (p.style.name or "").lower() if p.style is not None else ""
    return n.startswith("toc") or n.startswith("table of") or "فهرست" in n


def _attr(parent, tag, attr):
    if parent is None:
        return None
    el = parent.find(qn(tag))
    return None if el is None else el.get(qn(attr))


def _defaults(doc, kind):
    d = doc.styles.element.find(qn("w:docDefaults"))
    if d is None:
        return None
    holder = d.find(qn("w:rPrDefault" if kind == "r" else "w:pPrDefault"))
    if holder is None:
        return None
    return holder.find(qn("w:rPr" if kind == "r" else "w:pPr"))


def _levels(doc, par, kind, run=None):
    if kind == "r":
        if run is not None:
            yield run._r.rPr
        st = par.style
        while st is not None:
            yield st.element.rPr
            st = st.base_style
    else:
        yield par._p.pPr
        st = par.style
        while st is not None:
            yield st.element.pPr
            st = st.base_style
    yield _defaults(doc, kind)


def eff_rpr(doc, par, run, tag, attr):
    for lvl in _levels(doc, par, "r", run):
        v = _attr(lvl, tag, attr)
        if v is not None:
            return v
    return None


def eff_ppr(doc, par, tag, attr):
    for lvl in _levels(doc, par, "p"):
        v = _attr(lvl, tag, attr)
        if v is not None:
            return v
    return None


def eff_bidi(doc, par):
    for lvl in _levels(doc, par, "p"):
        if lvl is not None and lvl.find(qn("w:bidi")) is not None:
            el = lvl.find(qn("w:bidi"))
            return el.get(qn("w:val")) not in ("0", "false")
    return False


def font_ok(name, allowed):
    if name is None:
        return True
    n = name.lower().replace(" ", "")
    return any(n == a.lower().replace(" ", "") for a in allowed)


# ----------------------------------------------------------------------------
# قوانین متنی (تایپوگرافی، سبک علمی و انگلیسی)
# ----------------------------------------------------------------------------
MI_STEMS = ("کن|کرد|شو|شد|توان|باش|گوی|گفت|ده|داد|خواه|خواست|گیر|گرفت|بین|دید|رس|رو|"
            "نویس|پرداز|یاب|یافت|آید|آورد|ماند|ساز|شناس|پذیر")

SCIENTIFIC_UNITS_TAIL_RE = re.compile(
    r"^\s*(?:°\s*[cCfFkK]|(?:µ|u|m|c|k|d|n|p)?[gmldL]|(?:µ|u|m|k)?[vVaAwWjJ]|"
    r"sec|min|hr|hours?|days?|years?|hz|kHz|MHz|GHz|pa|kPa|MPa|GPa|mol|mmol|µmol|nmol|pmol|"
    r"rpm|bp|kb|mb|gb|tb|kda|kDa|iu|IU|cfu|CFU|ph|pH|%|‰|[xX]|fold)\b",
    re.IGNORECASE
)
SCIENTIFIC_HEAD_RE = re.compile(
    r"(?:(?:p|r|t|f|z)\s*[<>=~]|doi:|url:|https?://|www\.|n\s*=|ci\s*|ph\s*|covid-?)\s*$",
    re.IGNORECASE
)
ABBR_EXEMPT_RE = re.compile(r"(?:هـ\.[شق]|ه\.[شق])")

REDUPLICATION_WORDS = {
    "کم", "ذره", "پله", "تک", "رفته", "قطره", "دسته", "بخش", "صفحه",
    "سطر", "نکته", "لحظه", "ثانیه", "قدم", "سال", "روز", "ماه",
    "یواش", "آرام", "دور", "نرم", "تکه", "بار", "برگ", "دانه", "گروه",
    "مرحله", "سطح", "جزء", "تارتار", "یک", "دو"
}
REPEATED_WORD_RE = re.compile(r"(?<![" + FA + r"])([" + FA + r"]{2,})[ \u00a0]+\1(?![" + FA + "])")

TEXT_RULES = [
    ("arabic_letters", "error", re.compile("[يك]"),
     "حرف عربی به‌کار رفته؛ «ی» و «ک» فارسی تایپ شود (کیبورد فارسی استاندارد)."),
    ("arabic_digits", "error", re.compile("[٠-٩]"),
     "رقم عربی (٠١٢) به‌کار رفته؛ از ارقام فارسی (۰۱۲) استفاده کنید."),
    ("latin_digits", "warn",
     re.compile("[" + FA + r"]\s?[0-9]{1,4}(?![A-Za-z0-9])|(?<![A-Za-z0-9.,/-])[0-9]{1,4}\s?[" + FA + "]"),
     "رقم لاتین کنار متن فارسی؛ در متن فارسی ارقام فارسی بنویسید."),
    ("mi_prefix", "error",
     re.compile(r"(?<![" + FA + ZWNJ + r"])(ن?می)(?:" + MI_STEMS + r")"),
     "پیشوند «می/نمی» باید با نیم‌فاصله به فعل بچسبد: «می‌کند»."),
    ("mi_prefix", "error",
     re.compile(r"(?<![" + FA + ZWNJ + r"])ن?می[ \u00a0]+(?=[" + FA + r"]{2})"),
     "«می/نمی» با فاصله از فعل جدا شده؛ باید نیم‌فاصله باشد: «می‌کند»."),
    ("plural_space", "error",
     re.compile("[" + FA + r"]{2,}[ \u00a0]+(?:ها|های|هایی|هایش|هایم|هایمان|هایتان|هایشان)(?![" + FA + "])"),
     "«ها/های» با فاصله‌ی معمولی جدا شده؛ نیم‌فاصله بگذارید: «کتاب‌ها»."),
    ("comparative_space", "warn",
     re.compile("[" + FA + r"]{2,}[ \u00a0]+(?:تر|ترین)(?![" + FA + "])"),
     "«تر/ترین» باید با نیم‌فاصله به واژه بچسبد: «بزرگ‌تر»."),
    ("ye_space", "warn",
     re.compile("[" + FA + r"][ \u00a0]+ی(?![" + FA + "])"),
     "«ی» اضافه جدا از واژه است؛ «خانه‌ی» یا «خانهٔ»."),
    ("ye_space", "warn",
     re.compile("[" + FA + r"][ \u00a0]+ای(?![" + FA + "])"),
     "«ای» نکره جدا از واژه است؛ «نامه‌ای»."),
    ("space_before_punct", "warn",
     re.compile("(?<=[" + FA + r"\)])[ \u00a0]+[،؛؟!:\.](?!\.)"),
     "قبل از علائم نگارشی فاصله نیاید و بعد از آن‌ها یک فاصله بیاید."),
    ("no_space_after_punct", "warn",
     re.compile("[،؛؟](?=[" + FA + r"\w])|(?<=[" + FA + r"])\.(?=[" + FA + r"])|(?<=[" + FA + r"]):(?=[" + FA + "])"),
     "بعد از علامت نگارشی یک فاصله لازم است."),
    ("latin_punct", "warn",
     re.compile("(?<=[" + FA + r"])[,;?]"),
     "از علائم فارسی استفاده کنید: «،» «؛» «؟»."),
    ("paren_space", "warn",
     re.compile(r"\([ \u00a0]+|[ \u00a0]+\)|(?<=[" + FA + r"])\((?![A-Za-z0-9۰-۹])|\)(?=[" + FA + "])"),
     "داخل پرانتز فاصله نیاید؛ قبل از «(» و بعد از «)» یک فاصله بگذارید."),
    ("straight_quotes", "warn", re.compile(r'"[^"\n]{1,200}"'),
     "از گیومه‌ی فارسی « » استفاده کنید."),
    ("double_space", "warn", re.compile(r"(?<=\S)[ \u00a0]{2,}(?=\S)"),
     "بیش از یک فاصله بین کلمات."),
    ("leading_space", "warn", re.compile(r"^[ \t\u00a0]+(?=\S)"),
     "ابتدای پاراگراف با Space/Tab تورفتگی داده شده؛ از تنظیمات پاراگراف (First line) استفاده کنید."),
]


def _word_re(w):
    return re.compile("(?<![" + FA + ZWNJ + "])" + re.escape(w) + "(?![" + FA + ZWNJ + "])")


def check_text(par_index, text, loc, rules, in_refs=False, in_declaration=False, *, audit=None):
    from protected_content import ProtectedSpan
    if len(text) > rules.get("max_analysis_block_chars", 200000):
        return [Issue("analysis_limit", "warn", "Block exceeds the analysis budget; manual review required.",
                      loc=loc, par_index=par_index, confidence="B", extraction_status="analysis_skipped")]
    protected = ProtectedContentDetector.find_spans(text)
    if in_refs:
        # Bibliographic names and titles are quoted source records, not author prose.
        protected.append(ProtectedSpan(0, len(text), "bibliography_entry", text))
    for term in rules.get("protected_terms", []):
        if term:
            for m in re.finditer(re.escape(term), text):
                protected.append(ProtectedSpan(m.start(), m.end(), "configured_term", m.group()))
    issues = _check_text(par_index, text, loc, rules, in_refs, in_declaration)
    result, suppressed = [], []
    for issue in issues:
        if issue.rule_id in rules.get("disabled_rules", []):
            continue
        if issue.target_text:
            exact = (0 <= issue.char_start < issue.char_end <= len(text)
                     and text[issue.char_start:issue.char_end] == issue.target_text)
            issue.location_status = "exact" if exact else "unverified"
            if not exact:
                issue.confidence = "C"
                issue.char_start = issue.char_end = -1
            elif ProtectedContentDetector.is_protected(issue.char_start, issue.char_end, protected):
                continue
            else:
                issue.snippet = text[max(0, issue.char_start-18):min(len(text), issue.char_end+18)]
        elif issue.confidence == "A":
            issue.confidence = "B"
        if issue.rule_id in {"comparative_space", "ye_space", "latin_punct", "repeated_word", "EN_SVA"}:
            issue.confidence = "B"
        if issue.rule_id == "mi_prefix":
            following = text[issue.char_start:]
            if not re.match(r"ن?می[ \u00a0]*(?:" + MI_STEMS + r")(?:م|ی|د|یم|ید|ند|ت|تم|تی|تیم|تید|تند|ه)?(?![" + FA + r"])", following):
                issue.confidence = "C"
        if issue.confidence != "C" or not rules.get("strict", True):
            result.append(issue)
        else:
            suppressed.append(issue)
    if audit is not None:
        audit["ambiguous_candidates"] = audit.get("ambiguous_candidates", 0) + len(deduplicate_issues(suppressed))
    return deduplicate_issues(result)


@bounded_findings
def _check_text(par_index, text, loc, rules, in_refs=False, in_declaration=False):
    out = FindingBuffer(rules.get("max_findings_per_block", 500))
    if not text.strip():
        return out

    # استخراج محدوده‌های محافظت‌شده علمی (فرمول‌ها، یکاها، مواد شیمیایی، DOI، آمار و ارجاعات)
    prot_spans = ProtectedContentDetector.find_spans(text)

    # تایپوگرافی عمومی
    for rule, sev, rx, msg in TEXT_RULES:
        if rule in ("arabic_digits", "arabic_letters") or has_fa(text):
            for m in rx.finditer(text):
                # اگر با محدوده علمی محافظت‌شده تداخل داشته باشد، قاعده اعمال نشود
                if ProtectedContentDetector.is_protected(m.start(), m.end(), prot_spans):
                    continue

                # فیلتر استثناهای علمی ارقام لاتین (واحدها، درصدها، شاخص‌های آماری)
                if rule == "latin_digits":
                    after_pos = m.end()
                    tail = text[after_pos:after_pos + 12]
                    head = text[max(0, m.start() - 10):m.start()]
                    if SCIENTIFIC_UNITS_TAIL_RE.match(tail) or SCIENTIFIC_HEAD_RE.search(head):
                        continue
                # فیلتر اختصارات علمی و تاریخی (مانند هـ.ش یا ص.)
                if rule == "no_space_after_punct":
                    snippet = text[max(0, m.start() - 3):min(len(text), m.end() + 3)]
                    if ABBR_EXEMPT_RE.search(snippet):
                        continue

                conf = ConfidenceLevel.HIGH_A.value
                if rule in ("straight_quotes", "leading_space"):
                    conf = ConfidenceLevel.REVIEW_B.value

                cat = RULE_CATEGORIES.get(rule, "۱. غلط املایی و تایپوگرافی")
                out.append(Issue(
                    rule_id=rule,
                    severity=sev,
                    message=msg,
                    loc=loc,
                    snippet=snippet_of(text, m),
                    par_index=par_index,
                    suggestion="",
                    target_text=m.group(0),
                    category=cat,
                    confidence=conf,
                    char_start=m.start(),
                    char_end=m.end()
                ))

    # بررسی تکرار واژه با تفکیک واژگان تکرارشدنی مجاز (مانند کم‌کم یا ذره‌ذره)
    if has_fa(text):
        for m in REPEATED_WORD_RE.finditer(text):
            if ProtectedContentDetector.is_protected(m.start(), m.end(), prot_spans):
                continue
            w = m.group(1)
            cat = RULE_CATEGORIES.get("repeated_word", "۱. غلط املایی و تایپوگرافی")
            if w in REDUPLICATION_WORDS:
                out.append(Issue(
                    rule_id="repeated_word",
                    severity="info",
                    message=f"تکرار واژه «{w} {w}»؛ برای ساختار ترکیبی بهتر است با نیم‌فاصله پیوند داده شود: «{w}‌{w}».",
                    loc=loc,
                    snippet=snippet_of(text, m),
                    par_index=par_index,
                    suggestion=f"{w}‌{w}",
                    target_text=m.group(0),
                    category=cat,
                    confidence=ConfidenceLevel.REVIEW_B.value,
                    char_start=m.start(),
                    char_end=m.end()
                ))
            else:
                out.append(Issue(
                    rule_id="repeated_word",
                    severity="warn",
                    message=f"تکرار ناخواسته واژه «{w}» پشت‌سرهم؛ واژه دوم اضافی است.",
                    loc=loc,
                    snippet=snippet_of(text, m),
                    par_index=par_index,
                    suggestion=w,
                    target_text=m.group(0),
                    category=cat,
                    confidence=ConfidenceLevel.HIGH_A.value,
                    char_start=m.start(),
                    char_end=m.end()
                ))

    # بررسی انگلیسی در صورت وجود متن انگلیسی در پاراگراف
    en_words = [w for w in text.split() if LATIN_RE.match(w)]
    if len(en_words) >= 3:
        en_issues = en_checker.check_english_text(text, loc, par_index=par_index)
        for ei in en_issues:
            out.append(ei)

    if not has_fa(text):
        return out

    # نگاشتار مطمئن تحلیلی برای غلط‌های املایی رایج
    mapper = OffsetMapper.create_normalized(text)
    t = mapper.normalized_text

    for wrong, right in rules.get("common_mistakes", {}).items():
        rx = _word_re(wrong)
        for m in rx.finditer(t):
            orig_s, orig_e, ok = mapper.map_slice_to_original(m.start(), m.end())
            if ok and ProtectedContentDetector.is_protected(orig_s, orig_e, prot_spans):
                continue
            matched_orig = text[orig_s:orig_e] if ok else m.group(0)
            out.append(Issue(
                rule_id="common_mistake",
                severity="error",
                message=f"«{wrong}» ← پیشنهاد اصلاح: «{right}»",
                loc=loc,
                snippet=snippet_of(text, m),
                par_index=par_index,
                suggestion=right,
                target_text=matched_orig,
                category=RULE_CATEGORIES.get("common_mistake", "۱. غلط املایی و تایپوگرافی"),
                confidence=ConfidenceLevel.HIGH_A.value,
                char_start=orig_s if ok else -1,
                char_end=orig_e if ok else -1
            ))

    if in_refs:
        return out

    # واژگان محاوره‌ای
    for w in rules.get("informal_words", []):
        for m in _word_re(w).finditer(t):
            orig_s, orig_e, ok = mapper.map_slice_to_original(m.start(), m.end())
            if ok and ProtectedContentDetector.is_protected(orig_s, orig_e, prot_spans):
                continue
            matched_orig = text[orig_s:orig_e] if ok else m.group(0)
            out.append(Issue(
                rule_id="informal_word",
                severity="info",
                message=f"«{w}» محاوره‌ای است؛ معادل رسمی و آکادمیک بنویسید.",
                loc=loc,
                snippet=snippet_of(text, m),
                par_index=par_index,
                suggestion="معادل رسمی آکادمیک",
                target_text=matched_orig,
                category=RULE_CATEGORIES.get("informal_word", "۴. سبک و لحن علمی"),
                confidence=ConfidenceLevel.REVIEW_B.value,
                char_start=orig_s if ok else -1,
                char_end=orig_e if ok else -1
            ))

    # بررسی سبک علمی و مبالغه (ماژول تخصصی academic_checker)
    acad_issues = academic_checker.check_academic_style(text, loc, in_declaration, par_index=par_index)
    for ai in acad_issues:
        out.append(ai)

    lim = rules.get("limits", {})
    n_words = len(words(text))
    if n_words > lim.get("max_paragraph_words", 300):
        out.append(Issue(
            rule_id="long_paragraph",
            severity="info",
            message=f"پاراگراف {n_words} کلمه دارد؛ بهتر است برای خوانایی و انتقال بهتر پیام شکسته شود.",
            loc=loc,
            snippet=text[:60] + "…",
            par_index=par_index,
            category=RULE_CATEGORIES.get("long_paragraph", "۵. انسجام و ساختار فصول"),
            confidence=ConfidenceLevel.REVIEW_B.value
        ))

    for sent in re.split(r"[\.؟!؛\?]\s", text):
        if len(words(sent)) > lim.get("max_sentence_words", 60):
            out.append(Issue(
                rule_id="long_sentence",
                severity="info",
                message=f"جمله‌ای با {len(words(sent))} کلمه؛ برای درک ساده‌تر کوتاه‌تر شود.",
                loc=loc,
                snippet=sent[:60] + "…",
                par_index=par_index,
                category=RULE_CATEGORIES.get("long_sentence", "۵. انسجام و ساختار فصول"),
                confidence=ConfidenceLevel.REVIEW_B.value
            ))
    return out


# ----------------------------------------------------------------------------
# بررسی صفحه‌آرایی و استایل‌ها
# ----------------------------------------------------------------------------
@bounded_findings
def check_page(doc, rules, mode="thesis"):
    out = FindingBuffer()
    cfg = rules["page"]
    tol, mtol = cfg["size_tolerance_cm"], cfg["margin_tolerance_cm"]
    for i, s in enumerate(doc.sections, 1):
        if s.page_width is None or s.page_height is None:
            continue
        w, h = s.page_width / 360000, s.page_height / 360000
        if abs(w - cfg["width_cm"]) > tol or abs(h - cfg["height_cm"]) > tol:
            out.append(Issue("page_size", "error",
                             f"بخش {i}: اندازه‌ی صفحه {w:.1f}×{h:.1f} سانتی‌متر است؛ باید A4 "
                             f"({cfg['width_cm']}×{cfg['height_cm']}) باشد.", f"بخش {i}"))
        for side, attr in (("top", "top_margin"), ("bottom", "bottom_margin"),
                           ("right", "right_margin"), ("left", "left_margin")):
            v = getattr(s, attr)
            if mode == "tums":
                continue
            if v is None:
                continue
            cm, want = v / 360000, cfg["margins_cm"][side]
            if abs(cm - want) > mtol:
                nm = {"top": "بالا", "bottom": "پایین", "right": "راست", "left": "چپ"}[side]
                out.append(Issue("page_margin", "warn",
                                 f"بخش {i}: حاشیه‌ی {nm} {cm:.2f}cm است؛ باید {want}cm باشد.", f"بخش {i}"))

    # بررسی حاشیه‌های تخصصی TUMS
    if mode == "tums":
        tums_issues = tums_checker.check_tums_margins(doc.sections)
        for ti in tums_issues:
            out.append(ti)

    return out


@bounded_findings
def check_body_format(doc, rules, ctx):
    out = FindingBuffer()
    b = rules["body"]
    for i, p, heading in ctx:
        text = p.text
        if len(text.strip()) < b.get("min_chars_to_count_as_body", 40):
            continue
        if is_heading(p) or is_toc_style(p) or (p.style is not None and "caption" in (p.style.name or "").lower()):
            continue
        loc = f"زیر «{heading}»، پاراگراف {i + 1}"
        bad = defaultdict(set)
        for r in p.runs:
            if not r.text.strip():
                continue
            fa, en = has_fa(r.text), bool(LATIN_RE.search(r.text))
            if fa:
                name = eff_rpr(doc, p, r, "w:rFonts", "w:cs")
                sz = eff_rpr(doc, p, r, "w:szCs", "w:val")
                if not font_ok(name, b["fa_fonts"]):
                    bad["font_fa"].add(name)
                if sz and abs(int(sz) / 2 - b["fa_size_pt"]) > 0.1:
                    bad["size_fa"].add(f"{int(sz) / 2:g}")
            elif en:
                name = eff_rpr(doc, p, r, "w:rFonts", "w:ascii")
                sz = eff_rpr(doc, p, r, "w:sz", "w:val")
                if not font_ok(name, b["en_fonts"]):
                    bad["font_en"].add(name)
                if sz and abs(int(sz) / 2 - b["en_size_pt"]) > 0.1:
                    bad["size_en"].add(f"{int(sz) / 2:g}")
        snip = text[:50] + "…"
        for k, vals in bad.items():
            want = {"font_fa": b["fa_fonts"][0], "size_fa": b["fa_size_pt"],
                    "font_en": b["en_fonts"][0], "size_en": b["en_size_pt"]}[k]
            out.append(Issue(k, "error", f"یافت‌شده: {', '.join(map(str, vals))} — مطلوب: {want}", loc, snip, i))

        # فاصله‌ی خطوط
        line, lrule = None, "auto"
        for lvl in _levels(doc, p, "p"):
            sp = lvl.find(qn("w:spacing")) if lvl is not None else None
            if sp is not None and sp.get(qn("w:line")) is not None:
                line, lrule = int(sp.get(qn("w:line"))), sp.get(qn("w:lineRule")) or "auto"
                break
        if line is None:
            line, lrule = 240, "auto"
        if lrule == "auto":
            mult = line / 240
            target_spacing = b.get("line_spacing", 1.5)
            tol = b.get("line_spacing_tolerance", 0.2)
            if abs(mult - target_spacing) > tol:
                out.append(Issue("line_spacing", "warn",
                                 f"فاصله‌ی خطوط {mult:.2f} است؛ باید {target_spacing} باشد.", loc, snip, i))

        # تراز
        if b.get("alignment") == "justify" and has_fa(text):
            jc = eff_ppr(doc, p, "w:jc", "w:val")
            if jc not in ("both", "distribute", "thaiDistribute"):
                out.append(Issue("alignment", "warn", "متن دوطرفه (Justify) نیست.", loc, snip, i))

        # جهت
        if has_fa(text) and not eff_bidi(doc, p):
            out.append(Issue("rtl", "warn", "جهت پاراگراف راست‌به‌چپ (RTL) نیست.", loc, snip, i))
    return out


@bounded_findings
def check_fields(doc):
    out = FindingBuffer()
    instr = " ".join(e.text or "" for e in doc.element.body.iter(qn("w:instrText")))
    has_toc = "TOC" in instr
    if not has_toc:
        out.append(Issue("toc", "info",
                         "فهرست مطالب خودکار (References ▸ Table of Contents) پیدا نشد؛ اگر فهرست را دستی "
                         "نوشته‌اید، شماره‌صفحه‌ها با ویرایش به‌هم می‌خورد."))
    has_page = False
    for relationship in doc.part.rels.values():
        if relationship.is_external or not relationship.reltype.endswith(("/header", "/footer")):
            continue
        part = relationship.target_part
        x = getattr(part, "_element", None)
        if x is None:
            continue
        field_texts = [e.text or "" for e in x.iter(qn("w:instrText"))]
        field_texts += [e.get(qn("w:instr")) or "" for e in x.iter(qn("w:fldSimple"))]
        if any(value.strip().split() and value.strip().split()[0].upper() == "PAGE" for value in field_texts):
            has_page = True
    if not has_page:
        out.append(Issue("page_numbers", "warn", "شماره‌ی صفحه‌ی خودکار در سربرگ/پانویس پیدا نشد."))
    return out


# ----------------------------------------------------------------------------
# ساختار: بخش‌ها، فصل‌ها، چکیده، کپشن‌ها
# ----------------------------------------------------------------------------
ORD = {"اول": 1, "دوم": 2, "سوم": 3, "چهارم": 4, "پنجم": 5, "ششم": 6, "هفتم": 7, "هشتم": 8, "نهم": 9, "دهم": 10}
CHAPTER_RE = re.compile(r"^\s*فصل\s*(اول|دوم|سوم|چهارم|پنجم|ششم|هفتم|هشتم|نهم|دهم|[0-9۰-۹]+)")
CAPTION_RE = re.compile(r"^\s*(جدول|شکل|نمودار|تصویر)\s*(?:شماره\s*)?\(?([0-9۰-۹]+)(?:[\s\-–ـ\.]*([0-9۰-۹]+))?\)?")


def find_section(paragraphs, patterns):
    pats = [norm(x) for x in patterns]
    for i, p in enumerate(paragraphs):
        t = norm(p.text)
        if not t or (len(t.split()) > 10 and not is_heading(p)):
            continue
        if any(x in t for x in pats):
            return i
    return None


@bounded_findings
def check_structure(doc, rules, mode, stats):
    out = FindingBuffer()
    ps = doc.paragraphs
    m = rules["modes"].get(mode, rules["modes"].get("thesis", {}))

    for sec in m.get("required_sections", []):
        if find_section(ps, sec["patterns"]) is None:
            out.append(Issue("missing_section", "error", f"بخش «{sec['name']}» پیدا نشد."))
    for sec in m.get("recommended_sections", []):
        if find_section(ps, sec["patterns"]) is None:
            out.append(Issue("recommended_section", "info", f"بخش «{sec['name']}» پیدا نشد."))

    # ارزیابی چک‌لیست ۲۷گانه دانشگاه علوم پزشکی تهران
    if mode == "tums":
        checklist_status, tums_issues = tums_checker.check_tums_checklist(ps)
        stats["tums_checklist"] = checklist_status
        for ti in tums_issues:
            out.append(ti)

    # فصل‌ها و واژگان هر فصل
    nums, unstyled = [], []
    ch_indices = []
    for i, p in enumerate(ps):
        mm = CHAPTER_RE.match(norm(p.text).replace("ي", "ی"))
        if mm and len(p.text.split()) <= 12 and not is_toc_style(p) and "." * 4 not in p.text:
            v = mm.group(1)
            n = ORD.get(v) or int(to_latin_digits(v))
            nums.append(n)
            ch_indices.append((n, i))
            if not is_heading(p):
                unstyled.append(i)
    stats["chapters"] = len(nums)

    # شمارش کلمات هر فصل جهت بررسی سقف ۳۰ درصدی فصول ۱ و ۲ در TUMS
    ch_words = {}
    for idx, (ch_num, start_p) in enumerate(ch_indices):
        end_p = ch_indices[idx + 1][1] if idx + 1 < len(ch_indices) else len(ps)
        w_cnt = sum(len(words(p.text)) for p in ps[start_p:end_p])
        ch_words[ch_num] = w_cnt

    if mode == "tums":
        total_w = stats.get("words", 0)
        prop_issues = tums_checker.check_tums_chapter_proportions(ch_words, total_w)
        for pi in prop_issues:
            out.append(pi)

    if mode in ("thesis", "tums"):
        min_ch = rules.get("limits", {}).get("min_chapters_thesis", 5 if mode == "tums" else 3)
        if len(nums) < min_ch:
            out.append(Issue("chapters", "warn", f"تعداد {len(nums)} عنوان «فصل …» پیدا شد؛ پایان‌نامه استاندارد باید حداقل {min_ch} فصل کامل داشته باشد."))
        uniq = sorted(set(nums))
        if uniq and uniq != list(range(1, uniq[-1] + 1)):
            out.append(Issue("chapters", "warn", f"ترتیب/شماره‌ی فصل‌ها ناقص یا تکراری است: {nums}"))
    if unstyled:
        out.append(Issue("heading_style", "warn",
                         f"{len(unstyled)} عنوان فصل با استایل Heading ساخته نشده؛ این امر مانع ایجاد فهرست خودکار می‌شود.",
                         f"پاراگراف {unstyled[0] + 1}", ps[unstyled[0]].text[:50], unstyled[0]))

    # چکیده و کلیدواژه
    a_i = find_section(ps, ["چکیده", "چکیده فارسی"])
    k_i = find_section(ps, ["کلیدواژه", "کلید واژه", "واژگان کلیدی", "واژه های کلیدی", "کلمات کلیدی"])
    if a_i is not None:
        end = k_i if (k_i is not None and k_i > a_i) else min(a_i + 12, len(ps))
        n = sum(len(words(x.text)) for x in ps[a_i + 1:end])
        lo, hi = rules.get("limits", {}).get("abstract_words", [150, 350])
        if n and not lo <= n <= hi:
            out.append(Issue("abstract_length", "warn",
                             f"چکیده‌ی فارسی حدود {n} کلمه است؛ محدوده‌ی مطلوب {lo}–{hi} کلمه.",
                             f"پاراگراف {a_i + 1}", "", a_i))
    if k_i is not None:
        body = ps[k_i].text
        body = body.split(":", 1)[1] if ":" in body else (ps[k_i + 1].text if k_i + 1 < len(ps) else "")
        cnt = len([x for x in re.split(r"[،,؛;]", body) if x.strip()])
        lo, hi = rules.get("limits", {}).get("keywords_count", [3, 7])
        if cnt and not lo <= cnt <= hi:
            out.append(Issue("keywords_count", "warn",
                             f"{cnt} کلیدواژه یافت شد؛ تعداد مطلوب {lo}–{hi}.", f"پاراگراف {k_i + 1}", body[:50], k_i))

    # کپشن‌ها
    kinds = {"جدول": [], "شکل": [], "نمودار": []}
    for i, p in enumerate(ps):
        mm = CAPTION_RE.match(p.text)
        if mm and len(p.text.split()) < 40:
            raw_kind = mm.group(1)
            k = "شکل" if raw_kind == "تصویر" else raw_kind
            a = int(to_latin_digits(mm.group(2)))
            b = int(to_latin_digits(mm.group(3))) if mm.group(3) else None
            kinds[k].append((i, a, b))
    n_tables, n_imgs = len(doc.tables), len(doc.inline_shapes)
    stats["tables"], stats["images"] = n_tables, n_imgs
    if n_tables > len(kinds["جدول"]):
        out.append(Issue("caption_missing", "warn",
                         f"{n_tables} جدول در فایل هست ولی فقط {len(kinds['جدول'])} عنوان «جدول …» پیدا شد."))
    if n_imgs > len(kinds["شکل"]) + len(kinds["نمودار"]):
        out.append(Issue("caption_missing", "warn",
                         f"{n_imgs} تصویر در فایل هست ولی فقط {len(kinds['شکل']) + len(kinds['نمودار'])} "
                         "عنوان «شکل/نمودار …» پیدا شد."))
    for kind, items in kinds.items():
        last = {}
        for i, a, b in items:
            key = a if b is not None else 0
            seq = b if b is not None else a
            exp = last.get(key, 0) + 1
            if seq != exp:
                out.append(Issue("caption_numbering", "warn",
                                 f"{kind} {a}{'-' + str(b) if b else ''}: شماره‌ی بعدی باید {exp} می‌بود.",
                                 f"پاراگراف {i + 1}", ps[i].text[:50], i))
            last[key] = seq
    return out


# ----------------------------------------------------------------------------
# ارجاع‌دهی و منابع
# ----------------------------------------------------------------------------
REF_HEAD = ["منابع", "مراجع", "فهرست منابع", "references", "bibliography"]
AFTER_REFS_HEAD = [
    "پیوست", "پیوست‌ها", "ضمیمه", "ضمائم", "ضمایم", "appendix", "appendices",
    "چکیده انگلیسی", "abstract", "اظهارنامه", "تأییدیه", "صورتجلسه",
    "declaration", "فرم تعهد", "سپاسگزاری"
]
CITE_RE = re.compile(r"\(([^()\n]{2,90}?)[،,]\s*([0-9۰-۹]{4})[a-z]?(?:\s*[،,;؛][^()\n]{0,40})?\)")
NUM_CITE_RE = re.compile(r"\[([0-9٠-٩۰-۹]+(?:\s*[-–,،]\s*[0-9٠-٩۰-۹]+)*)\]")


def first_surname(s):
    s = norm(s)
    s = re.split(r"\s+(?:و|and|&|et al)\b|[،,;؛]", s)[0]
    return s.strip().split(" ")[0] if s.strip() else ""


def check_citations(doc, rules, mode="thesis"):
    from rule_registry import FindingBudgetExceeded
    try:
        return _check_citations(doc, rules, mode)
    except FindingBudgetExceeded as error:
        reference_index = next((i for i, p in enumerate(doc.paragraphs)
                                if any(norm(p.text) == norm(h) for h in REF_HEAD)), None)
        return error.findings + [Issue("findings_limit", "warn", "بررسی ارجاعات به سقف یافته‌ها رسید؛ بازبینی دستی لازم است.",
                                      confidence="B", extraction_status="analysis_limited")], reference_index


def _check_citations(doc, rules, mode="thesis"):
    """Separate candidate extraction, conservative matching, and review findings."""
    out, ref_i = FindingBuffer(), None
    paragraphs = doc.paragraphs
    for index in range(len(paragraphs)-1, -1, -1):
        text = norm(paragraphs[index].text)
        if any(text == norm(h) for h in REF_HEAD):
            ref_i = index
            break
    if ref_i is None:
        return out, None
    refs = []
    for index in range(ref_i+1,len(paragraphs)):
        text=paragraphs[index].text
        if any(norm(text)==norm(h) for h in AFTER_REFS_HEAD):
            break
        if text.strip():
            refs.append((index,text))
    if mode=="tums":
        out.extend(tums_checker.check_tums_references([text for _,text in refs]))

    normalized=[to_latin_digits(norm(text)) for _,text in refs]
    author_index = defaultdict(list)
    for reference_index, reference in enumerate(normalized):
        year_match = re.search(r"(?<![0-9])[0-9]{4}[a-z]?(?![0-9a-z])", reference)
        if year_match:
            author_tokens = set(re.findall(r"[\w\u200c'-]+", reference[:year_match.start()]))
            for author_token in author_tokens:
                author_index[(author_token, year_match.group())].append(reference_index)
    labels=[]
    for text in normalized:
        match=re.match(r"^\s*(?:\[(\d+)\]|(\d+)[.)])\s*",text)
        labels.append(int(match.group(1) or match.group(2)) if match else None)
    numbered=bool(labels) and all(n is not None for n in labels) and len(set(labels))==len(labels)
    author_year=re.compile(r"([^;؛()]+?)[,،]\s*([0-9٠-٩۰-۹]{4}[a-z]?)",re.IGNORECASE)
    narrative=re.compile(r"([A-Z][A-Za-z-]+(?:\s+et\s+al\.?)?)\s*\(([0-9٠-٩۰-۹]{4}[a-z]?)\)")
    for index,paragraph in enumerate(paragraphs[:ref_i]):
        text=paragraph.text
        candidates=[]
        for bracket in re.finditer(r"\([^()]*\)",text):
            for match in author_year.finditer(bracket.group()):
                candidates.append((match.group(1),match.group(2),
                                   bracket.start()+match.start(),bracket.start()+match.end()))
        for match in narrative.finditer(text):
            candidates.append((match.group(1),match.group(2),match.start(),match.end()))
        for author,year,left,right in candidates:
            surname=first_surname(author)
            year=to_latin_digits(year).lower()
            hits = author_index.get((surname, year), [])
            if len(hits)==1:
                continue
            rule="cite_ambiguous_ref" if hits else "cite_missing_ref"
            reason=("Multiple bibliography candidates; manual disambiguation required."
                    if hits else "No author/year candidate found in extracted references; review spelling, year suffix and bibliography structure.")
            out.append(Issue(rule,"warn",reason,f"Paragraph {index+1}",text[left:right],index,
                target_text=text[left:right],category="citation_consistency",
                confidence="C" if hits else "B",char_start=left,char_end=right,location_status="exact"))
        for match in NUM_CITE_RE.finditer(text):
            values=set()
            valid=True
            for item in re.split(r"[,،]",to_latin_digits(match.group(1))):
                bounds=re.split(r"[-–]",item)
                first,last=int(bounds[0]),int(bounds[-1])
                if first<=0 or last<first or last-first>1000:
                    valid=False
                    break
                values.update(range(first,last+1))
            if not valid:
                out.append(Issue("numeric_cite_ambiguous","info","Unreliable numeric range; manual review required.",
                                 f"Paragraph {index+1}",match.group(),index,confidence="C"))
                continue
            missing=values-set(labels) if numbered else {v for v in values if v>len(refs)}
            if missing and refs:
                out.append(Issue("numeric_cite","warn",
                    f"No extracted reference candidate for numbers {sorted(missing)}; bibliography numbering must be checked.",
                    f"Paragraph {index+1}",match.group(),index,target_text=match.group(),
                    category="citation_consistency",confidence="B",char_start=match.start(),
                    char_end=match.end(),location_status="exact"))
    return deduplicate_issues(out),ref_i


# ----------------------------------------------------------------------------
# ابزار یکپارچه‌سازی و رفع تکرار ایرادها
# ----------------------------------------------------------------------------
def deduplicate_issues(issues: List[Issue]) -> List[Issue]:
    """حذف هشدارهای تکراری یا همپوشان در یک محل واحد."""
    seen = set()
    deduped = []
    for it in issues:
        key = (it.block_id, it.source_part, it.loc, it.par_index, it.char_start, it.char_end,
               it.rule_id, it.target_text, it.message, it.suggestion)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(it)
    return deduped


# ----------------------------------------------------------------------------
# اجرای کامل روی Word (DOCX)
# ----------------------------------------------------------------------------
@immutable_input
def analyze(path, mode="thesis", rules=None):
    from document_blocks import extract_document_blocks
    started = time.monotonic()
    rules = _prepare_rules(rules, mode)
    validate_document(path, "docx", rules.get("max_input_bytes", 25 * 1024 * 1024))
    pre_hash = sha256_file(path)
    doc = Document(path)
    doc._review_source_path = os.path.abspath(path)
    doc._review_source_hash = pre_hash
    blocks = extract_document_blocks(doc, max_blocks=rules.get("max_document_blocks", 25000),
                                     max_total_chars=rules.get("max_document_chars", 4_000_000))
    body_blocks = {block.paragraph_index: block for block in blocks
                   if block.story == "body" and block.paragraph_index >= 0}
    issues, stats = [], {"mode": mode, "file_sha256": pre_hash, "ambiguous_candidates": 0}

    heading = "ابتدای سند"
    ctx = []
    paragraphs_texts = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        if is_heading(p) and t.strip():
            heading = t.strip()[:50]
        ctx.append((i, p, heading))
        paragraphs_texts.append(t)

    stats["paragraphs"] = len(ctx)
    stats["words"] = sum(len(words(p.text)) for p in doc.paragraphs)

    cit_issues, ref_i = check_citations(doc, rules, mode=mode)

    blank_run = 0
    analyzed_blocks = set()
    findings_budget = rules.get("max_findings", 10000)
    for i, p, h in ctx:
        block = body_blocks.get(i)
        text = block.text if block is not None else p.text
        if is_toc_style(p):
            continue
        if len(issues) >= findings_budget:
            break
        if not text.strip() and not p._p.xpath(".//w:drawing") and not p._p.xpath(".//w:br[@w:type='page']"):
            blank_run += 1
            if blank_run == 3:
                issues.append(Issue("blank_lines", "warn",
                                    "چند پاراگراف خالی پشت‌سرهم؛ برای فاصله از Spacing پاراگراف یا Page Break استفاده کنید.",
                                    f"زیر «{h}»، پاراگراف {i + 1}", "", i,
                                    category="۳. نشانه‌گذاری و فاصله‌گذاری",
                                    confidence=ConfidenceLevel.REVIEW_B.value))
            continue
        blank_run = 0
        in_refs = ref_i is not None and i > ref_i
        in_decl = any(kw in h.lower() for kw in ["تعهد", "اصالت", "سپاس", "تقدیم", "declaration"])
        found = check_text(i, text, f"زیر «{h}»، پاراگراف {i + 1}", rules, in_refs, in_declaration=in_decl, audit=stats)
        issues += _attach_block_issues(found, block) if block is not None else found
        if block is not None:
            analyzed_blocks.add(block.block_id)

    # بررسی و استخراج متن جدول‌ها (Tables Extraction)
    table_cells_count = 0
    for block in blocks:
        if block.story == "body" and block.paragraph_index >= 0:
            continue
        if not block.text.strip():
            continue
        if len(issues) >= findings_budget:
            break
        if block.story == "table":
            table_cells_count += 1
        found = check_text(-1, block.text, block.location, rules, audit=stats)
        issues += _attach_block_issues(found, block)
        analyzed_blocks.add(block.block_id)
    stats["tables"] = len(doc.tables)
    stats["table_cells_analyzed"] = table_cells_count

    # ارزیابی انسجام پاراگراف‌ها
    cohesion_issues = academic_checker.evaluate_cohesion(paragraphs_texts)
    for ci in cohesion_issues:
        issues.append(ci)

    issues += check_page(doc, rules, mode=mode)
    issues += check_body_format(doc, rules, ctx)
    issues += check_fields(doc)
    issues += check_structure(doc, rules, mode, stats)
    issues += cit_issues

    # راستی‌آزمایی عدم تغییر فایل ورودی در پایان پردازش
    post_hash = sha256_file(path)
    if pre_hash != post_hash:
        raise RuntimeError("نقض اصل تمامیت داده: فایل اصلی ورودی در حین پردازش تغییر کرده است!")

    for issue in issues:
        if issue.confidence == "A" and issue.location_status != "exact":
            issue.confidence = "B"
    stats["coverage_limits"] = [
        "وضعیت بخش‌های استخراج‌شده در جدول پوشش آمده است؛ فرمول‌ها و تغییرات ردگیری‌شده بررسی کامل ندارند. استخراج متن به معنی امکان درج امن کامنت نیست.",
        "شمارهٔ صفحهٔ Word از شمارهٔ پاراگراف قابل استنتاج نیست.",
        "تطبیق ارجاعات محتاطانه است و محتوای علمی منابع را تأیید نمی‌کند.",
        "نسخه و تاریخ همهٔ بندهای پروفایل دانشگاهی مستقلاً تأیید نشده است."
    ]
    stats["body_words"] = stats["words"]
    stats["words"] = sum(len(words(block.text)) for block in blocks)
    stats["block_count"] = len(blocks)
    stats["source_chars"] = sum(len(block.text) for block in blocks)
    stats["story_counts"] = dict(Counter(block.story for block in blocks))
    stats["partial_blocks"] = sum(block.extraction_status != "native" for block in blocks)
    stats["coverage"] = _document_coverage(blocks, analyzed_blocks, rules.get("max_analysis_block_chars", 200000))
    if len(issues) >= findings_budget:
        stats["skipped_blocks"] = sum(bool(block.text.strip()) and block.block_id not in analyzed_blocks for block in blocks)
        issues.append(Issue("findings_limit", "warn", "بودجهٔ یافته‌ها تکمیل شد؛ همهٔ بلوک‌های متن بررسی نشده‌اند.",
                            confidence="B", extraction_status="analysis_limited"))
    stats["protected_content_audit"] = _protected_audit(blocks, rules.get("max_analysis_block_chars", 200000))
    return doc, _finish_analysis(issues, stats, started, rules), stats


# ----------------------------------------------------------------------------
# اجرای بررسی روی فایل PDF
# ----------------------------------------------------------------------------
@immutable_input
def analyze_pdf(pdf_path, mode="thesis", rules=None):
    """تحلیل و استخراج متنی فایل PDF متناسب با اصول بند ۲ و ۱۰."""
    started = time.monotonic()
    rules = _prepare_rules(rules, mode)
    validate_document(pdf_path, "pdf", rules.get("max_input_bytes", 25 * 1024 * 1024))
    pre_hash = sha256_file(pdf_path)
    extracted = pdf_handler.extract_pdf_content(pdf_path, max_pages=rules.get("max_pdf_pages", 1000),
                                               max_chars=rules.get("max_document_chars", 4_000_000),
                                               max_blocks=rules.get("max_document_blocks", 25000))
    issues, stats = [], {"mode": mode, "is_pdf": True, "file_sha256": pre_hash, "ambiguous_candidates": 0}

    if not extracted.get("ok"):
        stats["analysis_status"] = "failed"
        stats["analysis_error_code"] = extracted.get("error_code", "read_error")
        issues.append(Issue("pdf_read_error", "error", extracted.get("error", "خطا در خواندن فایل PDF"),
                            confidence=ConfidenceLevel.HIGH_A.value))
        return None, _finish_analysis(issues, stats, started, rules), stats

    if extracted.get("no_text_document", extracted.get("is_scanned")):
        stats["analysis_status"] = "partial"
        issues.append(Issue("pdf_scanned_warning", "warn",
                            "متن قابل استخراج این PDF کم است؛ احتمال تصویرمحور بودن یا استخراج ناقص وجود دارد. طبق اصل ۱۰، استخراج متن با اطمینان کافی نیازمند سامانه OCR است و نمی‌توان درباره بخش‌های ناخوانا نظر قطعی داد.",
                            confidence=ConfidenceLevel.REVIEW_B.value))
        stats["is_scanned"] = extracted.get("is_scanned", False)
        stats["total_pages"] = extracted.get("total_pages", 0)
        stats["extraction_pages"] = [{"page": p["page_num"], "quality": p["quality"], "warnings": p["quality_warnings"]}
                                      for p in extracted.get("pages", [])]
        return None, _finish_analysis(issues, stats, started, rules), stats

    all_paragraphs = []
    total_words = 0
    par_idx = 0
    stats["skipped_blocks"] = 0

    for p_info in extracted.get("pages", []):
        p_num = p_info["page_num"]
        if p_info.get("quality") in {"image_only", "no_text"}:
            issues.append(Issue("pdf_page_unreadable", "warn", "متن قابل‌اعتماد این صفحه استخراج نشد؛ بررسی دستی یا OCR لازم است.",
                                f"صفحه {p_num}", confidence="B", extraction_status="no_reliable_text",
                                page_index=p_num - 1, story="pdf"))
        elif p_info.get("quality_warnings"):
            issues.append(Issue("pdf_extraction_review", "info", "ترتیب خواندن یا نویسه‌های استخراج‌شدهٔ این صفحه نیازمند بازبینی است.",
                                f"صفحه {p_num}", confidence="B", extraction_status="pdf_text_unverified",
                                page_index=p_num - 1, story="pdf"))
        for block in p_info.get("blocks", []):
            par_text = block["text"]
            all_paragraphs.append(par_text)
            w_cnt = len(words(par_text))
            total_words += w_cnt
            loc = f"صفحه {p_num}، پاراگراف {par_idx + 1}"
            if len(issues) >= rules.get("max_findings", 10000):
                stats["skipped_blocks"] += 1
                par_idx += 1
                continue
            found = check_text(-1, par_text, loc, rules, audit=stats)
            for issue in found:
                issue.block_id = block["block_id"]
                issue.source_part = f"pdf/page/{p_num}"
                issue.story = "pdf"
                issue.page_index = p_num - 1
                issue.paragraph_index = par_idx
            issues += found
            par_idx += 1

    stats["paragraphs"] = len(all_paragraphs)
    if stats["skipped_blocks"]:
        issues.append(Issue("findings_limit", "warn", "بودجهٔ یافته‌ها تکمیل شد؛ بعضی بلوک‌های PDF بررسی نشده‌اند.",
                            confidence="B", extraction_status="analysis_limited"))
    stats["words"] = total_words
    stats["total_pages"] = extracted.get("total_pages", 0)
    stats["block_count"] = par_idx
    stats["source_chars"] = extracted.get("total_chars", 0)
    stats["protected_content_audit"] = _protected_audit(all_paragraphs, rules.get("max_analysis_block_chars", 200000))
    stats["extraction_pages"] = [{"page": p["page_num"], "quality": p["quality"], "warnings": p["quality_warnings"]}
                                  for p in extracted.get("pages", [])]

    # ارزیابی انسجام
    cohesion_issues = academic_checker.evaluate_cohesion(all_paragraphs)
    for ci in cohesion_issues:
        issues.append(ci)

    # راستی‌آزمایی عدم تغییر فایل اصلی PDF
    post_hash = sha256_file(pdf_path)
    if pre_hash != post_hash:
        raise RuntimeError("نقض اصل تمامیت داده: فایل اصلی PDF تغییر یافته است!")

    for issue in issues:
        if issue.extraction_status == "native":
            issue.extraction_status = "pdf_reading_order_unverified"
        if issue.confidence == "A":
            issue.confidence = "B"
    stats["coverage_limits"] = ["ترتیب ستون‌ها و صحت نمادهای PDF مستقلاً تأیید نشده است؛ OCR انجام نشده و هر صفحه جداگانه ارزیابی می‌شود."]
    return None, _finish_analysis(issues, stats, started, rules), stats


# ----------------------------------------------------------------------------
# اجرای بررسی روی متن مستقیم ارسالی کاربر در چت
# ----------------------------------------------------------------------------
def analyze_text(raw_text, mode="thesis", rules=None):
    """بررسی متن مستقیم وارد شده در چت بله."""
    started = time.monotonic()
    rules = _prepare_rules(rules, mode)
    if len(raw_text) > rules.get("max_document_chars", 4_000_000):
        raise ValueError("Text exceeds its document analysis budget")
    issues, stats = [], {"mode": mode, "is_raw_text": True, "ambiguous_candidates": 0}
    paragraphs = [p for p in raw_text.split("\n\n") if p.strip()]
    if len(paragraphs) > rules.get("max_document_blocks", 25000):
        raise ValueError("Text exceeds its block analysis budget")

    stats["paragraphs"] = len(paragraphs)
    stats["words"] = len(words(raw_text))

    stats["skipped_blocks"] = 0
    for i, p in enumerate(paragraphs):
        if len(issues) >= rules.get("max_findings", 10000):
            stats["skipped_blocks"] += 1
            continue
        loc = f"پاراگراف {i + 1}"
        found = check_text(i, p, loc, rules, audit=stats)
        for issue in found:
            issue.block_id = f"text:p{i+1}"
            issue.story = "text"
            issue.paragraph_index = i
        issues += found

    if stats["skipped_blocks"]:
        issues.append(Issue("findings_limit", "warn", "بودجهٔ یافته‌ها تکمیل شد؛ بعضی بلوک‌ها بررسی نشده‌اند.",
                            confidence="B", extraction_status="analysis_limited"))

    cohesion_issues = academic_checker.evaluate_cohesion(paragraphs)
    for ci in cohesion_issues:
        issues.append(ci)

    stats["block_count"] = len(paragraphs)
    stats["source_chars"] = len(raw_text)
    stats["analysis_text_sha256"] = hashlib.sha256(raw_text.encode("utf-8")).hexdigest()
    stats["protected_content_audit"] = _protected_audit(paragraphs, rules.get("max_analysis_block_chars", 200000))
    return None, _finish_analysis(issues, stats, started, rules), stats


# ----------------------------------------------------------------------------
# تولید گزارش جامع و تفکیک‌شده (بر پایه سطوح قطعیت و شواهد قطعی)
# ----------------------------------------------------------------------------
def build_report(issues, stats, rules, examples=None):
    """Full text report by default; explicit preview limits are always disclosed."""
    if examples is not None and (isinstance(examples, bool) or not isinstance(examples, int) or examples < 0):
        raise ValueError("examples must be a nonnegative integer or None")
    issues = list(issues)
    mode_key = stats.get("mode", "thesis")
    mode_title = rules.get("modes", {}).get(mode_key, {}).get("title", mode_key)
    if mode_key == "tums":
        mode_title = "پایان‌نامه دانشگاه علوم پزشکی تهران (TUMS)"

    sev_count = Counter(it.severity for it in issues)
    conf_count = Counter(getattr(it, "confidence", "A") for it in issues)

    lines = [
        "══════════════════════════════════════════════════════════════",
        "  ویراستار پژوهش — گزارش بازبینی نگارشی و ساختاری",
        "══════════════════════════════════════════════════════════════",
        "",
        "الف) خلاصهٔ بررسی و مشخصات سند",
        "──────────────────────────────────────────────────────────────",
        f"• شیوه نامه / الگوی انتخابی: {mode_title}",
        f"• تعداد پاراگراف‌ها: {stats.get('paragraphs', 0)} | تعداد کلمات: {stats.get('words', 0)}",
    ]
    if stats.get("analysis_status") == "failed":
        lines.append("❌ بررسی سند تکمیل نشد؛ متن قابل‌اعتماد برای ارزیابی در دسترس نبود.")
    elif stats.get("analysis_status") == "partial":
        lines.append("⚠️ بررسی بخشی از سند انجام شد؛ بخش‌های بررسی‌نشده و محدودیت‌ها را در ادامه ببینید.")
    ambiguous = ambiguous_count(issues, stats, rules)
    if ambiguous is None:
        lines.append("• تعداد موارد مبهمِ کنارگذاشته‌شده ثبت نشده است؛ صفر فرض نشده‌اند.")
    elif stats.get("ambiguous_candidates", 0):
        lines.append(f"• {stats['ambiguous_candidates']} مورد مبهم در حالت سخت‌گیرانه وارد فهرست ایرادها نشده است؛ بررسی دستی متن همچنان لازم است.")
    if "total_pages" in stats:
        lines.append(f"• تعداد صفحات PDF: {stats['total_pages']}")
    if "tables" in stats or "images" in stats:
        t_cells = stats.get('table_cells_analyzed', 0)
        lines.append(f"• تعداد جدول‌های اصلی: {stats.get('tables', 0)} (بلوک‌های متن جدول بررسی‌شده: {t_cells}) | تعداد تصاویر: {stats.get('images', 0)}")
    for limitation in stats.get("coverage_limits", []):
        lines.append("• محدودیت بررسی: " + limitation)
    if "file_sha256" in stats:
        lines.append(f"• شناسه امنیتی یکپارچگی (SHA-256): {stats['file_sha256'][:16]}…")

    lines += [
        f"• شاخص شدت خطاها: {SEVERITY_ICON['error']} {sev_count['error']} مورد گزارش‌شده  |  "
        f"{SEVERITY_ICON['warn']} {sev_count['warn']} هشدار/اصلاحی  |  "
        f"{SEVERITY_ICON['info']} {sev_count['info']} پیشنهاد بهبود",
        f"• سطح اطمینان: A، اطمینان بالا: {conf_count['A']} | B، نیازمند بازبینی: {conf_count['B']} | C، مبهم: {ambiguous if ambiguous is not None else 'نامشخص'}",
        "• A/B/C درصد دقت نیستند؛ موارد C فقط برای بررسی دستی‌اند.",
        "",
        "ب) دسته‌بندی تفکیک‌شده ایرادها",
        "──────────────────────────────────────────────────────────────",
    ]

    # گروه‌بندی بر اساس دسته‌بندی‌های استاندارد
    cat_groups = defaultdict(list)
    for it in issues:
        cat = category_label(getattr(it, "category", "") or RULE_CATEGORIES.get(it.rule_id, "سایر موارد"))
        cat_groups[f"سطح {it.confidence} / {cat}"].append(it)

    for cat_name, cat_issues in sorted(cat_groups.items()):
        lines.append(f"📌 {cat_name}: {len(cat_issues)} مورد")

    annotation = safe_annotation_summary(stats.get("annotation_summary"))
    output_label = annotation_format_label(stats)
    records = {record.get("index"): record for record in (annotation or {}).get("records", [])}
    indices = {id(issue): index for index, issue in enumerate(issues)}
    if annotation:
        lines += ["", "کامنت‌های نسخهٔ بازبینی",
                  f"از {annotation.get('requested', len(issues))} مورد گزارش‌شده، {annotation.get('applied', 0)} مورد در {output_label} کامنت شد."]
        if not annotation.get("output_created"):
            lines.append("تولید امن نسخهٔ بازبینی تکمیل نشد؛ جزئیات یافته‌ها در همین گزارش آمده است.")
        if annotation.get("experimental"):
            lines.append("نشانه‌گذاری PDF آزمایشی است؛ نمایش یادداشت فارسی در PDFخوان شما باید بررسی شود.")
        for reason, count in annotation.get("reasons", {}).items():
            lines.append(f"• {ANNOTATION_REASONS.get(reason, reason)}: {count} مورد")
    if stats.get("coverage"):
        lines += ["", "پوشش بخش‌های سند (استخراج متن با امکان درج کامنت یکسان نیست)"]
        for story, info in stats["coverage"].items():
            lines.append(f"• {STORY_LABELS.get(story, story)}: {COVERAGE_LABELS.get(info.get('status'), 'نامشخص')}؛ "
                         f"{info.get('analyzed_blocks', 0)} از {info.get('present_blocks', 0)} بلوک متن بررسی شد.")

    # چک‌لیست دانشگاه علوم پزشکی تهران در صورت انتخاب TUMS
    if "tums_checklist" in stats and stats["tums_checklist"]:
        lines += [
            "",
            "ج) چک‌لیست انطباق با دستورالعمل دانشگاه علوم پزشکی تهران (TUMS)",
            "──────────────────────────────────────────────────────────────",
        ]
        for item in stats["tums_checklist"]:
            lines.append(f"  • {item['name']}: {item['status']}")

    lines += [
        "",
        "د) جزئیات ایرادها و پیشنهادهای اصلاحی (تفکیک سطح قطعیت)",
        "──────────────────────────────────────────────────────────────",
    ]

    if not issues:
        lines.append("در بخش‌های بررسی‌شده، موردی مطابق قواعد فعال شناسایی نشد.")
    else:
        issue_counter = 1
        for cat_name, cat_issues in sorted(cat_groups.items()):
            lines.append(f"\n【 {cat_name} 】")
            selected = cat_issues if examples is None else cat_issues[:examples]
            if examples is not None and len(selected) < len(cat_issues):
                lines.append(f"این نمایش به درخواست به {examples} نمونه محدود است؛ {len(cat_issues) - len(selected)} مورد این دسته در این نسخهٔ خلاصه نمایش داده نشده است.")
            for it in selected:
                rule_name = RULE_TITLES.get(it.rule_id, it.rule_id)
                conf_badge = f"【سطح {getattr(it, 'confidence', 'A')}】"
                row = f"[{issue_counter}] {SEVERITY_ICON.get(it.severity, 'ℹ️')} {it.rule_id} {conf_badge} — {rule_name}"
                if it.loc:
                    row += f"\n    📍 محل: {location_label(it.loc)}"
                row += f"\n    📝 شرح: {it.message}"
                if it.snippet:
                    row += f"\n    🔍 متن: «{it.snippet}»"
                if it.suggestion:
                    row += f"\n    💡 پیشنهاد اصلاح: «{it.suggestion}»"
                if it.confidence == "C":
                    row += "\n    بررسی دستی: این مورد مبهم است و به‌عنوان خطای قطعی گزارش نشده است."
                record = records.get(indices[id(it)])
                if record is not None and record.get("rule_id", it.rule_id) == it.rule_id:
                    state = f"در {output_label} کامنت شد" if record.get("status") == "applied" else ANNOTATION_REASONS.get(record.get("reason"), f"در {output_label} کامنت نشد")
                    row += f"\n    وضعیت {output_label}: {state}"
                lines.append(row + "\n")
                issue_counter += 1

    lines += [
        "",
        "هـ) تعهدنامه حفاظت از داده و راهنمای پژوهشگر",
        "──────────────────────────────────────────────────────────────",
        "۱. مطابق اصل تغییرناپذیری، فایل اصلی کاربر صرفاً خوانده شده و در مسیر تحلیل برای ذخیره‌سازی باز نمی‌شود و کنترل هش برای ورودی فایل فعال است.",
        "۲. پیشنهادها خودکار روی متن، فرمول‌ها، یکاها، ارقام و نام‌های خاص اعمال نشده‌اند؛ نسخهٔ بازبینی فقط نشانه‌گذاری دارد.",
        "۳. پیشنهادات سطح B و C صرفاً جنبه مشورتی داشته و تصمیم نهایی با نویسنده و استاد راهنما است.",
        "══════════════════════════════════════════════════════════════"
    ]
    lines.append("محدودیت‌ها: بررسی قاعده‌محور تأیید صحت علمی نیست؛ فرمول‌ها، تغییرات ردگیری‌شده و ترتیب خواندن PDF نیازمند بازبینی انسانی‌اند.")
    return "\n".join(lines)


def build_summary(issues, stats, rules):
    """A short student-facing receipt built after annotation has been validated."""
    issues = list(issues)
    conf = Counter(getattr(it, "confidence", "A") for it in issues)
    ambiguous = ambiguous_count(issues, stats, rules)
    annotation = safe_annotation_summary(stats.get("annotation_summary"))
    output_label = annotation_format_label(stats)
    title = ("❌ بررسی سند تکمیل نشد؛ علت در گزارش آمده است." if stats.get("analysis_status") == "failed"
             else "📄 نسخهٔ بازبینی آماده است." if annotation and annotation.get("output_created")
             else "📄 نتیجهٔ بررسی آماده است.")
    lines = [
        title,
        f"A، اطمینان بالا: {conf['A']} · B، نیازمند بازبینی: {conf['B']} · C، مبهم: {ambiguous if ambiguous is not None else 'نامشخص'}",
    ]
    if annotation:
        lines.append(f"از {annotation.get('requested', len(issues))} مورد گزارش‌شده، {annotation.get('applied', 0)} مورد در {output_label} کامنت شد؛ علت درج‌نشدن بقیه در گزارش آمده است.")
        if not annotation.get("output_created"):
            lines.append(f"تولید امن {output_label} تکمیل نشد؛ گزارش جایگزین را ببینید.")
    if stats.get("analysis_status") == "partial":
        lines.append("⚠️ بخشی از سند بررسی نشد یا پوشش کامل نداشت؛ جزئیات در گزارش است.")
    if not issues and stats.get("analysis_status") != "failed":
        lines.append("در بخش‌های بررسی‌شده، موردی مطابق قواعد فعال شناسایی نشد.")
    if ambiguous is None:
        lines.append("تعداد موارد مبهم ثبت نشده است؛ صفر فرض نشده‌اند.")
    lines.append("پیشنهادها خودکار اعمال نشده‌اند؛ موارد C فقط برای بررسی دستی‌اند.")
    if stats.get("report_available"):
        minutes = stats.get("report_retention_minutes")
        lines.append("گزارش کامل: /report" + (f" · تا {minutes} دقیقه در دسترس است." if minutes else ""))
    return "\n".join(lines)


def split_message(text, limit=3800):
    parts, cur = [], ""
    for block in text.split("\n\n"):
        if len(cur) + len(block) + 2 > limit and cur:
            parts.append(cur.rstrip())
            cur = ""
        while len(block) > limit:
            parts.append(block[:limit])
            block = block[limit:]
        cur += block + "\n\n"
    if cur.strip():
        parts.append(cur.rstrip())
    return parts


def annotate(doc, issues, out_path, max_comments=250, per_par=4):
    """
    ذخیره‌ی نسخه‌ی کامنت‌گذاری‌شده‌ی فایل کاربر در Word با موتور ایمن safe_annotator.
    قاعده اساسی: به هیچ عنوان به متن اصلی دست زده نمی‌شود (حفظ ۱۰۰٪ نویسه‌ها و استایل‌ها).
    """
    return safe_annotator.annotate_document(doc, issues, out_path, max_comments=max_comments, per_par=per_par)


if __name__ == "__main__":
    import sys
    from review_cli import main
    raise SystemExit(main(engine=sys.modules[__name__]))
