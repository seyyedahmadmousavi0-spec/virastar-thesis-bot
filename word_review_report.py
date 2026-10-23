"""A separate, complete Word review; the source document is never rewritten.

Some findings concern missing sections, protected content, or locations that
cannot safely accept a comment.  This companion artifact gives every returned
finding its own yellow entry and comment without inventing a source location.
"""
from collections import Counter
import gc
import hashlib
import io
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from doc_guard import (DocumentIntegrityError, atomic_write,
                       ensure_distinct_output, sha256_file)
from report_artifacts import (ANNOTATION_REASONS, CONFIDENCE_LABELS,
                              COVERAGE_LABELS, PRODUCT_VERSION, STORY_LABELS,
                              location_label, safe_annotation_summary,
                              safe_coverage, university_profile_lines)


# The analysis engine returns at most its 10,000 findings plus a limit notice.
# Oversize artifacts fail explicitly before publication, never lose entries.
MAX_REVIEW_FINDINGS = 10001
MAX_REVIEW_CHARACTERS = 16_000_000
MAX_REVIEW_BYTES = 25 * 1024 * 1024
TITLE = "گزارش کامل بازبینی"


class WordReviewReportError(ValueError):
    """A full Word review cannot be produced within the documented budget."""


def _xml_text(value):
    """Display forbidden XML characters visibly, preserving all valid Unicode.

    In particular, scientific symbols, Persian joining characters, accents,
    numbers and formulas are never normalized. The source is not modified.
    """
    result = []
    for character in str(value or ""):
        code = ord(character)
        if (code in (9, 10, 13) or 0x20 <= code <= 0xD7FF
                or 0xE000 <= code <= 0xFFFD or 0x10000 <= code <= 0x10FFFF):
            result.append(character)
        else:
            result.append(f"⟦U+{code:04X}⟧")
    return "".join(result)


def _rtl(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    if properties.find(qn("w:bidi")) is None:
        properties.append(OxmlElement("w:bidi"))
    return paragraph


def _run(paragraph, text, *, yellow=False, bold=False):
    run = paragraph.add_run(_xml_text(text))
    run.bold = bold
    properties = run._r.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:cs"), "Arial")
    if yellow:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _paragraph(document, text="", *, bold=False):
    paragraph = _rtl(document.add_paragraph())
    _run(paragraph, text, bold=bold)
    return paragraph


def _comment(document, anchor, text):
    comment = document.add_comment(anchor, text="", author="ویراستار پژوهش", initials="VP")
    paragraph = comment.paragraphs[0]
    _rtl(paragraph)
    _run(paragraph, text)
    return comment


def _short_comment(issue, index):
    # Do not truncate a qualifier, which could turn an uncertain claim into fact.
    reason = " ".join(str(issue.message or "").split())
    suggestion = " ".join(str(issue.suggestion or "").split())
    if not reason or len(reason) > 160:
        reason = "دلیل کامل در همین مدخل گزارش آمده است؛ آن را با متن اصلی بررسی کنید."
    if not suggestion or len(suggestion) > 160:
        suggestion = "پیشنهاد کامل همین مدخل را پیش از هر تغییر در متن اصلی بررسی کنید."
    return (f"مورد {index + 1} | {issue.rule_id} | {issue.confidence} — "
            f"{CONFIDENCE_LABELS[issue.confidence]}\n"
            f"دلیل: {reason}\nپیشنهاد: {suggestion}")


def _finding_id(issue, stats):
    identity = [stats.get("file_sha256") or stats.get("analysis_text_sha256")]
    identity.extend(getattr(issue, key, None) for key in
                    ("rule_id", "block_id", "source_part", "xml_path", "par_index",
                     "char_start", "char_end", "message"))
    # Match the other artifacts for valid text; represent lone surrogates only
    # for hashing so malformed input cannot become active XML or abort rendering.
    encoded = json.dumps(identity, ensure_ascii=False).encode("utf-8", errors="surrogatepass")
    return hashlib.sha256(encoded).hexdigest()[:20]


def _annotation_status(issue, index, records):
    record = records.get(index)
    if not record or any(record.get(key, getattr(issue, key)) != getattr(issue, key)
                         for key in ("rule_id", "block_id", "char_start", "char_end")):
        return "وضعیت درج کامنت در محل اصلی برای این مورد ثبت نشده است."
    if record.get("status") == "applied":
        if record.get("anchor_scope") == "paragraph":
            return "در نسخهٔ بازبینی سند اصلی، کل پاراگراف کامنت و زرد شد؛ این نشانه محل دقیق یک عبارت نیست."
        return "در محل تأییدشدهٔ نسخهٔ بازبینی سند اصلی کامنت شد."
    reason = ANNOTATION_REASONS.get(record.get("reason"), "محل امن برای کامنت‌گذاری تأیید نشد")
    return "در محل اصلی کامنت نشد: " + reason + ". این مدخل گزارش کامنت مستقل دارد."


def _issue_lines(issue, index, stats, records):
    lines = [f"محل در سند اصلی: {location_label(issue.loc)}",
             f"شناسهٔ نتیجه: {_finding_id(issue, stats)}",
             "دلیل: " + (issue.message or "دلیل متنی برای این مورد ثبت نشده است."),
             "پیشنهاد برای بازبینی: " + (issue.suggestion or "این مورد را با توجه به متن اصلی و بافت پژوهش بررسی کنید.")]
    # Only certified original witnesses are labelled exact. Unverified snippets
    # are context, never asserted to be a pinpointed defect in the source.
    if issue.location_status == "exact" and issue.target_text:
        lines.append("شاهد دقیق از متن اصلی: «" + issue.target_text + "»")
    elif issue.snippet:
        lines.append("متن پیرامونی ثبت‌شده؛ محل دقیق تأیید نشده: «" + issue.snippet + "»")
    else:
        lines.append("شاهد دقیق متنی برای این مورد ثبت نشده است؛ بررسی در سطح بخش یا سند است.")
    if issue.confidence == "C":
        lines.append("این مورد مبهم است؛ فقط بررسی دستی لازم است و خطای قطعی محسوب نمی‌شود.")
    if issue.protected_overlap:
        lines.append("این مورد با محتوای حفاظت‌شده تداخل دارد؛ اعداد، فرمول‌ها و محتوای علمی خودکار تغییر نمی‌کنند.")
    lines.append("وضعیت نسخهٔ بازبینی سند اصلی: " + _annotation_status(issue, index, records))
    return lines


def write_word_review(out_path, issues, stats, rules, *, source_path=None):
    """Atomically write one independent yellow entry and comment per finding.

    Returns a small delivery summary. The complete list is mandatory: an
    unsupported size, invalid report or source alias raises an explicit error
    rather than publishing a partial report. ``stats`` and findings are read-only.
    """
    destination = Path(out_path)
    if source_path is not None:
        ensure_distinct_output(source_path, destination)
    source_hash = sha256_file(source_path) if source_path is not None else None
    try:
        findings = []
        for issue in issues:
            if len(findings) >= MAX_REVIEW_FINDINGS:
                raise WordReviewReportError("تعداد موارد از سقف گزارش کامل Word بیشتر است؛ هیچ موردی حذف و هیچ گزارش ناقصی منتشر نشد.")
            findings.append(issue)
        annotation = safe_annotation_summary(stats.get("annotation_summary")) or {}
        records = {record.get("index"): record for record in annotation.get("records", [])}
        entries = []
        characters = 0
        for index, issue in enumerate(findings):
            title = (f"مورد {index + 1} | {issue.rule_id} | {issue.confidence} — "
                     f"{CONFIDENCE_LABELS[issue.confidence]}")
            lines = _issue_lines(issue, index, stats, records)
            comment = _short_comment(issue, index)
            characters += sum(len(_xml_text(value)) for value in (title, comment, *lines))
            if characters > MAX_REVIEW_CHARACTERS:
                raise WordReviewReportError("حجم متن از سقف گزارش کامل Word بیشتر است؛ هیچ موردی حذف و هیچ گزارش ناقصی منتشر نشد.")
            entries.append((title, lines, comment))

        document = Document()
        document.core_properties.title = TITLE
        document.core_properties.author = "ویراستار پژوهش"
        document.core_properties.subject = "گزارش مستقل همهٔ یافته‌های بازبینی؛ متن پایان‌نامه بازنویسی نشده است."
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(11)
        _paragraph(document, TITLE, bold=True)
        _paragraph(document, "نسخهٔ موتور: " + PRODUCT_VERSION)
        if source_hash is not None:
            _paragraph(document, "SHA-256 سند اصلی: " + source_hash)
        _paragraph(document, "این فایل، گزارش مستقل بازبینی است و نسخهٔ پایان‌نامه یا متن اصلاح‌شدهٔ آن نیست. "
                   "هر مورد گزارش، هایلایت زرد و کامنت مستقل دارد؛ این علامت‌گذاری به معنای تأیید خطا در محل نامطمئنِ متن اصلی نیست.")
        counts = Counter(issue.confidence for issue in findings)
        _paragraph(document, f"همهٔ {len(findings)} مورد گزارش‌شده در این فایل آمده‌اند؛ "
                   f"A، اطمینان بالا: {counts['A']} · B، نیازمند بازبینی: {counts['B']} · C، مبهم: {counts['C']}.")
        _paragraph(document, "پیشنهادها روی متن علمی اعمال نشده‌اند. موردهای C فقط برای بررسی دستی‌اند؛ "
                   "سطح اطمینان، درصد دقت نیست و بررسی قاعده‌محور صحت علمی پژوهش را تأیید نمی‌کند.")
        for line in university_profile_lines(stats):
            if line:
                _paragraph(document, line)
        if stats.get("analysis_status") in {"partial", "failed"}:
            _paragraph(document, "بررسی همهٔ محتوای سند تکمیل نشده است؛ این گزارش همهٔ موارد بازگردانده‌شدهٔ بررسی را شامل می‌شود.", bold=True)
        suppressed = stats.get("ambiguous_candidates", 0)
        if isinstance(suppressed, int) and not isinstance(suppressed, bool) and suppressed > 0:
            _paragraph(document, f"{suppressed} نامزد مبهم در حالت سخت‌گیرانه کنار گذاشته شده و جزئیات آن در خروجی تحلیل موجود نیست؛ "
                       "این تعداد با موردهای C دارای مدخل در این فایل یکسان نیست.")
        coverage = safe_coverage(stats.get("coverage"))
        for story, information in coverage.items():
            _paragraph(document, f"پوشش {STORY_LABELS[story]}: {COVERAGE_LABELS[information['status']]}")
        if not findings:
            _paragraph(document, "در بخش‌های بررسی‌شده، موردی مطابق قواعد فعال شناسایی نشد.")
        paragraph = anchor = None
        for title, lines, comment in entries:
            paragraph = _rtl(document.add_paragraph())
            paragraph.paragraph_format.keep_with_next = True
            anchor = _run(paragraph, title, yellow=True, bold=True)
            _comment(document, anchor, comment)
            for line in lines:
                _paragraph(document, line)

        output = io.BytesIO()
        document.save(output)
        payload = output.getvalue()
        # A large review must not retain two complete Word XML trees while it
        # is reopened for validation. Runs retain their document part too.
        del document, paragraph, anchor, output
        gc.collect()
        if len(payload) > MAX_REVIEW_BYTES:
            raise WordReviewReportError("حجم فایل از سقف گزارش کامل Word بیشتر است؛ هیچ گزارش ناقصی منتشر نشد.")
        reopened = Document(io.BytesIO(payload))
        comment_ids = [str(comment.comment_id) for comment in reopened.comments]
        start_ids = [node.get(qn("w:id")) for node in reopened.element.xpath("//w:commentRangeStart")]
        end_ids = [node.get(qn("w:id")) for node in reopened.element.xpath("//w:commentRangeEnd")]
        reference_ids = [node.get(qn("w:id")) for node in reopened.element.xpath("//w:commentReference")]
        yellow = [run for paragraph in reopened.paragraphs for run in paragraph.runs
                  if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        if (len(reopened.comments) != len(findings) or len(yellow) != len(findings)
                or len(set(comment_ids)) != len(findings)
                or start_ids != comment_ids or end_ids != comment_ids or reference_ids != comment_ids
                or [run.text for run in yellow] != [_xml_text(entry[0]) for entry in entries]
                or [comment.text for comment in reopened.comments] != [_xml_text(entry[2]) for entry in entries]):
            raise DocumentIntegrityError("Complete Word review validation failed")
        del reopened, yellow
        gc.collect()
        if source_path is not None and sha256_file(source_path) != source_hash:
            raise DocumentIntegrityError("Source changed while building Word review")
        atomic_write(destination, payload, source_path=source_path)
        return {"path": str(destination), "requested": len(findings), "comments": len(findings),
                "highlighted_entries": len(findings), "output_created": True}
    finally:
        if source_path is not None and sha256_file(source_path) != source_hash:
            raise DocumentIntegrityError("Source changed while building Word review")
