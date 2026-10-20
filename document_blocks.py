"""Read-only Word story extraction with source locations and coverage flags."""
from dataclasses import dataclass, field
import hashlib
from typing import Any

from lxml import etree
from docx.oxml.ns import qn

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_DELETED = {qn("w:del"), qn("w:moveFrom")}
_TRACKED = _DELETED | {qn("w:ins"), qn("w:moveTo")}
_FIELDS = {qn("w:fldChar"), qn("w:fldSimple"), qn("w:instrText")}
_MATH = {"{" + M + "}oMath", "{" + M + "}oMathPara"}
_ALTERNATE = "{" + MC + "}AlternateContent"


class DocumentExtractionLimit(ValueError):
    """A document cannot be fully extracted within the configured budgets."""


@dataclass(frozen=True)
class DocumentBlock:
    text: str
    block_id: str
    source_part: str
    xml_path: str
    story: str
    paragraph_index: int
    location: str
    extraction_status: str = "native"
    paragraph: Any = field(default=None, repr=False, compare=False)

    @property
    def original_text(self):
        return self.text

    @property
    def paragraph_element(self):
        return self.paragraph._p if self.paragraph is not None else None


def extract_document_blocks(doc, *, max_blocks=25000, max_total_chars=4000000,
                            max_part_bytes=16777216, max_nodes=1000000,
                            max_depth=96):
    """Return body, table, header/footer, note and textbox paragraphs.

    Direct body paragraphs keep python-docx indices; extras have index -1.
    Whitespace, tabs and breaks are retained. Deleted/moved-from text, field
    instructions, formulas and note separators are excluded. Field results and
    inserted/moved-to text remain with partial coverage flags. AlternateContent
    selects its first Choice, or Fallback, once and flags unverified support.
    Existing parts are read without accessing section header/footer properties,
    which could create parts. Limits raise instead of silently truncating.
    No layout or page numbers are inferred.
    """
    limits = (max_blocks, max_total_chars, max_part_bytes, max_nodes, max_depth)
    if any(not isinstance(value, int) or isinstance(value, bool) or value <= 0 for value in limits):
        raise ValueError("Extraction budgets must be positive integers")
    blocks = []
    total_chars = 0
    nodes_seen = 0
    body_paragraphs = {id(p._p): (i, p) for i, p in enumerate(doc.paragraphs)}

    def tick(depth):
        nonlocal nodes_seen
        nodes_seen += 1
        if depth > max_depth or nodes_seen > max_nodes:
            raise DocumentExtractionLimit("Document XML traversal budget exceeded")

    def selected_children(node):
        if node.tag == _ALTERNATE:
            choices = node.findall("{" + MC + "}Choice")
            fallback = node.find("{" + MC + "}Fallback")
            return list(choices[0]) if choices else (list(fallback) if fallback is not None else [])
        return list(node)

    def inherited_flags(node):
        flags = set()
        parent = node.getparent()
        while parent is not None:
            if parent.tag in _TRACKED:
                flags.add("tracked_changes")
            if parent.tag == _ALTERNATE:
                flags.add("alternate_content")
            parent = parent.getparent()
        return flags

    def paragraph_content(paragraph, depth):
        pieces = []
        flags = inherited_flags(paragraph)

        def collect(node, level):
            tick(level)
            tag = node.tag
            if tag in _TRACKED:
                flags.add("tracked_changes")
            if tag in _DELETED or tag == qn("w:delText"):
                return
            if tag == qn("w:txbxContent"):
                return
            if tag in _FIELDS:
                flags.add("fields")
            if tag in _MATH:
                flags.add("math")
                return
            if tag == _ALTERNATE:
                flags.add("alternate_content")
            if tag == qn("w:t"):
                pieces.append(node.text or "")
                return
            if tag == qn("w:tab"):
                pieces.append("\t")
            elif tag in {qn("w:br"), qn("w:cr")}:
                pieces.append("\n")
            elif tag == qn("w:noBreakHyphen"):
                pieces.append("\u2011")
            elif tag == qn("w:softHyphen"):
                pieces.append("\u00ad")
            for child in selected_children(node):
                collect(child, level + 1)

        collect(paragraph, depth)
        return "".join(pieces), flags

    def add_paragraph(node, part_name, story, location, depth):
        nonlocal total_chars
        text, flags = paragraph_content(node, depth)
        index, paragraph = body_paragraphs.get(id(node), (-1, None)) if story == "body" else (-1, None)
        xml_path = node.getroottree().getpath(node)
        block_id = hashlib.sha256((part_name + "\0" + xml_path).encode("utf-8")).hexdigest()[:24]
        if len(blocks) >= max_blocks or total_chars + len(text) > max_total_chars:
            raise DocumentExtractionLimit("Document block/text extraction budget exceeded")
        total_chars += len(text)
        blocks.append(DocumentBlock(text=text, block_id=block_id, source_part=part_name,
                                    xml_path=xml_path, story=story, paragraph_index=index,
                                    location=location,
                                    extraction_status=("partial:" + ",".join(sorted(flags))) if flags else "native",
                                    paragraph=paragraph))

    def textboxes_in(node, part_name, location, depth):
        tick(depth)
        if node.tag in _DELETED:
            return
        if node.tag == qn("w:txbxContent"):
            walk(node, part_name, "textbox", location + " / textbox", depth + 1)
            return
        for child in selected_children(node):
            textboxes_in(child, part_name, location, depth + 1)

    def walk(container, part_name, story, location, depth=0):
        tick(depth)
        counts = {}
        for node in selected_children(container):
            if not isinstance(node.tag, str) or node.tag in _DELETED:
                continue
            local = etree.QName(node).localname
            counts[local] = counts.get(local, 0) + 1
            position = counts[local]
            if node.tag == qn("w:p"):
                paragraph_location = location + " / paragraph " + str(position)
                add_paragraph(node, part_name, story, paragraph_location, depth + 1)
                textboxes_in(node, part_name, paragraph_location, depth + 1)
            elif node.tag == qn("w:tbl"):
                next_story = "table" if story == "body" else story
                walk(node, part_name, next_story, location + " / table " + str(position), depth + 1)
            elif node.tag == qn("w:tr"):
                walk(node, part_name, story, location + " / row " + str(position), depth + 1)
            elif node.tag == qn("w:tc"):
                # Physical tc XML avoids row.cells' repeated merged-cell proxies.
                walk(node, part_name, story, location + " / cell " + str(position), depth + 1)
            elif node.tag == qn("w:txbxContent"):
                walk(node, part_name, "textbox", location + " / textbox " + str(position), depth + 1)
            elif node.tag not in {qn("w:sectPr"), qn("w:tblPr"), qn("w:tblGrid"), qn("w:trPr"), qn("w:tcPr")}:
                walk(node, part_name, story, location, depth + 1)

    main_part = str(doc.part.partname)
    if len(doc.part.blob) > max_part_bytes:
        raise DocumentExtractionLimit("Main document part exceeds the extraction budget")
    walk(doc.element.body, main_part, "body", "body")
    seen_parts = {main_part}
    for part in sorted(doc.part.package.parts, key=lambda item: str(item.partname)):
        part_name = str(part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        kind = str(getattr(part, "content_type", "")).rsplit(".", 1)[-1]
        stories = {"header+xml": "header", "footer+xml": "footer",
                   "footnotes+xml": "footnote", "endnotes+xml": "endnote"}
        story = stories.get(kind)
        if story is None:
            continue
        blob = part.blob
        if len(blob) > max_part_bytes:
            raise DocumentExtractionLimit("Word story part exceeds the extraction budget")
        root = getattr(part, "element", None)
        if root is None:
            parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
            root = etree.fromstring(blob, parser=parser)
        if story in {"footnote", "endnote"}:
            note_tag = "{" + W + "}" + story
            for note in root:
                if note.tag != note_tag or note.get(qn("w:type"), "normal") != "normal":
                    continue
                note_id = note.get(qn("w:id"), "?")
                walk(note, part_name, story, story + " " + note_id)
        else:
            walk(root, part_name, story, story + " " + part_name.rsplit("/", 1)[-1])
    return blocks